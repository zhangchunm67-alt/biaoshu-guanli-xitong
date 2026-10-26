"""Image review flows directly into text review without industry selection."""
import base64
import io
import json
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path

from docx import Document
from PIL import Image

import engine
import server


class TextReviewFlowTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='bid-text-flow-')
        self.root = Path(self.temp.name)
        self.httpd = server.create_server(0, self.root, seed=False)
        self.store = self.httpd.store
        # Mirror the current library: images plus an industry-free text category.
        for category in self.store.all('category'):
            category['enabled'] = category['id'] == 'company_profile'
            self.store.put('category', category)
        self.store.db.execute("DELETE FROM objects WHERE kind='industry'")
        self.store.put('category', dict(id='text_template', name='文书模板',
            layer='special', industry_id='', aliases=['投标函'], enabled=True, sort_order=100))
        self.store.db.commit()
        picture = io.BytesIO()
        Image.new('RGB', (150, 100), 'navy').save(picture, format='PNG')
        self.store.material(dict(name='公司介绍图片', category_id='company_profile',
            subject=engine.EXPECTED_SUBJECT, mode='page', valid_until='长期',
            files=[dict(name='company.png', data=base64.b64encode(picture.getvalue()).decode())]))
        self.text = self.store.material(dict(name='投标函模板', category_id='text_template',
            subject=engine.EXPECTED_SUBJECT, mode='text', valid_until='长期',
            content='文字材料回归正文'))
        doc = Document()
        doc.add_heading('公司简介', 1)
        doc.add_paragraph('公司简介原文。')
        doc.add_heading('投标函', 1)
        doc.add_paragraph('投标函原文。')
        source = io.BytesIO()
        doc.save(source)
        self.project = self.store.create_project(dict(name='四步流程验证', filename='framework.docx',
            data=base64.b64encode(source.getvalue()).decode()))
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()
        self.base = f'http://127.0.0.1:{self.httpd.server_port}'
        self.opener = urllib.request.build_opener(urllib.request.ProxyHandler({}))

    def tearDown(self):
        self.httpd.shutdown()
        self.thread.join(timeout=5)
        self.httpd.server_close()
        self.store.close()
        self.temp.cleanup()

    def prepare(self):
        request = urllib.request.Request(self.base + '/api/projects/' + self.project['id'] + '/text-materials',
            data=b'{}', headers={'Content-Type': 'application/json', 'Origin': self.base}, method='POST')
        with self.opener.open(request, timeout=10) as response:
            return json.load(response)

    def test_direct_text_step_loads_library_and_exports_both_material_types(self):
        self.assertEqual([], self.store.all('industry'))
        common = self.project['mappings']
        with self.assertRaisesRegex(ValueError, '文字材料核对'):
            self.store.prepare_export(self.project['id'], {'mode': 'draft'})
        project = self.prepare()
        self.assertEqual(common, [r for r in project['mappings'] if r['layer'] == 'common'])
        self.assertTrue(project['special_selection_confirmed'])
        self.assertEqual(['text_template'], project['special_category_ids'])
        rows = [r for r in project['mappings'] if r['layer'] == 'special']
        self.assertEqual(1, len(rows))
        self.assertEqual('投标函', rows[0]['title'])
        self.store.update_mappings(project['id'], {'mappings': [dict(id=rows[0]['id'],
            action='fill', material_ids=[self.text['id']], note='保留文字核对备注')]})
        template = self.store.path(self.store.require('project', project['id'])['template_storage'])
        original = template.read_bytes()
        exported = self.store.export(project['id'], {'mode': 'draft'})
        result = Document(self.store.path('exports/' + exported['id'] + '.docx'))
        self.assertEqual(2, exported['filled_count'])
        self.assertEqual(1, len(result.inline_shapes))
        self.assertEqual(1, [p.text for p in result.paragraphs].count('文字材料回归正文'))
        self.assertEqual(original, template.read_bytes())

    def test_reentering_text_review_preserves_selections_and_notes(self):
        first = self.prepare()
        row = next(r for r in first['mappings'] if r['layer'] == 'special')
        saved = self.store.update_mappings(first['id'], {'mappings': [dict(id=row['id'],
            action='fill', material_ids=[self.text['id']], note='商务已核对')]})
        second = self.prepare()
        self.assertEqual(saved['mappings'], second['mappings'])
        self.assertEqual(1, len([r for r in second['mappings'] if r['layer'] == 'special']))

    def test_empty_text_library_categories_still_allow_the_text_step_and_export(self):
        category = self.store.require('category', 'text_template')
        category['enabled'] = False
        self.store.put('category', category)
        self.store.db.commit()
        project = self.prepare()
        self.assertEqual([], project['special_category_ids'])
        self.assertTrue(project['special_selection_confirmed'])
        self.assertFalse(any(r['layer'] == 'special' for r in project['mappings']))
        self.assertEqual(1, self.store.export(project['id'], {'mode': 'draft'})['filled_count'])

    def test_empty_legacy_industry_columns_are_not_text_review_items(self):
        for category in self.store.all('category'):
            if category['layer'] == 'special' and category['id'] != 'text_template':
                category['enabled'] = True
                self.store.put('category', category)
        self.store.db.commit()
        self.assertEqual(['text_template'], self.prepare()['special_category_ids'])

    def test_blocked_image_material_still_prevents_advancing(self):
        row = self.project['mappings'][0]
        self.store.update_mappings(self.project['id'], {'mappings': [dict(id=row['id'],
            action='fill', material_ids=[])]})
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.prepare()
        self.assertEqual(400, error.exception.code)
        self.assertIn('通用材料核对', json.load(error.exception)['error'])
        self.assertFalse(self.store.require('project', self.project['id']).get('special_selection_confirmed'))


if __name__ == '__main__':
    unittest.main()
