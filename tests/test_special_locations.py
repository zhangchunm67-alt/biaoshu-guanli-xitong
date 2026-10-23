"""Special evidence recognition, actual export anchors and legacy project upgrades."""
import base64
import copy
import io
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import engine
from server import Store, SUBJECT


class SpecialLocationTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='bid-special-locator-test-')
        self.root = Path(self.tmp.name)
        self.store = Store(self.root / 'data', seed=False)
        self.categories = self.store.all('category')

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def scan(self, doc, selected=('I01', 'I02', 'I03', 'I04')):
        path = self.root / 'template.docx'
        doc.save(path)
        return engine.scan_document(path, self.categories, [], selected)

    @staticmethod
    def rows(result):
        return {r['category_id']: r for r in result['mappings'] if r['layer'] == 'special'}

    def test_real_title_variants_and_export_after_form_notes(self):
        doc = Document()
        doc.add_heading('第六章 投标文件格式', 1)
        doc.add_heading('3.同类项目实施经验', 2)
        doc.add_paragraph('同类业绩汇总表')
        doc.add_table(rows=2, cols=2).cell(0, 0).text = '合同名称'
        doc.add_paragraph('（1）按照项目顺序，附上对应的合同证明材料。')
        doc.add_heading('1.拟派服务团队成员汇总表', 2)
        doc.add_table(rows=2, cols=2)
        doc.add_heading('2.拟派团队成员简历表', 2)
        doc.add_table(rows=2, cols=2)
        note = '1.根据团队要求，先附简历，再附上相应人员的资质、经验和证明材料。'
        doc.add_paragraph(note)
        doc.add_heading('其他材料', 2)
        result = self.scan(doc)
        rows = self.rows(result)
        self.assertEqual('同类业绩汇总表', rows['I01']['title'])
        self.assertEqual('2.拟派团队成员简历表', rows['I02']['title'])
        self.assertEqual('not_found', rows['I03']['location_match']['status'])
        records, _ = engine._inspect(doc, self.categories)
        record = next(r for r in records if r['id'] == rows['I02']['locator_id'])
        self.assertEqual(note, engine._text(record['_anchor']))
        material = {'id': 'evidence', 'name': '人员证明', 'category_id': 'I02',
                    'subject': SUBJECT, 'mode': 'text', 'content': '插入的人员材料', 'files': [],
                    'valid_until': '长期', 'years': []}
        row = dict(rows['I02'], material_ids=['evidence'], action='fill', confirmed=True)
        output = self.root / 'export.docx'
        engine.export_document(self.root / 'template.docx', output, [row], [material],
                               'reviewed', categories=self.categories)
        output_doc = Document(output)
        inserted = next(p for p in output_doc.paragraphs if p.text == '插入的人员材料')
        self.assertEqual(note, engine._text(inserted._p.getprevious()))
        self.assertEqual('其他材料', engine._text(inserted._p.getnext()))

    def test_directory_scoring_contract_prose_and_table_cells_are_not_targets(self):
        doc = Document()
        doc.add_paragraph('同类业绩汇总表\t20', style='TOC Heading')
        doc.add_heading('第三章 评标办法', 1)
        doc.add_heading('保险项目业绩与案例', 2)
        doc.add_heading('第四章 合同条款及格式', 1)
        doc.add_heading('人员简历表', 2)
        doc.add_heading('第六章 投标文件格式', 1)
        doc.add_paragraph('应提交同类业绩汇总表及证明材料，不得遗漏。')
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '人员简历表'
        doc.add_paragraph('同类业绩汇总表')
        doc.add_heading('拟派团队成员简历表', 2)
        rows = self.rows(self.scan(doc))
        self.assertEqual('同类业绩汇总表', rows['I01']['title'])
        self.assertEqual('拟派团队成员简历表', rows['I02']['title'])
        self.assertEqual(1, len(rows['I01']['location_match']['candidates']))
        self.assertEqual(1, len(rows['I02']['location_match']['candidates']))

    def test_all_industries_share_variants_but_not_generic_qualification_parents(self):
        for prefix in 'IFBW':
            with self.subTest(prefix=prefix):
                doc = Document()
                for title in ['类似项目业绩一览表', '项目负责人简历表',
                              '广播电视节目制作经营许可证', '平台代理授权证明']:
                    doc.add_paragraph(title)
                ids = tuple(prefix + f'{i:02}' for i in range(1, 5))
                rows = self.rows(self.scan(doc, ids))
                self.assertTrue(all(r['locator_id'] and r['location_match']['status'] == 'matched'
                                    for r in rows.values()))
        doc = Document()
        for title in ['资格证明文件', '营业执照', 'ISO相关资质证书', '授权委托书']:
            doc.add_heading(title, 1)
        rows = self.rows(self.scan(doc))
        self.assertFalse(rows['I03']['locator_id'])
        self.assertFalse(rows['I04']['locator_id'])

    def test_repeated_targets_and_shared_industry_targets_are_ambiguous(self):
        doc = Document()
        doc.add_heading('同类业绩汇总表', 1)
        doc.add_paragraph('第一组')
        doc.add_heading('同类业绩汇总表', 1)
        rows = self.rows(self.scan(doc))
        self.assertFalse(rows['I01']['locator_id'])
        self.assertEqual('ambiguous', rows['I01']['location_match']['status'])
        self.assertEqual(2, len(rows['I01']['location_match']['candidates']))
        doc = Document()
        doc.add_heading('同类业绩汇总表', 1)
        rows = self.rows(self.scan(doc, ('I01', 'B01')))
        self.assertTrue(all(not r['locator_id'] for r in rows.values()))
        self.assertTrue(all(r['location_match']['status'] == 'ambiguous' for r in rows.values()))

    def test_explicit_industry_titles_and_table_tokens(self):
        doc = Document()
        doc.add_heading('银行同类业绩汇总表', 1)
        doc.add_heading('保险同类业绩汇总表', 1)
        doc.add_table(rows=1, cols=1).cell(0, 0).text = '{{保险项目业绩与案例}}'
        rows = self.rows(self.scan(doc, ('I01', 'B01')))
        self.assertEqual('银行同类业绩汇总表', rows['B01']['title'])
        self.assertTrue(rows['I01']['locator_id'].startswith('ph:'))
        self.assertNotEqual(rows['B01']['locator_id'], rows['I01']['locator_id'])

    def test_legacy_backfill_preserves_materials_notes_manual_positions_and_confirmations(self):
        doc = Document()
        doc.add_heading('公司介绍', 1)
        doc.add_paragraph('通用正文')
        doc.add_paragraph('同类业绩汇总表')
        doc.add_heading('拟派团队成员简历表', 1)
        doc.add_heading('补充说明', 1)
        raw = io.BytesIO()
        doc.save(raw)
        p = self.store.create_project({'name': '旧项目', 'filename': 'template.docx',
                                      'data': base64.b64encode(raw.getvalue()).decode()})
        self.store.update_mappings(p['id'], {'mappings': [dict(id=r['id'], action='skip', confirmed=True)
                                                         for r in p['mappings']]})
        p = self.store.select_special_categories(p['id'], {'industry_ids': ['insurance']})
        material = self.store.material({'name': '已选业绩', 'category_id': 'I01', 'subject': SUBJECT,
                                       'mode': 'text', 'content': '已选材料正文', 'valid_until': '长期'})
        initial = {r['category_id']: r for r in p['mappings']}
        manual = next(h['id'] for h in p['headings'] if h['title'] == '补充说明')
        self.store.update_mappings(p['id'], {'mappings': [
            dict(id=initial['I02']['id'], locator_id=manual, action='skip', confirmed=True, note='人工位置'),
            dict(id=initial['I03']['id'], action='skip', confirmed=True)]})
        legacy = self.store.require('project', p['id'])
        common = copy.deepcopy([r for r in legacy['mappings'] if r['layer'] == 'common'])
        rows = {r['category_id']: r for r in legacy['mappings']}
        rows['I01'].update(locator_id='', material_ids=[material['id']], note='保留备注')
        rows['I01'].pop('location_match', None)
        legacy.pop('special_locator_version')
        self.store.put('project', legacy)
        self.store.migrate_project(legacy)
        upgraded = self.store.public_project(self.store.require('project', p['id']))
        rows = {r['category_id']: r for r in upgraded['mappings']}
        self.assertEqual('同类业绩汇总表', rows['I01']['title'])
        self.assertTrue(rows['I01']['locator_id'])
        self.assertEqual([material['id']], rows['I01']['material_ids'])
        self.assertEqual('保留备注', rows['I01']['note'])
        self.assertFalse(rows['I01']['confirmed'])
        self.assertEqual(manual, rows['I02']['locator_id'])
        self.assertTrue(rows['I02']['confirmed'])
        self.assertTrue(rows['I03']['confirmed'])
        self.assertEqual(common, [r for r in upgraded['mappings'] if r['layer'] == 'common'])
        self.assertEqual(upgraded, self.store.public_project(self.store.require('project', p['id'])))
        cleared = self.store.update_mappings(p['id'], {'mappings': [dict(id=rows['I01']['id'], locator_id='')]})
        again = self.store.public_project(self.store.require('project', p['id']))
        self.assertEqual('', next(r for r in again['mappings'] if r['category_id'] == 'I01')['locator_id'])


if __name__ == '__main__':
    unittest.main()
