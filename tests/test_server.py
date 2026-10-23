"""持久化、文件边界及同源保护测试；全部使用临时数据目录。"""
import base64
import io
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
import zipfile

from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import server


def image_upload(name='凭证.png', color='white'):
    raw = io.BytesIO()
    Image.new('RGB', (40, 20), color).save(raw, 'PNG')
    return {'name': name, 'data': base64.b64encode(raw.getvalue()).decode()}


class ServerTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='bid-server-test-')
        self.httpd = server.create_server(port=0, data_dir=self.tmp.name)
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()
        self.base = f'http://127.0.0.1:{self.httpd.server_port}'

    def tearDown(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join()
        self.httpd.store.close()
        self.tmp.cleanup()

    def request(self, method, path, payload=None, headers=None, raw=False):
        values = {'Content-Type': 'application/json', 'Origin': self.base}
        values.update(headers or {})
        req = urllib.request.Request(self.base + path, data=json.dumps(payload).encode() if payload is not None else None, method=method, headers=values)
        with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(req, timeout=10) as response:
            body = response.read()
            return (body if raw else json.loads(body)), dict(response.headers)

    def profile_project(self):
        from docx import Document
        doc = Document()
        doc.add_heading('公司介绍', 1)
        doc.add_heading('营业执照', 1)
        raw = io.BytesIO()
        doc.save(raw)
        project, _ = self.request('POST', '/api/projects', {
            'name': '材料联动测试', 'filename': 'template.docx',
            'data': base64.b64encode(raw.getvalue()).decode()})
        return project

    def profile(self, **overrides):
        body = {'name': '公司介绍', 'subject': server.SUBJECT,
                'category_id': 'company_profile', 'mode': 'text', 'content': '公司业务介绍'}
        body.update(overrides)
        return self.request('POST', '/api/materials', body)[0]

    def test_upload_fills_confirmed_missing_and_preserves_other_confirmations(self):
        p = self.profile_project()
        rows = [dict(m, confirmed=True) for m in p['mappings']]
        self.request('PATCH', '/api/projects/' + p['id'] + '/mappings', {'mappings': rows})
        material = self.profile()
        p, _ = self.request('GET', '/api/projects/' + p['id'])
        row = next(m for m in p['mappings'] if m['category_id'] == 'company_profile')
        self.assertEqual(('fill', 'ready', False), (row['action'], row['status'], row['confirmed']))
        self.assertEqual([material['id']], row['material_ids'])
        self.assertTrue(all(m['confirmed'] for m in p['mappings'] if m['id'] != row['id']))
        self.assertEqual([], row['issues'])
        self.request('PATCH', '/api/projects/' + p['id'] + '/mappings', {'mappings': [dict(row, confirmed=True)]})
        self.request('PATCH', '/api/materials/' + material['id'], {'content': '修订介绍'})
        p, _ = self.request('GET', '/api/projects/' + p['id'])
        self.assertFalse(next(m for m in p['mappings'] if m['id'] == row['id'])['confirmed'])

    def test_manual_skip_missing_and_selected_material_are_preserved(self):
        for action in ('skip', 'missing'):
            p = self.profile_project()
            row = next(m for m in p['mappings'] if m['category_id'] == 'company_profile')
            # Explicitly change the action before returning to missing, to record manual intent.
            self.request('PATCH', '/api/projects/' + p['id'] + '/mappings', {'mappings': [dict(row, action='skip')]})
            self.request('PATCH', '/api/projects/' + p['id'] + '/mappings', {'mappings': [dict(row, action=action)]})
            self.profile(name='后补材料' + action)
            updated, _ = self.request('GET', '/api/projects/' + p['id'])
            result = next(m for m in updated['mappings'] if m['id'] == row['id'])
            self.assertEqual(action, result['action'])
        p = self.profile_project()
        row = next(m for m in p['mappings'] if m['category_id'] == 'company_profile')
        chosen = self.httpd.store.all('material')[0]['id']
        self.request('PATCH', '/api/projects/' + p['id'] + '/mappings', {'mappings': [dict(row, action='fill', material_ids=[chosen], confirmed=True)]})
        self.profile(name='新版本')
        updated, _ = self.request('GET', '/api/projects/' + p['id'])
        result = next(m for m in updated['mappings'] if m['id'] == row['id'])
        self.assertEqual([chosen], result['material_ids'])
        self.assertTrue(result['confirmed'])

    def test_legacy_missing_recovers_and_ambiguous_versions_require_selection(self):
        p = self.profile_project()
        material = self.profile()
        stored = self.httpd.store.require('project', p['id'])
        row = next(m for m in stored['mappings'] if m['category_id'] == 'company_profile')
        row.update(action='missing', material_ids=[], confirmed=True)
        self.httpd.store.put('project', stored)
        self.httpd.store.db.commit()
        updated, _ = self.request('GET', '/api/projects/' + p['id'])
        result = next(m for m in updated['mappings'] if m['id'] == row['id'])
        self.assertEqual([material['id']], result['material_ids'])
        self.assertFalse(result['confirmed'])
        self.profile(name='另一版本')
        row.update(action='missing', material_ids=[], confirmed=True)
        self.httpd.store.put('project', stored)
        self.httpd.store.db.commit()
        updated, _ = self.request('GET', '/api/projects/' + p['id'])
        result = next(m for m in updated['mappings'] if m['id'] == row['id'])
        self.assertEqual(('fill', 'blocked', False), (result['action'], result['status'], result['confirmed']))
        self.assertEqual([], result['material_ids'])

    def test_invalid_material_not_automatically_selected_and_append_recovers(self):
        p = self.profile_project()
        with self.assertRaises(urllib.error.HTTPError):
            self.profile(subject='其他公司')
        material = self.profile(mode='page', content='', files=[image_upload()])
        self.request('PATCH', '/api/materials/' + material['id'], {'removed_file_ids': [material['files'][0]['id']]})
        updated, _ = self.request('GET', '/api/projects/' + p['id'])
        row = next(m for m in updated['mappings'] if m['category_id'] == 'company_profile')
        self.assertEqual('blocked', row['status'])
        self.assertFalse(row['confirmed'])
        self.request('POST', '/api/materials/' + material['id'] + '/files', {'files': [image_upload()]})
        updated, _ = self.request('GET', '/api/projects/' + p['id'])
        row = next(m for m in updated['mappings'] if m['category_id'] == 'company_profile')
        self.assertEqual([material['id']], row['material_ids'])
        self.assertEqual('ready', row['status'])

    def test_form_previews_linked_project_issues_without_writing(self):
        p = self.profile_project()
        material = self.profile()
        feedback, _ = self.request('POST', '/api/materials/validate', {
            'material_id': material['id'], 'category_id': 'iso_cert', 'valid_until': '长期'})
        self.assertEqual({}, feedback['errors'])
        self.assertTrue(any('所属栏目不符' in warning for warning in feedback['project_warnings']))
        self.assertEqual('company_profile', self.httpd.store.require('material', material['id'])['category_id'])

    def test_material_feedback_and_save_reject_invalid_metadata(self):
        body = {'name': 'ISO证书', 'category_id': 'iso_cert', 'subject': server.SUBJECT,
                'mode': 'text', 'content': '证书正文', 'valid_until': '2028/12/31'}
        for invalid in ('2028/12/31', '2028-02-30', '20281231', '2000-01-01'):
            body['valid_until'] = invalid
            feedback, _ = self.request('POST', '/api/materials/validate', body)
            self.assertIn('valid_until', feedback['errors'])
            with self.assertRaises(urllib.error.HTTPError):
                self.request('POST', '/api/materials', body)
        self.assertEqual([], self.httpd.store.all('material'))
        body['valid_until'] = '长期'
        feedback, _ = self.request('POST', '/api/materials/validate', body)
        self.assertEqual({}, feedback['errors'])
        saved, _ = self.request('POST', '/api/materials', body)
        with self.assertRaises(urllib.error.HTTPError):
            self.request('PATCH', '/api/materials/' + saved['id'], {'valid_until': '无效日期'})
        self.assertEqual('长期', self.httpd.store.require('material', saved['id'])['valid_until'])
        feedback, _ = self.request('POST', '/api/materials/validate', {**body, 'category_id': 'audit_report', 'years': [], 'valid_until': '', 'content': '', 'subject': '其他公司'})
        self.assertTrue({'years', 'content', 'subject'}.issubset(feedback['errors']))
        self.assertIn('valid_until', feedback['warnings'])

    def test_batch_remove_preserves_originals_order_and_invalidates_confirmation(self):
        p = self.profile_project()
        material = self.profile(mode='page', content='', files=[image_upload(str(i)+'.png') for i in range(3)])
        ids = [f['id'] for f in material['files']]
        p, _ = self.request('GET', '/api/projects/' + p['id'])
        row = next(m for m in p['mappings'] if m['category_id'] == 'company_profile')
        self.request('PATCH', '/api/projects/' + p['id'] + '/mappings', {'mappings': [dict(row, confirmed=True)]})
        result, _ = self.request('PATCH', '/api/materials/' + material['id'], {'removed_file_ids': [ids[0], ids[2]]})
        self.assertEqual([ids[1]], [f['id'] for f in result['files']])
        p, _ = self.request('GET', '/api/projects/' + p['id'])
        self.assertFalse(next(m for m in p['mappings'] if m['id'] == row['id'])['confirmed'])
        for fid in ids:
            self.assertTrue(self.request('GET', '/api/files/' + fid, raw=True)[0])
        result, _ = self.request('PATCH', '/api/materials/' + material['id'], {'removed_file_ids': [ids[1]]})
        self.assertEqual([], result['files'])
        p, _ = self.request('GET', '/api/projects/' + p['id'])
        self.assertEqual('blocked', next(m for m in p['mappings'] if m['id'] == row['id'])['status'])

    def test_no_seed_material_upload_reorder_rotate_and_persistence(self):
        state, _ = self.request('GET', '/api/bootstrap')
        self.assertEqual(10, len([c for c in state['categories'] if c['layer']=='common']))
        self.assertEqual(7, sum(c['enabled'] for c in state['categories'] if c['layer']=='common'))
        self.assertEqual([], state['materials'])
        self.assertEqual(4, len(state['industries']))
        self.assertEqual('not_imported', state['stats']['source_status']['state'])
        category, _ = self.request('POST', '/api/categories', {'name': '补充资质', 'aliases': ['新资质']})
        material, _ = self.request('POST', '/api/materials', {'name': '资质包', 'subject': server.SUBJECT, 'category_id': category['id'], 'years': [2026], 'files': [image_upload(), image_upload('第二页.png', 'blue')]})
        self.assertEqual(2, len(material['files']))
        first, second = [f['id'] for f in material['files']]
        material, _ = self.request('PATCH', '/api/materials/' + material['id'], {'file_order': [second, first], 'rotations': {second: 90}})
        self.assertEqual(second, material['files'][0]['id'])
        self.assertEqual(90, material['files'][0]['rotation'])
        self.assertEqual(2, material['revision'])
        self.assertNotIn(self.tmp.name, json.dumps(material))
        self.assertNotIn('storage', json.dumps(material))
        raw, response_headers = self.request('GET', material['files'][0]['url'] + '?thumb=1', raw=True)
        with Image.open(io.BytesIO(raw)) as picture:
            self.assertEqual((40, 20), picture.size)
        self.assertNotIn('Access-Control-Allow-Origin', response_headers)
        self.request('POST', '/api/industries', {'name': '证券'})
        self.httpd.store.db.commit()
        reopened = server.Store(self.tmp.name)
        try:
            count = reopened.db.execute("SELECT count(*) FROM objects WHERE kind='material'").fetchone()[0]
            self.assertEqual(1, count)
        finally:
            reopened.close()
        material, _ = self.request('PATCH', '/api/materials/' + material['id'], {'removed_file_ids': [first]})
        self.assertEqual([second], [file['id'] for file in material['files']])
        self.assertIsNotNone(self.httpd.store.get('file', first))

    def test_reject_invalid_files_and_private_paths(self):
        base = {'name': '材料', 'subject': server.SUBJECT, 'category_id': 'business_license'}
        for file in [image_upload('../bad.png'), {'name': 'bad.png', 'data': base64.b64encode(b'not an image').decode()}, {'name': 'bad.svg', 'data': base64.b64encode(b'<svg/>').decode()}]:
            with self.assertRaises(urllib.error.HTTPError) as error:
                self.request('POST', '/api/materials', {**base, 'files': [file]})
            self.assertEqual(400, error.exception.code)
        for path in ['/api/files/../../server.py', '/data/store.sqlite3', '/server.py']:
            with self.assertRaises(urllib.error.HTTPError):
                self.request('GET', path)
        state, _ = self.request('GET', '/api/bootstrap')
        self.assertEqual([], state['materials'])

    def test_cross_origin_host_and_form_post_are_blocked(self):
        for headers in [{'Origin': 'https://example.com'}, {'Origin': 'null'}, {'Host': 'evil.example'}, {'Sec-Fetch-Site': 'cross-site'}]:
            with self.assertRaises(urllib.error.HTTPError) as error:
                self.request('POST', '/api/industries', {'name': '非法修改'}, headers=headers)
            self.assertEqual(403, error.exception.code)
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.request('POST', '/api/industries', {'name': '非法修改'}, headers={'Content-Type': 'text/plain'})
        self.assertEqual(400, error.exception.code)
        state, _ = self.request('GET', '/api/bootstrap')
        self.assertEqual(4, len(state['industries']))

    def test_foreign_file_rotation_and_duplicate_order_rejected(self):
        common = {'subject': server.SUBJECT, 'category_id': 'business_license', 'files': [image_upload()]}
        first, _ = self.request('POST', '/api/materials', {**common, 'name': '材料一'})
        second, _ = self.request('POST', '/api/materials', {**common, 'name': '材料二'})
        for update in [{'rotations': {second['files'][0]['id']: 90}}, {'file_order': []}, {'file_order': [first['files'][0]['id'], first['files'][0]['id']]}]:
            with self.assertRaises(urllib.error.HTTPError) as error:
                self.request('PATCH', '/api/materials/' + first['id'], update)
            self.assertEqual(400, error.exception.code)


class DocumentValidationTests(unittest.TestCase):
    def docx(self, extra=None):
        raw = io.BytesIO()
        with zipfile.ZipFile(raw, 'w', zipfile.ZIP_DEFLATED) as archive:
            archive.writestr('[Content_Types].xml', '<Types/>')
            archive.writestr('word/document.xml', '<document/>')
            for name, content in extra or []:
                archive.writestr(name, content)
        return raw.getvalue()

    def test_zip_traversal_macro_and_expansion_rejected(self):
        for additions in [[('../escape.txt', 'x')], [('word/vbaProject.bin', 'x')], [('word/embeddings/blob.bin', 'x')], [('word/bomb.xml', b'x' * (2 * 1024 * 1024))]]:
            with self.assertRaises(ValueError):
                server.check_docx('template.docx', self.docx(additions))

    def test_foreign_extension_and_damaged_word_rejected(self):
        for name, raw in [('document.zip', self.docx()), ('document.docx', b'damaged')]:
            with self.assertRaises(ValueError):
                server.check_docx(name, raw)


if __name__ == '__main__':
    unittest.main()
