"""P0: source evidence, attachment boundaries and actual image placement."""
import copy
import hashlib
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

import engine
from server import Store


class RequirementLocationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='bid-p0-test-')
        self.root = Path(self.temp.name)
        self.path = self.root / 'source.docx'
        self.store = Store(self.root / 'data', seed=False)
        self.categories = self.store.all('category')

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def scan(self, doc, materials=()):
        doc.save(self.path)
        return engine.scan_document(self.path, self.categories, list(materials), ['I01','I02','I03','I04'])

    def row(self, result, cid):
        rows = [r for r in result['mappings'] if r['category_id']==cid]
        self.assertEqual(1, len(rows))
        return rows[0]

    def test_table_request_keeps_evidence_without_using_scoring_cell(self):
        d = Document()
        d.add_paragraph('第三章 评分标准')
        t = d.add_table(rows=1, cols=2)
        t.cell(0,0).text = '提供广播电视节目制作经营许可证得2分。'
        x = self.scan(d)
        r = self.row(x, 'I03')
        self.assertFalse(r['locator_id'])
        self.assertEqual('optional', r['requirement']['status'])
        self.assertEqual(t.cell(0,0).text, r['requirement']['evidence'][0]['text'])
        self.assertGreater(len(r['requirement']['evidence'][0]['source_path']), 1)

    def test_attachment_number_ends_previous_form_and_export(self):
        d = Document()
        d.add_paragraph('附件2')
        d.add_paragraph('业绩情况汇总表')
        d.add_table(rows=1, cols=2).cell(0,0).text = '合同名称'
        d.add_paragraph('后附相应合同复印件。')
        d.add_paragraph('附件3')
        d.add_paragraph('项目负责人简历表')
        d.add_paragraph('不得插入到此人员栏目。')
        x = self.scan(d)
        r = self.row(x, 'I01')
        m = dict(id='case', category_id='I01', name='案例', subject=engine.EXPECTED_SUBJECT,
                 mode='text', content='测试插入材料', files=[], years=[], valid_until='长期')
        out = self.root/'output.docx'
        original = self.path.read_bytes()
        engine.export_document(self.path, out, [dict(r, action='fill', confirmed=True, material_ids=['case'])], [m], 'reviewed', self.categories)
        texts = [p.text for p in Document(out).paragraphs]
        self.assertLess(texts.index('测试插入材料'), texts.index('附件3'))
        self.assertGreater(texts.index('测试插入材料'), texts.index('后附相应合同复印件。'))
        self.assertEqual(original, self.path.read_bytes())

    def test_next_numbered_requirement_is_not_owned_by_previous_material(self):
        d = Document()
        d.add_paragraph('第四章 投标文件格式')
        d.add_paragraph('3.提供依法缴纳税收和社会保障资金的证明；')
        d.add_paragraph('4.提供信用中国查询截图；')
        x = self.scan(d)
        r = self.row(x, 'tax_social')
        records,_ = engine._inspect(d,self.categories)
        target = next(z for z in records if z['id']==r['locator_id'])
        self.assertEqual(d.paragraphs[1]._p, target['_anchor'])

    def test_file_composition_entry_can_be_filled_directly(self):
        d=Document();d.add_paragraph('三、报价文件的编制');d.add_paragraph('2.报价文件的组成')
        d.add_paragraph('（1）营业执照');d.add_paragraph('（2）制造商授权书（本项目无需提供）')
        r=self.row(self.scan(d),'business_license')
        self.assertTrue(r['locator_id'].startswith('at:'))
        self.assertEqual('（1）营业执照',r['title'])

    def test_appendix_export_without_built_in_heading_styles(self):
        d=Document()
        d.styles['Heading 1'].delete();d.styles['Heading 2'].delete()
        d.add_paragraph('企业应具备独立法人资格，并提供有效营业执照复印件。')
        r=self.row(self.scan(d),'business_license')
        m=dict(id='license',category_id='business_license',name='执照',subject=engine.EXPECTED_SUBJECT,
               mode='text',content='测试材料',files=[],years=[],valid_until='长期')
        out=self.root/'appendix.docx'
        engine.export_document(self.path,out,[dict(r,action='fill',confirmed=True,material_ids=['license'])],[m],'reviewed',self.categories)
        text=[p.text for p in Document(out).paragraphs]
        self.assertEqual(1,text.count('营业执照'))
        self.assertEqual(1,text.count('通用材料附件（请商务核对）'))
        self.assertIn('测试材料',text)

    def test_new_attachment_chapter_ends_contract_and_personnel_boundary(self):
        d=Document();d.add_paragraph('第四章 合同格式');d.add_paragraph('合同正文')
        d.add_paragraph('第六章 附 件');d.add_paragraph('（二）拟派人员简历')
        d.add_table(rows=1,cols=1).cell(0,0).text='姓名职称'
        d.add_paragraph('6-5 投标人近三年财务报表');d.add_paragraph('财务内容')
        x=self.scan(d);r=self.row(x,'I02');self.assertTrue(r['locator_id'])
        records,_=engine._inspect(d,self.categories)
        record=next(z for z in records if z['id']==r['locator_id'])
        self.assertEqual('姓名职称',engine._text(record['_anchor']))

    def test_role_specific_empty_tables_receive_only_their_images(self):
        d = Document()
        d.add_paragraph('第三部分 谈判响应文件格式')
        d.add_paragraph('法定代表人授权委托书')
        d.add_paragraph('本人张三系本单位法定代表人，现委托李四，为我方代理人。')
        d.add_paragraph('附：')
        d.add_paragraph('法定代表人身份证复印件。')
        d.add_table(rows=1, cols=2)
        d.add_paragraph('委托代理人身份证复印件。')
        d.add_table(rows=1, cols=2)
        d.add_paragraph('（三）投标人认为需要提供的其他材料')
        d.add_paragraph('其他内容')
        mats=[]
        for cid, person, color in [('legal_identity','张三','red'),('authorization','李四','blue')]:
            path=self.root/(cid+'.png');Image.new('RGB',(120,80),color).save(path)
            mats.append(dict(id=cid,category_id=cid,name=person+('（法定代表人）' if cid=='legal_identity' else '（委托代理人）')+'身份证',subject=engine.EXPECTED_SUBJECT,
                             mode='side_by_side',files=[{'path':str(path),'rotation':0}],years=[],valid_until='长期'))
        x=self.scan(d,mats)
        rows=[dict(self.row(x,cid),action='fill',confirmed=True,material_ids=[cid]) for cid in ['legal_identity','authorization']]
        self.assertTrue(all(r['locator_id'].startswith('at:') for r in rows))
        out=self.root/'roles.docx'
        engine.export_document(self.path,out,rows,mats,'reviewed',self.categories)
        result=Document(out)
        for i,cid in enumerate(['legal_identity','authorization']):
            blips=list(result.tables[i]._tbl.iter(qn('a:blip')))
            self.assertEqual(1,len(blips))
            part=result.part.related_parts[blips[0].get(qn('r:embed'))]
            self.assertEqual((self.root/(cid+'.png')).read_bytes(),part.blob)

    def test_identity_picture_slot_stops_before_instructions_and_signatures(self):
        d=Document();d.add_paragraph('第四章 投标文件格式')
        label=d.add_paragraph('附：法定代表人身份证电子件：')
        d.add_paragraph('');d.add_paragraph('说明：');d.add_paragraph('若法人亲自参加，应提供相关证明。')
        x=self.scan(d);r=self.row(x,'legal_identity')
        records,_=engine._inspect(d,self.categories)
        target=next(z for z in records if z['id']==r['locator_id'])
        self.assertEqual(label._p,target['_anchor'])

    def test_audit_year_from_body_and_or_alternatives(self):
        d=Document();d.add_paragraph('第四章 投标文件格式');d.add_paragraph('财务审计报告')
        d.add_paragraph('提供2024或2025年度经审计的财务报告。')
        x=self.scan(d)
        r=self.row(x,'audit_report')
        self.assertEqual([],r['required_years'])
        self.assertEqual([[2024],[2025]],r['requirement']['year_options'])
        m=dict(id='audit',category_id='audit_report',name='2024年审计报告',subject=engine.EXPECTED_SUBJECT,mode='text',content='报告',files=[],years=[2024],valid_until='长期')
        r.update(action='fill',material_ids=['audit'])
        checked=engine.validate_mappings([r],[m],x['headings'],self.categories)[0]
        self.assertEqual('ready',checked['status'])

    def audit_doc(self):
        d=Document();d.add_paragraph('第四章 投标文件格式');d.add_paragraph('财务审计报告')
        d.add_paragraph('提供2023或2024年度经审计的财务报告。')
        return d

    def audit(self, key, year):
        return dict(id=key,category_id='audit_report',name=f'{year}年审计报告',subject=engine.EXPECTED_SUBJECT,
                    mode='text',content='报告',files=[],years=[year],valid_until='长期')

    def test_alternative_audit_years_recommend_the_latest_available(self):
        cases = [
            ([('a22',2022),('a23',2023),('a24',2024)], ['a24']),
            ([('a22',2022),('a23',2023)], ['a23']),      # Newest acceptable year actually held.
            ([('a24',2024),('b24',2024),('a23',2023)], []),  # Two 2024 versions: manual choice.
            ([('a22',2022),('a21',2021)], []),            # No acceptable year at all.
        ]
        for held, expected in cases:
            with self.subTest(held=held):
                r=self.row(self.scan(self.audit_doc(),[self.audit(k,y) for k,y in held]),'audit_report')
                self.assertEqual(expected,r['material_ids'])
                if expected:
                    self.assertEqual(('fill','ready'),(r['action'],r['status']),r['issues'])
                    self.assertTrue(any('原文接受 2023 或 2024' in w for w in r['warnings']),r['warnings'])
                else:
                    self.assertNotEqual('ready',r['status'])

    def test_material_sync_recommends_latest_alternative_year(self):
        import base64,io
        raw=io.BytesIO();self.audit_doc().save(raw)
        for key,year in [('x',2023),('y',2024)]:
            body={k:v for k,v in self.audit(key,year).items() if k not in {'id','files'}}
            self.store.material(body)
        latest=next(m['id'] for m in self.store.all('material') if m['years']==[2024])
        p=self.store.create_project(dict(name='年度二选一',filename='t.docx',data=base64.b64encode(raw.getvalue()).decode()))
        self.assertEqual([latest],self.row(p,'audit_report')['material_ids'])
        # A project saved before the recommendation existed is filled by the sync, too.
        stored=self.store.require('project',p['id'])
        row=next(r for r in stored['mappings'] if r['category_id']=='audit_report')
        row.update(action='missing',material_ids=[],confirmed=False)
        synced=self.store.sync_material_matches(stored)
        self.assertEqual([latest],self.row(synced,'audit_report')['material_ids'])

    def test_definition_personal_social_and_contract_are_not_corporate_requests(self):
        d=Document()
        for text in ['如无分支机构（具有营业执照及正式员工）的，填写无分支机构。',
                     '须提供供应商为拟派项目负责人缴纳的近6个月内任意一个月的社保缴纳证明资料。',
                     '合同双方均同意，甲方支付税款后应向乙方提供税收缴纳证明。',
                     '领购招标文件前在平台提交授权委托书。']:
            d.add_paragraph(text)
        x=self.scan(d)
        for cid in ['business_license','tax_social','I04']:
            self.assertEqual('not_detected',self.row(x,cid)['requirement']['status'])

    def test_specific_case_form_preferred_to_generic_mixed_material_list(self):
        d=Document();d.add_paragraph('附件1')
        d.add_paragraph('后附公司简介、案例展示及营业执照复印件。')
        d.add_paragraph('附件2');d.add_paragraph('业绩情况汇总表')
        d.add_table(rows=1,cols=2).cell(0,0).text='合同名称'
        d.add_paragraph('后附项目合同复印件。')
        self.assertEqual('业绩情况汇总表',self.row(self.scan(d),'I01')['title'])

    def test_upgrade_refreshes_automatic_but_keeps_manual_materials_and_note(self):
        d=Document();d.add_paragraph('第四章 投标文件格式');d.add_paragraph('营业执照复印件')
        x=self.scan(d);r=self.row(x,'business_license')
        source=self.store.root/'templates/t.docx';source.write_bytes(self.path.read_bytes())
        old=dict(r,locator_id='ca:business_license:old',location_match={'status':'appendix'},confirmed=True,material_ids=['chosen'],note='保留备注')
        project=dict(id='old',template_storage='templates/t.docx',template_hash=hashlib.sha256(source.read_bytes()).hexdigest(),headings=[],mappings=[old])
        self.store.sync_common_locations(project)
        self.assertEqual(r['locator_id'],old['locator_id'])
        self.assertFalse(old['confirmed'])
        self.assertEqual(['chosen'],old['material_ids']);self.assertEqual('保留备注',old['note'])
        old.update(location_match={'status':'manual'},confirmed=True)
        project.pop('common_locator_version')
        before=copy.deepcopy(old)
        self.store.sync_common_locations(project)
        self.assertEqual(before,old)


if __name__=='__main__':
    unittest.main()
