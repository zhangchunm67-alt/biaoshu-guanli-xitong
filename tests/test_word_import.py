"""Word upload compatibility, provenance, conversion and failure cleanup."""
import base64
import hashlib
import io
from pathlib import Path
import subprocess
import sys
import tempfile
import unittest
from unittest.mock import patch
import zipfile

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import server
import word_import


def template_bytes():
    document = Document()
    document.add_heading('公司介绍', 1)
    document.add_paragraph('上传测试正文')
    table = document.add_table(rows=2, cols=2)
    for cell, text in zip((cell for row in table.rows for cell in row.cells),
                          ['材料', '数量', '保留表格', '1']):
        cell.text = text
    data = io.BytesIO()
    document.save(data)
    return data.getvalue()


class WordImportTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='bid-import-test-')
        self.store = server.Store(self.tmp.name)

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def upload(self, name, raw):
        return self.store.create_project({
            'name': 'Word 上传测试', 'filename': name,
            'data': base64.b64encode(raw).decode(),
        })

    def test_doc_suffix_with_docx_content_preserves_source_and_matches(self):
        raw = template_bytes()
        project = self.upload('招标模板.DOC', raw)
        self.assertEqual('docx_with_doc_suffix', project['template_import_kind'])
        self.assertEqual(hashlib.sha256(raw).hexdigest(), project['template_hash'])
        self.assertEqual(project['template_hash'], project['template_source_hash'])
        stored = self.store.require('project', project['id'])
        for field in ('template_storage', 'template_source_storage'):
            self.assertNotIn(field, project)
            self.assertEqual(raw, self.store.path(stored[field]).read_bytes())
        self.assertTrue(any(row['locator_id'] for row in project['mappings']))

    def test_2_5_mb_word_upload(self):
        buffer = io.BytesIO(template_bytes())
        # Unreferenced custom part makes a valid package of realistic upload size.
        with zipfile.ZipFile(buffer, 'a', compression=zipfile.ZIP_STORED) as package:
            package.writestr('custom/upload-test.bin', b'x' * 2_500_000)
        raw = buffer.getvalue()
        project = self.upload('大模板.docx', raw)
        self.assertGreater(len(raw), 2_500_000)
        self.assertEqual(hashlib.sha256(raw).hexdigest(), project['template_hash'])

    def test_invalid_macro_and_oversize_uploads_leave_no_project(self):
        raw = template_bytes()
        macro = io.BytesIO(raw)
        with zipfile.ZipFile(macro, 'a') as package:
            package.writestr('word/vbaProject.bin', b'macro')
        for name, data in [('bad.doc', b'bad'), ('renamed.doc', macro.getvalue()),
                           ('bad.docm', raw), ('empty.doc', b'')]:
            with self.subTest(name=name), self.assertRaises(ValueError):
                self.upload(name, data)
        with patch.object(server, 'MAX_DOCX', 100), self.assertRaises(ValueError):
            self.upload('large.doc', raw)
        self.assertEqual([], self.store.all('project'))
        self.assertEqual([], list((self.store.root / 'templates').iterdir()))

    def test_failed_scan_cleans_source_and_normalized_template(self):
        with patch.object(self.store, 'scan', side_effect=ValueError('扫描失败')):
            with self.assertRaisesRegex(ValueError, '扫描失败'):
                self.upload('有效模板.doc', template_bytes())
        self.assertEqual([], self.store.all('project'))
        self.assertEqual([], list((self.store.root / 'templates').iterdir()))

    @unittest.skipUnless(word_import.find_converter(), 'Local LibreOffice is not installed')
    def test_real_legacy_conversion_and_missing_converter(self):
        root = Path(self.tmp.name)
        source = root / 'fixture.docx'
        source.write_bytes(template_bytes())
        profile = (root / 'fixture-profile').as_uri()
        subprocess.run([
            word_import.find_converter(), f'-env:UserInstallation={profile}',
            '--headless', '--convert-to', 'doc:MS Word 97', '--outdir', str(root), str(source),
        ], check=True, timeout=60, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
        legacy = (root / 'fixture.doc').read_bytes()
        # Old WPS files can include a trailing ZIP. It must not override the OLE type.
        hybrid = io.BytesIO(legacy)
        with zipfile.ZipFile(hybrid, 'a') as package:
            package.writestr('drs/ink.xml', '<ink/>')
        legacy = hybrid.getvalue()
        with patch.object(word_import, 'find_converter', return_value=None):
            with self.assertRaisesRegex(ValueError, '未找到本机'):
                self.upload('旧版.doc', legacy)
        with patch.object(word_import.subprocess, 'run', side_effect=subprocess.TimeoutExpired('soffice', 60)):
            with self.assertRaisesRegex(ValueError, '转换超时'):
                self.upload('旧版.doc', legacy)
        project = self.upload('旧版.doc', legacy)
        stored = self.store.require('project', project['id'])
        self.assertEqual('converted_doc', project['template_import_kind'])
        self.assertEqual(legacy, self.store.path(stored['template_source_storage']).read_bytes())
        self.assertEqual(hashlib.sha256(legacy).hexdigest(), project['template_source_hash'])
        document = Document(self.store.path(stored['template_storage']))
        self.assertIn('上传测试正文', '\n'.join(p.text for p in document.paragraphs))
        # The converter can change legacy table layout; the UI warns about that.
        # Verify that all source cell contents survive, without promising layout identity.
        contents = ''.join(document.element.itertext())
        for text in ('材料', '数量', '保留表格', '1'):
            self.assertIn(text, contents)
        self.assertTrue(any('旧版 .doc 转换' in warning for warning in project['warnings']))


if __name__ == '__main__':
    unittest.main()
