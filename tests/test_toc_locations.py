"""Directory titles never receive material images; real body sections do."""
import copy
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from PIL import Image

import engine


class TocLocationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='bid-toc-test-')
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.docx'
        self.output = self.root / 'output.docx'
        self.categories = copy.deepcopy(engine.DEFAULT_CATEGORIES)
        self.categories.append(dict(id='I01', code='I01', layer='special', enabled=True,
                                    name='同类项目业绩汇总表', aliases=[], industry_id='insurance'))
        self.materials = []
        for cid, color in [('authorization', 'red'), ('business_license', 'green'),
                           ('company_profile', 'blue'), ('I01', 'gold')]:
            image = self.root / (cid + '.png')
            Image.new('RGB', (180, 120), color).save(image)
            self.materials.append(dict(id=cid, category_id=cid, name=cid + '图片',
                                       subject=engine.EXPECTED_SUBJECT, mode='page',
                                       files=[dict(path=str(image), name=image.name, rotation=0)],
                                       years=[], valid_until='长期', review_status='checked'))

    def tearDown(self):
        self.temp.cleanup()

    def scan(self, document):
        document.save(self.source)
        return engine.scan_document(self.source, self.categories, self.materials, ['I01'])

    @staticmethod
    def locator(document, paragraph):
        path = engine._xml_path(paragraph._p, document.element.body)
        return 'p:' + '.'.join(map(str, path)) + ':' + engine._digest(paragraph._p)

    def assert_directory_unselectable(self, document, scanned, paragraphs):
        selectable = {h['id'] for h in scanned['headings']}
        selected = {row['locator_id'] for row in scanned['mappings']}
        for paragraph in paragraphs:
            locator = self.locator(document, paragraph)
            self.assertNotIn(locator, selectable)
            self.assertNotIn(locator, selected)
            for cid in [c['id'] for c in self.categories]:
                self.assertNotIn('at:' + cid + ':' + locator, selectable)
                self.assertNotIn('at:' + cid + ':' + locator, selected)

    def export_and_assert_windows(self, document, scanned, counts, windows):
        """Check actual picture XML and image bytes in each expected body window."""
        source_bytes = self.source.read_bytes()
        source_text = [p.text for p in document.paragraphs if p.text.strip()]
        rows = []
        for cid, count in counts.items():
            matching = [r for r in scanned['mappings'] if r['category_id'] == cid]
            self.assertEqual(count, len(matching), (cid, matching))
            for row in matching:
                self.assertTrue(row['locator_id'], row)
                rows.append(dict(row, action='fill', material_ids=[cid], confirmed=True))
        checked = engine.validate_mappings(rows, self.materials, scanned['headings'], self.categories)
        self.assertTrue(all(r['status'] == 'ready' for r in checked), checked)
        engine.export_document(self.source, self.output, checked, self.materials,
                               mode='reviewed', categories=self.categories)
        result = Document(self.output)
        children = list(result.element.body)
        self.assertEqual(source_text, [p.text for p in result.paragraphs if p.text.strip()])
        self.assertEqual(sum(counts.values()), len(list(result.element.body.iter(qn('a:blip')))))
        first_body = next(i for i, node in enumerate(children) if engine._text(node) == windows[0][0])
        self.assertFalse(any(list(node.iter(qn('a:blip'))) for node in children[:first_body]))
        for start, end, cid in windows:
            left = next(i for i, node in enumerate(children) if engine._text(node) == start)
            right = next(i for i, node in enumerate(children) if i > left and engine._text(node) == end)
            pictures = [blip for node in children[left:right] for blip in node.iter(qn('a:blip'))]
            self.assertEqual(1, len(pictures), (start, end))
            image_part = result.part.related_parts[pictures[0].get(qn('r:embed'))]
            self.assertEqual((self.root / (cid + '.png')).read_bytes(), image_part.blob)
        self.assertEqual(source_bytes, self.source.read_bytes())
        with ZipFile(self.output) as archive:
            self.assertIsNone(archive.testzip())

    @staticmethod
    def add_body_sections(document):
        titles = [('（一）授权委托书', 'authorization'),
                  ('（二）营业执照', 'business_license'),
                  ('（三）公司简介', 'company_profile'),
                  ('（四）同类项目业绩汇总表', 'I01')]
        for title, cid in titles:
            document.add_paragraph(title)
            document.add_paragraph('本授权书声明：现委托李四办理本项目投标事项。' if cid == 'authorization'
                                   else '本节原有内容保持不变。')
            if cid == 'authorization':
                document.add_paragraph('附：被授权人身份证复印件')
        document.add_paragraph('（五）结束附件')
        ends = [title for title, _ in titles[1:]] + ['（五）结束附件']
        return [(title, end, cid) for (title, cid), end in zip(titles, ends)]

    def test_standard_toc_style_is_not_selectable_or_filled(self):
        document = Document()
        document.styles.add_style('TOC 1', WD_STYLE_TYPE.PARAGRAPH)
        entry = document.add_paragraph('授权委托书', style='TOC 1')
        document.add_paragraph('一、授权委托书')
        document.add_paragraph('本授权书声明：现委托李四办理本项目投标事项。')
        document.add_paragraph('附：被授权人身份证复印件')
        document.add_paragraph('二、结束附件')
        scanned = self.scan(document)
        self.assert_directory_unselectable(document, scanned, [entry])
        self.export_and_assert_windows(document, scanned, {'authorization': 1},
                                       [('一、授权委托书', '二、结束附件', 'authorization')])

    @staticmethod
    def add_field_char(paragraph, kind, instruction=''):
        run = paragraph.add_run()
        char = OxmlElement('w:fldChar')
        char.set(qn('w:fldCharType'), kind)
        run._r.append(char)
        if instruction:
            text = OxmlElement('w:instrText')
            text.text = instruction
            separate = OxmlElement('w:fldChar')
            separate.set(qn('w:fldCharType'), 'separate')
            run._r.extend([text, separate])

    def test_toc_field_spanning_plain_paragraphs_excludes_result_entries(self):
        document = Document()
        self.add_field_char(document.add_paragraph(), 'begin', ' TOC \\o "1-3" ')
        entries = [document.add_paragraph(title) for title in ['授权委托书', '营业执照']]
        self.add_field_char(document.add_paragraph(), 'end')
        windows = self.add_body_sections(document)
        scanned = self.scan(document)
        self.assert_directory_unselectable(document, scanned, entries)
        self.export_and_assert_windows(document, scanned,
                                       {cid: 1 for _, _, cid in windows}, windows)

    def test_toc_field_end_merged_into_first_body_title_keeps_that_title(self):
        # The merged heading may keep the plain style or inherit the TOC style.
        for style in (None, 'TOC 1'):
            with self.subTest(style=style):
                document = Document()
                if style:
                    document.styles.add_style(style, WD_STYLE_TYPE.PARAGRAPH)
                self.add_field_char(document.add_paragraph(style=style), 'begin', ' TOC \\o "1-3" ')
                entries = [document.add_paragraph(title, style=style) for title in ['营业执照', '公司简介']]
                # Deleting the blank line after a Word TOC leaves its field end here.
                merged = document.add_paragraph(style=style)
                self.add_field_char(merged, 'end')
                merged.add_run('（一）营业执照')
                document.add_paragraph('本节原有内容保持不变。')
                document.add_paragraph('（二）公司简介')
                document.add_paragraph('本节原有内容保持不变。')
                document.add_paragraph('（三）结束附件')
                scanned = self.scan(document)
                self.assert_directory_unselectable(document, scanned, entries)
                self.assertIn(self.locator(document, merged), {h['id'] for h in scanned['headings']})
                self.export_and_assert_windows(document, scanned, {'business_license': 1, 'company_profile': 1},
                                               [('（一）营业执照', '（二）公司简介', 'business_license'),
                                                ('（二）公司简介', '（三）结束附件', 'company_profile')])

    def test_unrepeated_chapters_keep_same_named_blank_sections(self):
        document = Document()
        for line in ['第一章 资格证明文件', '营业执照', '公司简介', '第二章 其他证明文件', '营业执照',
                     '本节原有内容保持不变。', '公司简介', '本节原有内容保持不变。', '三、结束附件']:
            document.add_paragraph(line)
        scanned = self.scan(document)
        for cid in ('business_license', 'company_profile'):
            self.assertEqual(2, len([r for r in scanned['mappings'] if r['category_id'] == cid and r['locator_id']]))
        # Chapter titles that repeat later are still an unlabelled contents list.
        document = Document()
        entries = [document.add_paragraph(t) for t in ['第一章 营业执照', '第二章 公司简介']]
        for line in ['第一章 营业执照', '本节原有内容保持不变。', '第二章 公司简介', '本节原有内容保持不变。',
                     '第三章 结束附件']:
            document.add_paragraph(line)
        scanned = self.scan(document)
        self.assert_directory_unselectable(document, scanned, entries)

    def test_manual_directory_without_page_numbers_covers_common_and_special(self):
        document = Document()
        document.add_paragraph('目录')
        entries = [document.add_paragraph(title) for title in
                   ['授权委托书', '营业执照', '公司简介', '同类项目业绩汇总表']]
        windows = self.add_body_sections(document)
        scanned = self.scan(document)
        self.assert_directory_unselectable(document, scanned, entries)
        self.export_and_assert_windows(document, scanned,
                                       {cid: 1 for _, _, cid in windows}, windows)

    def test_unlabelled_short_contents_list_before_ordered_body_is_excluded(self):
        document = Document()
        document.add_paragraph('第五部分 投标文件格式')
        document.add_paragraph('一、商务部分')
        entries = [document.add_paragraph(title) for title in
                   ['授权委托书', '营业执照', '公司简介', '同类项目业绩汇总表']]
        windows = self.add_body_sections(document)
        scanned = self.scan(document)
        self.assert_directory_unselectable(document, scanned, entries)
        self.export_and_assert_windows(document, scanned,
                                       {cid: 1 for _, _, cid in windows}, windows)

    def test_manual_contents_with_nested_introduction_keeps_both_body_titles(self):
        document = Document()
        document.add_paragraph('第五部分 投标文件格式')
        document.add_paragraph('一、商务部分')
        entries = [document.add_paragraph(title) for title in [
            '投标函', '投标报价一览表', '法定代表人身份证明', '授权委托书',
            '资格证明文件，包括：', '投标人基本情况表；', '拟投入本项目主要人员汇总表；',
            '拟投入本项目主要人员简历表；', '类似项目业绩一览表；', '技术需求偏离表',
            '服务要求偏离表；', '商务要求偏离表；', '财务状况；', '纳税证明；',
            '缴纳社会保障资金证明', '信用记录；', '承诺书；', '离职人员任职情况承诺函；',
            '投标保证金递交凭证；', '投标人关联关系单位情况说明；',
            '投标人认为应提交的其他资料。',
        ]]
        numbered_body_title = document.add_paragraph('（一）投标函')
        repeated_body_title = document.add_paragraph('投标函')
        document.add_paragraph('本函用于提交本次投标响应。')
        document.add_paragraph('（二）投标报价一览表')
        document.add_paragraph('本节报价内容由商务人员核对。')
        document.add_paragraph('（三）法定代表人身份证明')
        document.add_paragraph('兹证明本单位法定代表人的身份。')
        document.add_paragraph('（四）授权委托书格式')
        document.add_paragraph('本授权书声明：现委托李四办理本项目投标事项。')
        document.add_paragraph('附：被授权人身份证复印件')
        document.add_paragraph('（五）资格证明文件')
        document.add_paragraph('（1）投标人基本情况表')
        document.add_paragraph('本节基本情况由商务人员核对。')
        document.add_paragraph('二、服务方案部分')
        document.add_paragraph('服务方案内容。')

        scanned = self.scan(document)
        self.assert_directory_unselectable(document, scanned, entries)
        records, _ = engine._inspect(document, self.categories)
        retained = {record['id'] for record in records}
        self.assertIn(self.locator(document, numbered_body_title), retained)
        self.assertIn(self.locator(document, repeated_body_title), retained)
        self.export_and_assert_windows(document, scanned, {'authorization': 1},
                                       [('（四）授权委托书格式', '（五）资格证明文件', 'authorization')])

    def test_two_directory_authorizations_leave_both_real_body_forms(self):
        document = Document()
        document.add_paragraph('目录')
        entries = []
        for section in ['一、商务部分', '二、服务方案部分']:
            entries.append(document.add_paragraph(section))
            entries.append(document.add_paragraph('授权委托书'))
        document.add_paragraph('一、商务部分')
        document.add_paragraph('授权委托书')
        document.add_paragraph('本授权书声明：现委托李四办理本项目商务事项。')
        document.add_paragraph('附：被授权人身份证复印件')
        document.add_paragraph('二、服务方案部分')
        document.add_paragraph('授权委托书')
        document.add_paragraph('本授权书声明：现委托李四办理本项目服务事项。')
        document.add_paragraph('附：被授权人身份证复印件')
        document.add_paragraph('三、结束附件')
        scanned = self.scan(document)
        self.assert_directory_unselectable(document, scanned, entries)
        # Use body-only numbered markers because directory section names repeat.
        windows = [('一、商务部分', '二、服务方案部分', 'authorization'),
                   ('二、服务方案部分', '三、结束附件', 'authorization')]
        self.export_and_assert_repeated_sections(document, scanned, windows)

    def export_and_assert_repeated_sections(self, document, scanned, windows):
        source_bytes = self.source.read_bytes()
        rows = [dict(r, action='fill', material_ids=['authorization'], confirmed=True)
                for r in scanned['mappings'] if r['category_id'] == 'authorization']
        self.assertEqual(2, len(rows), rows)
        engine.export_document(self.source, self.output, rows, self.materials,
                               mode='reviewed', categories=self.categories)
        result = Document(self.output)
        children = list(result.element.body)
        indices = {title: [i for i, node in enumerate(children) if engine._text(node) == title]
                   for title in {title for start, end, _ in windows for title in [start, end]}}
        first_body = indices[windows[0][0]][-1]
        self.assertFalse(any(list(node.iter(qn('a:blip'))) for node in children[:first_body]))
        for start, end, cid in windows:
            left, right = indices[start][-1], indices[end][-1]
            pictures = [b for node in children[left:right] for b in node.iter(qn('a:blip'))]
            self.assertEqual(1, len(pictures))
            part = result.part.related_parts[pictures[0].get(qn('r:embed'))]
            self.assertEqual((self.root / (cid + '.png')).read_bytes(), part.blob)
        self.assertEqual(2, len(list(result.element.body.iter(qn('a:blip')))))
        self.assertEqual([p.text for p in document.paragraphs if p.text.strip()],
                         [p.text for p in result.paragraphs if p.text.strip()])
        self.assertEqual(source_bytes, self.source.read_bytes())

    def test_material_list_without_a_directory_or_later_body_remains_fillable(self):
        document = Document()
        titles = [('被授权人身份证复印件（加盖公章）；', 'authorization'),
                  ('营业执照副本复印件（加盖公章）；', 'business_license'),
                  ('公司简介；', 'company_profile')]
        for title, _ in titles:
            document.add_paragraph(title)
        document.add_paragraph('（四）结束附件')
        scanned = self.scan(document)
        ends = [title for title, _ in titles[1:]] + ['（四）结束附件']
        self.export_and_assert_windows(document, scanned, {cid: 1 for _, cid in titles},
                                       [(title, end, cid) for (title, cid), end in zip(titles, ends)])

    def test_same_name_real_body_forms_are_not_deduplicated(self):
        document = Document()
        document.add_paragraph('第五部分 投标文件格式')
        for section in ['一、商务部分', '二、服务方案部分']:
            document.add_paragraph(section)
            document.add_paragraph('授权委托书')
            document.add_paragraph('本授权书声明：现委托李四办理本项目投标事项。')
            document.add_paragraph('附：被授权人身份证复印件')
        document.add_paragraph('三、结束附件')
        scanned = self.scan(document)
        self.export_and_assert_windows(document, scanned, {'authorization': 2},
                                       [('一、商务部分', '二、服务方案部分', 'authorization'),
                                        ('二、服务方案部分', '三、结束附件', 'authorization')])

    def test_legacy_directory_locator_is_rejected_on_export(self):
        document = Document()
        document.add_paragraph('目录')
        entry = document.add_paragraph('授权委托书')
        document.add_paragraph('授权委托书')
        document.add_paragraph('本授权书声明：现委托李四办理本项目投标事项。')
        document.add_paragraph('附：被授权人身份证复印件')
        scanned = self.scan(document)
        source_bytes = self.source.read_bytes()
        row = next(r for r in scanned['mappings'] if r['category_id'] == 'authorization')
        locator = self.locator(document, entry)
        for old_locator in [locator, 'at:authorization:' + locator]:
            with self.subTest(locator=old_locator):
                legacy = dict(row, locator_id=old_locator, action='fill',
                              material_ids=['authorization'], confirmed=True)
                with self.assertRaisesRegex(ValueError, '有效填充位置'):
                    engine.export_document(self.source, self.output, [legacy], self.materials,
                                           mode='reviewed', categories=self.categories)
                self.assertFalse(self.output.exists())
                self.assertEqual(source_bytes, self.source.read_bytes())


if __name__ == '__main__':
    unittest.main()
