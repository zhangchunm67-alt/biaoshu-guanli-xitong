"""Project-level special evidence, confirmation gates and real DOCX exports."""
import base64
import io
import json
from pathlib import Path
import sys
import tempfile
import unittest

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from server import Store, SUBJECT
from test_server import image_upload


class ProjectSpecialTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='bid-project-special-')
        self.store = Store(self.tmp.name)
        for category in self.store.all('category'):
            if category['layer'] == 'common' and category['id'] != 'company_profile':
                category['enabled'] = False
                self.store.put('category', category)
        self.profile = self.store.material(dict(name='公司介绍', category_id='company_profile',
            subject=SUBJECT, mode='text', content='通用材料正文', valid_until='长期'))
        doc = Document()
        for title in ['公司介绍', '保险项目业绩与案例', '银行人员经历与专业证明', '补充说明']:
            doc.add_heading(title, 1)
        raw = io.BytesIO()
        doc.save(raw)
        self.p = self.store.create_project(dict(name='专用材料流程测试', filename='template.docx',
            data=base64.b64encode(raw.getvalue()).decode()))

    def tearDown(self):
        self.store.close()
        self.tmp.cleanup()

    def current(self):
        return self.store.public_project(self.store.require('project', self.p['id']))

    def common_confirm(self):
        return self.store.update_mappings(self.p['id'], {'mappings': [
            dict(id=r['id'], confirmed=True) for r in self.current()['mappings'] if r['layer'] == 'common']})

    def select(self, ids):
        return self.store.select_special_categories(self.p['id'], {'category_ids': ids})

    def evidence(self, category='I01', **kwargs):
        value = dict(name='专用证明', category_id=category, subject=SUBJECT,
                     mode='text', content='专用材料正文', valid_until='长期')
        value.update(kwargs)
        return self.store.material(value)

    def fill(self, material, category='I01', **kwargs):
        row = next(r for r in self.current()['mappings'] if r['category_id'] == category)
        row.update(action='fill', material_ids=[material['id']], confirmed=True, **kwargs)
        return self.store.update_mappings(self.p['id'], {'mappings': [row]})

    def test_sequence_and_explicit_empty_selection(self):
        with self.assertRaisesRegex(ValueError, '通用材料'):
            self.select(['I01'])
        self.common_confirm()
        with self.assertRaisesRegex(ValueError, '专用材料分类'):
            self.store.export(self.p['id'], {'mode': 'reviewed'})
        for body in ({'category_ids': []}, {'category_ids': ['company_profile']},
                     {'category_ids': ['I01', 'I01']}, {'category_ids': ['I01'], 'no_special_materials': True}):
            with self.assertRaises(ValueError):
                self.store.select_special_categories(self.p['id'], body)
        p = self.store.select_special_categories(self.p['id'], {'category_ids': [], 'no_special_materials': True})
        self.assertTrue(p['special_selection_confirmed'])
        self.assertEqual(1, self.store.export(self.p['id'], {'mode': 'reviewed'})['filled_count'])

    def test_selection_preserves_common_and_requires_manual_evidence(self):
        self.evidence()
        before = self.common_confirm()['mappings']
        p = self.select(['I01', 'B02'])
        self.assertEqual(before, [r for r in p['mappings'] if r['layer'] == 'common'])
        self.assertEqual(2, len([r for r in p['mappings'] if r['layer'] == 'special']))
        self.assertTrue(all(not r['confirmed'] and not r['material_ids']
                            for r in p['mappings'] if r['layer'] == 'special'))
        with self.assertRaises(ValueError):
            self.store.export(self.p['id'], {'mode': 'draft'})

    def test_person_and_cross_industry_link_export_and_persistence(self):
        material = self.evidence('F02', subject='张测试', subject_kind='person',
                                 linked_category_ids=['B02'], mode='page', files=[image_upload()])
        self.common_confirm()
        self.select(['B02'])
        p = self.fill(material, 'B02')
        self.assertTrue(all(r['confirmed'] for r in p['mappings']))
        self.assertEqual(material['id'], p['selected_special_materials'][0]['id'])
        for _ in range(2):
            export = self.store.export(self.p['id'], {'mode': 'reviewed'})
            doc = Document(self.store.path('exports/' + export['id'] + '.docx'))
            self.assertEqual(1, len(doc.inline_shapes))
            self.assertEqual(1, sum('通用材料正文' in p.text for p in doc.paragraphs))
            report = json.loads(self.store.path('exports/' + export['id'] + '.json').read_text())
            self.assertEqual(['B02'], report['special_category_ids'])
            self.assertEqual({self.profile['id'], material['id']}, {m['id'] for m in report['materials']})
        self.store.close()
        self.store = Store(self.tmp.name)
        self.assertTrue(all(r['confirmed'] for r in self.current()['mappings']))
        self.assertEqual(['B02'], self.current()['special_category_ids'])

    def test_material_edit_invalidates_only_affected_confirmation(self):
        material = self.evidence()
        self.common_confirm()
        self.select(['I01'])
        self.fill(material)
        self.store.material({'content': '更新后的专用材料'}, material['id'])
        p = self.current()
        self.assertTrue(next(r for r in p['mappings'] if r['layer'] == 'common')['confirmed'])
        self.assertFalse(next(r for r in p['mappings'] if r['layer'] == 'special')['confirmed'])
        with self.assertRaises(ValueError):
            self.store.export(self.p['id'], {'mode': 'reviewed'})
        self.fill(material)
        self.store.material({'archived': True}, material['id'])
        row = next(r for r in self.current()['mappings'] if r['layer'] == 'special')
        self.assertEqual('blocked', row['status'])
        self.assertTrue(any('归档停用' in issue for issue in row['issues']))

    def test_category_changes_and_rescan_keep_scope(self):
        material = self.evidence()
        self.common_confirm()
        self.select(['I01'])
        self.fill(material)
        p = self.select(['I01', 'B02'])
        self.assertFalse(any(r['confirmed'] for r in p['mappings'] if r['layer'] == 'special'))
        self.assertTrue(next(r for r in p['mappings'] if r['layer'] == 'common')['confirmed'])
        p = self.select(['B02'])
        self.assertNotIn('I01', [r['category_id'] for r in p['mappings']])
        p = self.store.rescan(self.p['id'])
        self.assertEqual(['B02'], p['special_category_ids'])
        self.assertFalse(any(r['confirmed'] for r in p['mappings']))

    def test_wrong_category_expired_qualification_and_missing_position_block(self):
        wrong = self.evidence('F01')
        self.common_confirm()
        self.select(['I01', 'B02'])
        with self.assertRaisesRegex(ValueError, '所属栏目'):
            self.fill(wrong)
        expired = self.evidence('B02', subject='张测试', valid_until='2020-01-01')
        with self.assertRaisesRegex(ValueError, '已过期'):
            self.fill(expired, 'B02')
        valid = self.evidence()
        with self.assertRaisesRegex(ValueError, '填充位置'):
            self.fill(valid, locator_id='')

    def test_old_project_must_choose_special_scope(self):
        self.common_confirm()
        project = self.store.require('project', self.p['id'])
        project.pop('special_category_ids', None)
        project.pop('special_selection_confirmed', None)
        self.store.put('project', project)
        p = self.current()
        self.assertFalse(p['special_selection_confirmed'])
        self.assertTrue(p['mappings'][0]['confirmed'])
        with self.assertRaisesRegex(ValueError, '专用材料分类'):
            self.store.export(self.p['id'], {'mode': 'draft'})

    def test_industry_selection_expands_all_columns_and_preserves_common(self):
        before = self.common_confirm()['mappings']
        p = self.store.select_special_categories(self.p['id'], {'industry_ids': ['insurance', 'bank']})
        expected = ['I01', 'I02', 'I03', 'I04', 'B01', 'B02', 'B03', 'B04']
        self.assertEqual(expected, p['special_category_ids'])
        self.assertEqual(['insurance', 'bank'], p['special_industry_ids'])
        self.assertEqual(before, [r for r in p['mappings'] if r['layer'] == 'common'])
        self.assertEqual(expected, [r['category_id'] for r in p['mappings'] if r['layer'] == 'special'])
        self.store.update_mappings(self.p['id'], {'mappings': [
            dict(id=r['id'], action='skip', confirmed=True) for r in p['mappings'] if r['layer'] == 'special']})
        repeated = self.store.select_special_categories(self.p['id'], {'industry_ids': ['bank', 'insurance']})
        self.assertTrue(all(r['confirmed'] for r in repeated['mappings']))
        self.store.close()
        self.store = Store(self.tmp.name)
        self.assertEqual(['insurance', 'bank'], self.current()['special_industry_ids'])

    def test_industry_request_validation_and_explicit_none(self):
        with self.assertRaisesRegex(ValueError, '通用材料'):
            self.store.select_special_categories(self.p['id'], {'industry_ids': ['insurance']})
        self.common_confirm()
        for body in ({'industry_ids': 'insurance'}, {'industry_ids': ['insurance', 'insurance']},
                     {'industry_ids': ['unknown']}, {'industry_ids': []},
                     {'industry_ids': ['insurance'], 'category_ids': ['I01']},
                     {'industry_ids': ['insurance'], 'no_special_materials': True}):
            with self.subTest(body=body), self.assertRaises(ValueError):
                self.store.select_special_categories(self.p['id'], body)
        p = self.store.select_special_categories(self.p['id'], {'industry_ids': [], 'no_special_materials': True})
        self.assertEqual([], p['special_industry_ids'])
        self.assertEqual([], p['special_category_ids'])
        self.assertTrue(p['special_selection_confirmed'])

    def test_legacy_column_selection_reads_back_as_industries_without_mutation(self):
        self.common_confirm()
        p = self.select(['I01', 'B02'])
        original = self.store.require('project', p['id'])
        self.assertEqual(['insurance', 'bank'], self.current()['special_industry_ids'])
        self.assertEqual(original, self.store.require('project', p['id']))
        self.assertEqual(['I01', 'B02'], self.current()['special_category_ids'])

    def test_common_revision_can_be_saved_then_reconfirmed(self):
        material = self.evidence()
        self.common_confirm()
        self.select(['I01'])
        p = self.fill(material)
        common = next(r for r in p['mappings'] if r['layer'] == 'common')
        p = self.store.update_mappings(self.p['id'], {'mappings': [
            dict(common, note='重新核对', confirmed=False)]})
        self.assertFalse(any(r['confirmed'] for r in p['mappings']))
        self.common_confirm()
        self.fill(material)
        self.assertEqual(2, self.store.export(self.p['id'], {'mode': 'reviewed'})['filled_count'])


if __name__ == '__main__':
    unittest.main()
