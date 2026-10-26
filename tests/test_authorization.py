import base64
import copy
import sys
import tempfile
import unittest
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import engine


class AuthorizationImagesTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.root = Path(self.tmp.name)
        self.source, self.output = self.root / 'input.docx', self.root / 'output.docx'
        self.categories = copy.deepcopy(engine.DEFAULT_CATEGORIES)
        self.materials = []
        for cid, who in [('legal_identity', '甲某'), ('authorization', '乙某')]:
            role = '法定代表人' if cid == 'legal_identity' else '委托代理人'
            files = []
            for side in ('人像面', '国徽面'):
                p = self.root / (who + side + '.png')
                Image.new('RGB', (300, 180), 'white').save(p)
                files.append(dict(name=p.name, path=str(p), rotation=0))
            self.materials.append(dict(id=cid, name=who + '（' + role + '）身份证',
                category_id=cid, subject=engine.EXPECTED_SUBJECT, mode='side_by_side',
                files=files, years=[], valid_until='2046-02-04', content=''))

    def tearDown(self):
        self.tmp.cleanup()

    def prepare(self, document):
        document.save(self.source)
        scan = engine.scan_document(self.source, self.categories, self.materials)
        rows = [r for r in scan['mappings'] if r['category_id'] == 'authorization']
        self.assertEqual(1, len(rows))
        row = dict(rows[0], action='fill', material_ids=['legal_identity', 'authorization'], confirmed=True)
        return row, scan

    def checked(self, row, scan, materials=None):
        return engine.validate_mappings([row], materials or self.materials, scan['headings'], self.categories)[0]

    def form(self):
        doc = Document()
        doc.add_heading('第六章 投标文件格式', 1)
        doc.add_paragraph('法定代表人授权委托书')
        doc.add_paragraph('其他目录项')
        doc.add_heading('附件5：法定代表人授权书格式', 2)
        doc.add_heading('法定代表人授权委托书', 3)
        p = doc.add_paragraph('本授权书声明：注册于（国家或地区的名称）的（公司名称）的在下面签字的')
        p.add_run('（法定代表人姓名、').bold = True
        p.add_run('职务）代表本公司授权（单位名称）的（被授权人的姓名、职务）为合法代理人，就（项目名称/项目编号）处理竞价事宜。')
        doc.add_paragraph('本授权书于   年  月  日签字生效，特此声明。')
        doc.add_paragraph('单位名称：（盖公章）')
        doc.add_paragraph('日期：')
        doc.add_paragraph('附：被授权人身份证复印件')
        doc.add_paragraph('附件6：资格证明文件')
        doc.add_paragraph('下面原文不得改变')
        return doc

    def test_matching_entry_receives_only_selected_images_without_generating_body(self):
        doc = Document()
        doc.add_paragraph('被授权人身份证复印件（加盖公章）；')
        row, scan = self.prepare(doc)
        source_bytes = self.source.read_bytes()
        row = self.checked(row, scan)
        self.assertEqual('ready', row['status'], row['issues'])
        self.assertTrue(row['locator_id'].startswith('at:'))
        report = engine.export_document(self.source, self.output, [row], self.materials, categories=self.categories)
        result = Document(self.output)
        text = '\n'.join(p.text for p in result.paragraphs)
        self.assertNotIn('本人甲某', text)
        self.assertNotIn('现委托乙某', text)
        self.assertEqual(4, len(result.inline_shapes))
        self.assertEqual(4, report['results'][0]['file_count'])
        self.assertNotIn('authorization', report['results'][0])
        self.assertEqual(source_bytes, self.source.read_bytes())
        with ZipFile(self.output) as z:
            self.assertIsNone(z.testzip())

    def test_blank_form_keeps_all_text_and_formatting_with_images_before_next_attachment(self):
        doc = self.form()
        row, scan = self.prepare(doc)
        row['material_ids'] = ['authorization']
        row = self.checked(row, scan)
        self.assertEqual('ready', row['status'], row['issues'])
        self.assertTrue(row['locator_id'].startswith('at:authorization:'))
        engine.export_document(self.source, self.output, [row], self.materials, categories=self.categories, mode='reviewed')
        result = Document(self.output)
        original = [p.text for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(original, [p.text for p in result.paragraphs if p.text.strip()])
        original_runs = [[(r.text, r.bold) for r in p.runs] for p in doc.paragraphs if p.text.strip()]
        self.assertEqual(original_runs, [[(r.text, r.bold) for r in p.runs] for p in result.paragraphs if p.text.strip()])
        self.assertEqual(2, len(result.inline_shapes))
        children = list(result.element.body)
        image_index = next(i for i, e in enumerate(children) if list(e.iter(qn('w:drawing'))))
        following_index = next(i for i, e in enumerate(children) if '附件6' in engine._text(e))
        self.assertLess(image_index, following_index)

    def test_images_do_not_require_both_roles_or_personal_fields(self):
        doc = Document()
        doc.add_paragraph('被授权人身份证复印件；')
        row, scan = self.prepare(doc)
        for ids in [['authorization'], ['legal_identity'], ['legal_identity', 'authorization']]:
            checked = self.checked(dict(row, material_ids=ids), scan)
            self.assertEqual('ready', checked['status'], checked['issues'])
        self.assertEqual('blocked', self.checked(dict(row, material_ids=[]), scan)['status'])

    def test_retired_settings_are_discarded_and_require_one_new_confirmation(self):
        row, scan = self.prepare(self.form())
        row['authorization'] = dict(enabled=True, signed_on='', project_name='不得写入')
        row['authorization_preview'] = {'body_lines': ['不得写入']}
        with self.assertRaisesRegex(ValueError, '重新核对'):
            engine.export_document(self.source, self.output, [row], self.materials, 'reviewed', categories=self.categories)
        engine.export_document(self.source, self.output, [row], self.materials, 'draft', categories=self.categories)
        self.assertNotIn('不得写入', '\n'.join(p.text for p in Document(self.output).paragraphs))
        checked = self.checked(row, scan)
        self.assertEqual('ready', checked['status'], checked['issues'])
        self.assertFalse(checked['confirmed'])
        self.assertNotIn('authorization', checked)
        self.assertNotIn('authorization_preview', checked)
        checked['confirmed'] = True
        engine.export_document(self.source, self.output, [checked], self.materials, categories=self.categories)
        self.assertNotIn('不得写入', '\n'.join(p.text for p in Document(self.output).paragraphs))

    def test_letter_body_takes_only_the_text_template(self):
        categories = self.categories + [dict(id='text_template', name='文书模板', aliases=[],
                                             layer='special', enabled=True)]
        materials = self.materials + [dict(id='letter', name='法定代表人授权委托书模板',
            category_id='text_template', subject=engine.EXPECTED_SUBJECT, mode='text',
            content='授权委托书文字模板正文', files=[], years=[], valid_until='长期')]
        for agent_slot in (True, False):
            doc = Document()
            doc.add_heading('法定代表人授权委托书', 1)
            doc.add_paragraph('本授权书声明：本人甲某系北京腾宝科技有限公司的法定代表人，现委托乙某，为我方代理人。')
            doc.add_paragraph('日期：')
            if agent_slot:
                doc.add_paragraph('附：被授权人身份证复印件')
            doc.add_heading('附件6：资格证明文件', 1)
            doc.save(self.source)
            scan = engine.scan_document(self.source, categories, materials, ['text_template'])
            image = next(r for r in scan['mappings'] if r['category_id'] == 'authorization')
            letter = next(r for r in scan['mappings'] if r['category_id'] == 'text_template')
            self.assertEqual('法定代表人授权委托书', letter['title'])
            self.assertFalse(any(h['id'].startswith('ca:authorization') for h in scan['headings']))
            rows = [dict(letter, action='fill', material_ids=['letter'])]
            if agent_slot:
                self.assertTrue(image['locator_id'].startswith('at:authorization:'))
                rows.append(dict(image, action='fill', material_ids=['authorization']))
            else:
                self.assertEqual(('', 'missing'), (image['locator_id'], image['status']))
                self.assertIn('文字模板', image['location_match']['reason'])
                rows.append(image)
            checked = engine.validate_mappings(rows, materials, scan['headings'], categories)
            self.assertEqual(['ready', 'ready' if agent_slot else 'missing'],
                             [r['status'] for r in checked], checked)
            engine.export_document(self.source, self.output, checked, materials, categories=categories)
            text = [p.text for p in Document(self.output).paragraphs]
            self.assertEqual(1, text.count('授权委托书文字模板正文'))
            self.assertEqual(2 if agent_slot else 0, len(Document(self.output).inline_shapes))
            self.assertEqual(not agent_slot, any('待补：被授权人身份证附件' in t for t in text))

    def test_agent_id_slots_need_a_copy_and_accept_passports(self):
        categories = self.categories + [dict(id='text_template', name='文书模板', aliases=[],
                                             layer='special', enabled=True)]
        cases = {'附：被授权人护照复印件': True, '附：委托代理人护照复印件': True,
                 '授权代理人身份证/护照复印件（正反面）': True, '附：被授权人身份证复印件（原件备查）': True,
                 '附：被授权人身份证号码：________': False, '附：被授权人身份证的真实性声明': False,
                 '附：被授权人身份证原件（开标现场查验）': False}
        for region in ([], ['第六章 投标文件格式']):
            for label, expected in cases.items():
                with self.subTest(region=bool(region), label=label):
                    doc = Document()
                    for line in region + ['附件3：法定代表人授权书', '本授权书声明：现委托乙某办理投标事项。',
                                          '日期：', label, '附件4：资格证明文件']:
                        doc.add_paragraph(line)
                    doc.save(self.source)
                    scan = engine.scan_document(self.source, categories, self.materials, ['text_template'])
                    image = next(r for r in scan['mappings'] if r['category_id'] == 'authorization')
                    letter = next(r for r in scan['mappings'] if r['category_id'] == 'text_template')
                    self.assertEqual('附件3：法定代表人授权书', letter['title'])
                    self.assertEqual(expected, image['locator_id'].startswith('at:authorization:'), image)
                    if expected:
                        checked = self.checked(dict(image, action='fill', material_ids=['authorization']), scan)
                        self.assertEqual('ready', checked['status'], checked['issues'])
                    else:
                        self.assertEqual(('', 'missing'), (image['locator_id'], image['status']))

    def test_expired_missing_or_text_materials_still_block(self):
        row, scan = self.prepare(self.form())
        for change in [{'valid_until': '2000-01-01'}, {'files': []}, {'mode': 'text', 'content': '正文'}]:
            materials = copy.deepcopy(self.materials)
            materials[1].update(change)
            self.assertEqual('blocked', self.checked(row, scan, materials)['status'])

    def test_existing_images_and_conflicting_agent_still_block(self):
        doc = self.form()
        doc.paragraphs[10].insert_paragraph_before().add_run().add_picture(self.materials[0]['files'][0]['path'])
        row, scan = self.prepare(doc)
        self.assertTrue(any('已含证件图片' in e for e in self.checked(row, scan)['issues']))
        doc = Document()
        doc.add_heading('法定代表人授权委托书', 1)
        doc.add_paragraph('本人甲某系北京腾宝科技有限公司的法定代表人，现委托丙某，为我方代理人。')
        doc.add_paragraph('附：被授权人身份证复印件')
        row, scan = self.prepare(doc)
        self.assertTrue(any('丙某' in e and '乙某' in e for e in self.checked(row, scan)['issues']))

    def test_store_migrates_old_project_and_ignores_cached_client_fields(self):
        from server import Store
        store = Store(self.root / 'store', seed=False)
        try:
            ids = []
            for material in self.materials:
                body = {k: v for k, v in material.items() if k not in {'id', 'files'}}
                body['files'] = [dict(name=f['name'], data=base64.b64encode(Path(f['path']).read_bytes()).decode()) for f in material['files']]
                ids.append(store.material(body)['id'])
            self.form().save(self.source)
            project = store.create_project(dict(name='图片填充', filename='template.docx', data=base64.b64encode(self.source.read_bytes()).decode()))
            raw = store.require('project', project['id'])
            row = next(r for r in raw['mappings'] if r['category_id'] == 'authorization')
            row.update(action='fill', material_ids=ids, confirmed=True, note='保留备注',
                       authorization={'enabled': True}, authorization_preview={'body_lines': ['旧正文']})
            locator, mapping_id = row['locator_id'], row['id']
            store.put('project', raw)
            store.db.commit()
            reopened = store.public_project(store.require('project', project['id']))
            migrated = next(r for r in reopened['mappings'] if r['id'] == mapping_id)
            self.assertEqual('ready', migrated['status'], migrated['issues'])
            self.assertFalse(migrated['confirmed'])
            self.assertEqual(ids, migrated['material_ids'])
            self.assertEqual(locator, migrated['locator_id'])
            self.assertEqual('保留备注', migrated['note'])
            self.assertNotIn('authorization', migrated)
            saved = store.update_mappings(project['id'], {'mappings': [dict(migrated, confirmed=True, authorization={'enabled': True, 'signed_on': ''})]})
            result = next(r for r in saved['mappings'] if r['id'] == mapping_id)
            self.assertTrue(result['confirmed'])
            self.assertNotIn('authorization', result)
            persisted = next(r for r in store.require('project', project['id'])['mappings'] if r['id'] == mapping_id)
            self.assertNotIn('authorization', persisted)
        finally:
            store.close()


if __name__ == '__main__':
    unittest.main()
