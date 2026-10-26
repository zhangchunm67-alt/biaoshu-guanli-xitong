"""Original material entries, semantic variants, image placement and upgrades."""
import copy
import hashlib
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

import engine
from server import Store


class EntryMatchingTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='bid-entry-test-')
        self.root = Path(self.temp.name)
        self.path = self.root / 'source.docx'
        self.store = Store(self.root / 'data', seed=False)
        self.categories = self.store.all('category')

    def tearDown(self):
        self.store.close()
        self.temp.cleanup()

    def scan(self, doc, materials=(), special=('I01', 'I02', 'I03', 'I04')):
        doc.save(self.path)
        return engine.scan_document(self.path, self.categories, list(materials), special)

    def row(self, scanned, cid):
        rows = [r for r in scanned['mappings'] if r['category_id'] == cid]
        self.assertEqual(1, len(rows))
        return rows[0]

    def picture(self, cid):
        path = self.root / (cid + '.png')
        Image.new('RGB', (160, 100), 'navy').save(path)
        return dict(id=cid, category_id=cid, name='法人身份证' if cid == 'legal_identity' else cid,
                    subject=engine.EXPECTED_SUBJECT, mode='page', years=[], valid_until='长期',
                    files=[dict(path=str(path), rotation=0)])

    def test_repeated_submission_entries_insert_one_picture_after_first_match(self):
        doc = Document()
        for text in ['标书内容至少包括：', '公司简介；', '营业执照副本复印件（加盖公章）；',
                     '组织机构代码副本复印件（加盖公章）；', '投标函',
                     '•营业执照副本复印件（加盖公章）；']:
            doc.add_paragraph(text)
        material = self.picture('business_license')
        scanned = self.scan(doc, [material])
        row = self.row(scanned, 'business_license')
        self.assertEqual('ready', row['status'])
        self.assertEqual(['business_license'], row['material_ids'])
        self.assertEqual('entry', row['location_match']['method'])
        self.assertEqual(1, row['location_match']['alternative_count'])
        before = hashlib.sha256(self.path.read_bytes()).hexdigest()
        out = self.root / 'output.docx'
        engine.export_document(self.path, out, [dict(row, confirmed=True)], [material], 'reviewed', self.categories)
        result = Document(out)
        body = list(result.element.body)
        left = next(i for i, e in enumerate(body) if engine._text(e) == doc.paragraphs[2].text)
        right = next(i for i, e in enumerate(body) if engine._text(e) == doc.paragraphs[3].text)
        self.assertEqual(1, len(list(result.element.body.iter(qn('a:blip')))))
        self.assertEqual(1, sum(len(list(e.iter(qn('a:blip')))) for e in body[left+1:right]))
        self.assertFalse(any('通用材料附件' in p.text for p in result.paragraphs))
        self.assertEqual(before, hashlib.sha256(self.path.read_bytes()).hexdigest())
        with zipfile.ZipFile(out) as archive:
            self.assertIsNone(archive.testzip())

    def test_nonidentical_names_in_submission_list_cover_material_families(self):
        labels = {
            'business_license': '工商营业执照扫描件（盖章）；',
            'legal_identity': '单位负责人身份证正反面复印件；',
            'authorization': '被授权人身份证复印件（加盖公章）；',
            'audit_report': '2024年度财务会计报告（扫描件）；',
            'tax_social': '纳税及社保缴费证明；',
            'company_profile': '单位基本概况；',
            'iso_cert': '质量管理体系认证证书（复印件）；',
            'I01': '既往执行案例（罗列）；',
            'I02': '核心团队履历；',
            'I03': '广播电视节目制作经营许可证复印件；',
            'I04': '媒体合作协议（扫描件）；',
        }
        for cid, title in labels.items():
            with self.subTest(category=cid):
                doc = Document()
                doc.add_paragraph('投标文件的组成')
                doc.add_paragraph(title)
                row = self.row(self.scan(doc), cid)
                self.assertTrue(row['locator_id'].startswith('at:'), row)
                self.assertEqual(title, row['title'])
                self.assertEqual('matched', row['location_match']['status'])
                self.assertEqual('entry', row['location_match']['method'])
                if cid == 'audit_report':
                    self.assertEqual([2024], row['required_years'])

    def test_actual_response_slot_preferred_over_earlier_submission_list(self):
        doc = Document()
        doc.add_paragraph('投标文件的组成')
        doc.add_paragraph('营业执照副本复印件；')
        doc.add_paragraph('第四章 投标文件格式')
        doc.add_paragraph('营业执照')
        row = self.row(self.scan(doc), 'business_license')
        self.assertEqual('营业执照', row['title'])
        self.assertTrue(row['locator_id'].startswith('p:3:'))

    def test_table_entry_fills_its_own_cell_without_changing_other_cells(self):
        doc = Document()
        table = doc.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '材料名称'
        table.cell(0, 1).text = '提交要求'
        table.cell(1, 0).text = '工商营业执照副本复印件'
        table.cell(1, 1).text = '加盖公章'
        material = self.picture('business_license')
        row = self.row(self.scan(doc, [material]), 'business_license')
        self.assertEqual('ready', row['status'])
        out = self.root / 'table.docx'
        engine.export_document(self.path, out, [dict(row, confirmed=True)], [material], 'reviewed', self.categories)
        table = Document(out).tables[0]
        self.assertEqual(1, len(list(table.cell(1, 0)._tc.iter(qn('a:blip')))))
        self.assertEqual('加盖公章', table.cell(1, 1).text)
        self.assertEqual(0, len(list(table.cell(1, 1)._tc.iter(qn('a:blip')))))

    def test_identity_entry_accepts_picture_without_generating_body(self):
        material = self.picture('legal_identity')
        for label in ['法人身份证扫描件；', '法人身份证扫描件']:
            with self.subTest(label=label):
                doc = Document()
                doc.add_paragraph(label)
                row = self.row(self.scan(doc, [material]), 'legal_identity')
                self.assertEqual('ready', row['status'])
                out = self.root / 'identity.docx'
                engine.export_document(self.path, out, [dict(row, confirmed=True)], [material], 'reviewed', self.categories)
                self.assertEqual(1, len(list(Document(out).element.body.iter(qn('a:blip')))))
                self.assertNotIn('兹证明', ''.join(p.text for p in Document(out).paragraphs))

    def test_similar_but_different_materials_are_not_bidder_license_or_identity(self):
        doc = Document()
        for title in ['无需提供营业执照；', '会计师事务所营业执照复印件；',
                      '组织机构代码副本复印件；', '税务登记证副本复印件；',
                      '纳税信用等级证明；', '项目经理身份证复印件；']:
            doc.add_paragraph(title)
        scanned = self.scan(doc)
        for cid in ['business_license', 'legal_identity', 'tax_social']:
            self.assertFalse(self.row(scanned, cid)['locator_id'])

    def test_cross_industry_generic_entry_is_not_silently_duplicated(self):
        doc = Document()
        doc.add_paragraph('投标文件的组成')
        doc.add_paragraph('既往执行案例（罗列）；')
        scanned = self.scan(doc, special=['I01', 'B01'])
        for cid in ['I01', 'B01']:
            row = self.row(scanned, cid)
            self.assertFalse(row['locator_id'])
            self.assertEqual('ambiguous', row['location_match']['status'])

    def test_upgrade_retains_selected_materials_notes_and_manual_decisions(self):
        doc = Document()
        doc.add_paragraph('营业执照副本复印件（加盖公章）；')
        fresh = self.row(self.scan(doc), 'business_license')
        old = dict(copy.deepcopy(fresh), id='old', locator_id='ca:business_license:old',
                   title='文末新增材料区 · 营业执照', material_ids=['chosen'], note='用户备注',
                   auto_match=False, action='fill', confirmed=True,
                   location_match={'status': 'appendix'})
        for decision in ['automatic', 'manual', 'skip', 'missing']:
            with self.subTest(decision=decision):
                row = copy.deepcopy(old)
                if decision == 'manual':
                    row['location_match']['status'] = 'manual'
                elif decision in {'skip', 'missing'}:
                    row['action'] = decision
                project = dict(mappings=[row])
                self.store.refresh_location_rows(project, {'mappings': [fresh], 'requirements': {}}, 'common')
                self.assertEqual(['chosen'], row['material_ids'])
                self.assertEqual('用户备注', row['note'])
                if decision == 'automatic':
                    self.assertEqual(fresh['locator_id'], row['locator_id'])
                    self.assertFalse(row['confirmed'])
                else:
                    self.assertEqual(old['locator_id'], row['locator_id'])
                    self.assertTrue(row['confirmed'])


if __name__ == '__main__':
    unittest.main()
