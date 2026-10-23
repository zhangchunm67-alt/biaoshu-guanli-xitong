"""Common material requests, append-only attachment areas and legacy upgrades."""
import copy
import hashlib
import sys
import tempfile
import unittest
from pathlib import Path
from unittest.mock import patch

from docx import Document

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import engine
from location_preview import LocationPreview
from server import Store


class CommonLocationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='bid-common-test-')
        self.root = Path(self.temp.name)
        self.source = self.root / 'source.docx'
        self.output = self.root / 'output.docx'
        self.categories = copy.deepcopy(engine.DEFAULT_CATEGORIES)

    def tearDown(self):
        self.temp.cleanup()

    def material(self, category, **extra):
        return dict(id=category, category_id=category, name='法人身份证' if category == 'legal_identity' else category,
                    subject=engine.EXPECTED_SUBJECT, mode='text', content='已确认的' + category + '材料',
                    files=[], years=[], valid_until='长期', **extra)

    def scan(self, doc, materials=()):
        doc.save(self.source)
        return engine.scan_document(self.source, self.categories, list(materials))

    def test_punctuated_and_annotated_material_list_is_filled_at_each_entry(self):
        doc = Document()
        doc.add_paragraph('公司简介；')
        doc.add_paragraph('营业执照副本复印件（加盖公章）；')
        materials = [self.material(c) for c in ('business_license', 'company_profile')]
        original_text = [p.text for p in doc.paragraphs]
        scanned = self.scan(doc, materials)
        rows = [dict(r, confirmed=True) for r in scanned['mappings'] if r['status'] == 'ready']
        self.assertEqual(2, len(rows))
        self.assertTrue(all(r['location_match']['method'] == 'entry' for r in rows))
        before = hashlib.sha256(self.source.read_bytes()).hexdigest()
        engine.export_document(self.source, self.output, rows, materials, categories=self.categories)
        result = [p.text for p in Document(self.output).paragraphs]
        self.assertNotIn('通用材料附件（请商务核对）', result)
        for title,cid in zip(original_text,['company_profile','business_license']):
            self.assertEqual('已确认的'+cid+'材料',result[result.index(title)+1])
        engine.export_document(self.source, self.output, rows, materials, categories=self.categories)
        self.assertEqual(1, [p.text for p in Document(self.output).paragraphs].count('已确认的business_license材料'))
        self.assertEqual(before, hashlib.sha256(self.source.read_bytes()).hexdigest())

    def test_id_attachment_request_does_not_claim_proof_body_completed(self):
        doc = Document()
        doc.add_paragraph('投标人须提供营业执照复印件、法人身份证复印件并加盖公章。')
        material = self.material('legal_identity')
        scanned = self.scan(doc, [material])
        row = next(r for r in scanned['mappings'] if r['category_id'] == 'legal_identity')
        self.assertEqual('ready', row['status'])
        self.assertTrue(any('仅装配身份证附件' in w for w in row['warnings']))
        self.assertIn('证明正文另行核对', row['title'])
        material['valid_until'] = '2001-01-01'
        checked = engine.validate_mappings([row], [material], scanned['headings'], self.categories)[0]
        self.assertEqual('blocked', checked['status'])

    def test_scoring_title_is_not_an_attachment_target(self):
        doc = Document()
        doc.add_heading('第三章 评标办法', 1)
        doc.add_heading('营业执照', 2)
        doc.add_paragraph('投标人应提供营业执照复印件。')
        scanned = self.scan(doc, [self.material('business_license')])
        row = next(r for r in scanned['mappings'] if r['category_id'] == 'business_license')
        self.assertTrue(row['locator_id'].startswith('ca:'))
        self.assertEqual('appendix', row['location_match']['status'])

    def test_real_response_form_preferred_over_scoring_section(self):
        doc = Document()
        doc.add_heading('第三章 评标办法', 1)
        doc.add_heading('营业执照', 2)
        doc.add_heading('第六章 投标文件格式', 1)
        doc.add_heading('营业执照副本（加盖公章）', 2)
        scanned = self.scan(doc, [self.material('business_license')])
        rows = [r for r in scanned['mappings'] if r['category_id'] == 'business_license']
        self.assertEqual(1, len(rows))
        self.assertEqual('营业执照副本（加盖公章）', rows[0]['title'])
        self.assertTrue(rows[0]['locator_id'].startswith('p:'))

    def test_negative_request_and_auditor_license_do_not_select_bidder_license(self):
        doc = Document()
        doc.add_paragraph('本项目无需提供营业执照。')
        doc.add_paragraph('审计报告应包含会计师事务所营业执照。')
        scanned = self.scan(doc, [self.material('business_license')])
        row = next(r for r in scanned['mappings'] if r['category_id'] == 'business_license')
        self.assertFalse(row['locator_id'])
        self.assertEqual('missing', row['status'])

    def test_request_in_table_is_not_filled_into_scoring_cell(self):
        doc = Document()
        table=doc.add_table(rows=2, cols=2)
        table.cell(0,0).text='评分项目'
        table.cell(1,0).text='须提供营业执照复印件'
        scanned = self.scan(doc, [self.material('business_license')])
        row = next(r for r in scanned['mappings'] if r['category_id'] == 'business_license')
        self.assertTrue(row['locator_id'].startswith('ca:'))

    def test_project_managers_id_does_not_select_legal_representatives_id(self):
        doc = Document()
        doc.add_paragraph('请提供项目负责人身份证复印件。')
        scanned = self.scan(doc, [self.material('legal_identity')])
        row = next(r for r in scanned['mappings'] if r['category_id'] == 'legal_identity')
        self.assertFalse(row['locator_id'])
        self.assertFalse(row['material_ids'])

    def test_appendix_preview_describes_new_area_without_fabricating_original_marker(self):
        doc = Document()
        doc.add_paragraph('请说明企业发展历程，并提供企业简介。')
        scanned = self.scan(doc)
        row = next(r for r in scanned['mappings'] if r['category_id'] == 'company_profile')
        preview = LocationPreview()
        try:
            with patch.object(preview, '_build', return_value={'preview_id': 'test', 'page_count': 3}):
                result = preview.describe({'template': self.source, 'template_hash': hashlib.sha256(self.source.read_bytes()).hexdigest(),
                    'template_name': 'source.docx', 'locator_id': row['locator_id'], 'action': 'fill',
                    'materials': [], 'categories': self.categories})
            self.assertEqual(3, result['initial_page'])
            self.assertEqual([], result['markers'])
            self.assertIn('新增通用材料区', result['description'])
        finally:
            preview.close()

    def test_legacy_automatic_missing_upgrades_but_manual_skip_is_preserved(self):
        store = Store(self.root / 'data', seed=False)
        try:
            doc = Document()
            doc.add_paragraph('公司简介；')
            path = store.root / 'templates/template.docx'
            doc.save(path)
            scanned = engine.scan_document(path, store.all('category'), [])
            row = next(r for r in scanned['mappings'] if r['category_id'] == 'company_profile')
            row.update(locator_id='', confirmed=True, action='missing')
            project = dict(id='legacy', template_storage='templates/template.docx', template_hash=hashlib.sha256(path.read_bytes()).hexdigest(),
                           mappings=[row], headings=[], special_category_ids=[])
            upgraded = store.sync_common_locations(project)
            self.assertTrue(upgraded['mappings'][0]['locator_id'].startswith('at:'))
            self.assertFalse(upgraded['mappings'][0]['confirmed'])
            project.pop('common_locator_version')
            project['mappings'][0].update(locator_id='', action='skip', confirmed=True)
            kept = store.sync_common_locations(project)['mappings'][0]
            self.assertEqual(('skip', True, ''), (kept['action'], kept['confirmed'], kept['locator_id']))
        finally:
            store.close()


if __name__ == '__main__':
    unittest.main()
