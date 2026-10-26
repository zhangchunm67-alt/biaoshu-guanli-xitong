"""Actual page coordinates, layout preservation, unsaved choices and read-only API."""
import base64
import copy
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import threading
import unittest
from unittest.mock import patch
import urllib.error
import urllib.request

from docx import Document
from docx.oxml.ns import qn

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import engine
import server
from location_preview import LocationPreview, _render_pdf
from word_import import find_converter


def make_fixture(path):
    doc = Document()
    doc.sections[0].header.paragraphs[0].text = '原页预览验收'
    doc.add_heading('公司介绍', 1)
    doc.add_paragraph('第一处原文，材料必须插在这段之后。')
    doc.add_page_break()
    doc.add_heading('公司介绍', 1)
    doc.add_paragraph('第二处原文，不能与第一处混淆。')
    doc.add_heading('业绩附件', 1)
    table = doc.add_table(rows=2, cols=2)
    table.style = 'Table Grid'
    table.cell(0, 0).text = '项目'
    table.cell(0, 1).text = '证明'
    table.cell(1, 0).text = '示例项目'
    table.cell(1, 1).text = '{{同类项目业绩}}'
    doc.add_page_break()
    doc.add_heading('身份证附件', 1)
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    doc.add_paragraph()
    doc.save(path)


@unittest.skipUnless(find_converter(), 'Local LibreOffice is required for page rendering')
class RenderedLocationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.TemporaryDirectory(prefix='bid-location-test-')
        cls.root = Path(cls.tmp.name)
        cls.source = cls.root / 'fixture.docx'
        make_fixture(cls.source)
        cls.raw = cls.source.read_bytes()
        cls.categories = copy.deepcopy(engine.DEFAULT_CATEGORIES) + [
            {'id': 'I01', 'name': '同类项目业绩', 'aliases': [], 'layer': 'special', 'enabled': True}]
        cls.scan = engine.scan_document(cls.source, cls.categories, [], ['I01'])
        cls.preview = LocationPreview()

    @classmethod
    def tearDownClass(cls):
        cls.preview.close()
        cls.tmp.cleanup()

    def snapshot(self, locator, **overrides):
        return dict({'template': self.source, 'template_hash': hashlib.sha256(self.raw).hexdigest(),
                     'template_name': self.source.name, 'categories': self.categories,
                     'locator_id': locator, 'materials': [], 'action': 'fill'}, **overrides)

    def test_repeated_titles_land_on_different_original_pages(self):
        rows = [r for r in self.scan['mappings'] if r['category_id'] == 'company_profile']
        first, second = [self.preview.describe(self.snapshot(r['locator_id'])) for r in rows]
        self.assertEqual([1], first['target_pages'])
        self.assertEqual([2], second['target_pages'])
        self.assertEqual(first['preview_id'], second['preview_id'])
        self.assertTrue(first['markers'][0]['y'] > .1)
        self.assertEqual(self.raw, self.source.read_bytes())

    def test_unselected_position_shows_original_pages_without_guessing_a_marker(self):
        result = self.preview.describe(self.snapshot(''))
        self.assertEqual(1, result['initial_page'])
        self.assertEqual(3, result['page_count'])
        self.assertEqual([], result['target_pages'])
        self.assertEqual([], result['markers'])
        self.assertTrue(self.preview.page(result['preview_id'], 1).startswith(b'\x89PNG'))
        self.assertEqual(self.raw, self.source.read_bytes())

    def test_wrapped_paragraph_marker_is_below_its_last_line(self):
        import pypdfium2 as pdfium
        document = Document()
        document.add_heading('同类项目业绩', 1)
        document.add_paragraph('请按项目顺序提供合同及履约证明材料，并核对签章与实施日期。' * 7 + '说明末尾。')
        source = self.root / 'wrapped.docx'
        document.save(source)
        scan = engine.scan_document(source, self.categories, [], ['I01'])
        locator = next(r['locator_id'] for r in scan['mappings'] if r['category_id'] == 'I01')
        result = self.preview.describe(self.snapshot(locator, template=source,
            template_hash=hashlib.sha256(source.read_bytes()).hexdigest()))
        pdf = pdfium.PdfDocument(self.preview.cache[result['preview_id']]['pdf'])
        page = pdf[0]
        text = page.get_textpage()
        try:
            bottom = min(text.get_charbox(i)[1] for i in range(text.count_chars())
                         if text.get_text_range(i, 1).strip())
            self.assertEqual([1], result['target_pages'])
            self.assertGreater(result['markers'][0]['y'] * page.get_height(), page.get_height() - bottom)
        finally:
            text.close()
            page.close()
            pdf.close()

    def test_placeholder_and_empty_table_targets_are_in_correct_cells(self):
        row = next(r for r in self.scan['mappings'] if r['category_id'] == 'I01')
        result = self.preview.describe(self.snapshot(row['locator_id']))
        self.assertEqual([2], result['target_pages'])
        self.assertGreater(result['markers'][0]['x'], .5)
        self.assertEqual('cell', result['markers'][0]['kind'])
        table_box = self.preview._page_tables(self.preview.cache[result['preview_id']], 2)[0]
        self.assertLess(table_box['ys'][1], result['markers'][0]['y'])
        self.assertLess(result['markers'][0]['y'], table_box['ys'][2])
        self.assertIn('占位符', result['description'])
        heading = next(h for h in self.scan['headings'] if h['title'] == '身份证附件')
        pair = {'mode': 'side_by_side', 'files': [{}, {}]}
        cells = self.preview.describe(self.snapshot(heading['id'], materials=[pair]))
        self.assertEqual([3], cells['target_pages'])
        self.assertEqual(['cell', 'cell'], [m['kind'] for m in cells['markers']])
        self.assertLess(cells['markers'][0]['x'], cells['markers'][1]['x'])

    def test_marker_is_after_own_content_and_agrees_with_export(self):
        document = engine._read_template(self.source)
        records, _ = engine._inspect(document, self.categories)
        record = next(r for r in records if r['title'] == '公司介绍')
        self.assertEqual('第一处原文，材料必须插在这段之后。', engine._text(record['_anchor']))
        engine._insert_materials(document, record, [{'mode': 'text', 'content': '新插入材料'}])
        paragraph = next(p for p in document.paragraphs if p.text == '新插入材料')
        self.assertEqual(record['_anchor'], paragraph._p.getprevious())
        preview = self.preview.describe(self.snapshot(record['id']))
        names = self.preview.cache[preview['preview_id']]['names']
        end_name = names[record['_path']]['E']
        title_location = self.preview.cache[preview['preview_id']]['destinations'][end_name]
        self.assertGreater(preview['markers'][0]['y'], title_location['y'])

    def test_append_after_empty_table_uses_rendered_bottom_border(self):
        heading = next(h for h in self.scan['headings'] if h['title'] == '身份证附件')
        result = self.preview.describe(self.snapshot(heading['id'], materials=[
            {'mode': 'text', 'content': '追加材料'}]))
        self.assertEqual([3], result['target_pages'])
        self.assertEqual('table', result['markers'][0]['kind'])
        pair = self.preview.describe(self.snapshot(heading['id'], materials=[
            {'mode': 'side_by_side', 'files': [{}, {}]}]))
        self.assertGreater(result['markers'][0]['y'], max(m['y'] for m in pair['markers']))

    def test_bookmarks_do_not_change_rendered_page_layout(self):
        import pypdfium2 as pdfium
        from pypdf import PdfReader
        row = next(r for r in self.scan['mappings'] if r['category_id'] == 'company_profile')
        result = self.preview.describe(self.snapshot(row['locator_id']))
        baseline_dir = self.root / 'baseline'
        baseline_dir.mkdir()
        baseline = _render_pdf(self.source, baseline_dir)
        self.assertEqual(len(PdfReader(baseline).pages), result['page_count'])
        pdf = pdfium.PdfDocument(baseline)
        try:
            for index in range(len(pdf)):
                page = pdf[index]
                try:
                    width, height = page.get_size()
                    scale = min(1.7, 1700 / width, (5_000_000 / (width * height)) ** .5)
                    bitmap = page.render(scale=scale)
                    try:
                        from PIL import Image
                        actual = Image.open(io.BytesIO(self.preview.page(result['preview_id'], index + 1)))
                        self.assertEqual(bitmap.to_pil().tobytes(), actual.tobytes())
                    finally:
                        bitmap.close()
                finally:
                    page.close()
        finally:
            pdf.close()

    def test_card_previews_share_one_recognition_and_detect_changed_files(self):
        rows = [r for r in self.scan['mappings'] if r['category_id'] == 'company_profile']
        self.preview._recognized.clear()
        with patch.object(engine, '_inspect', wraps=engine._inspect) as inspect:
            first, second = [self.preview.describe(self.snapshot(r['locator_id'])) for r in rows]
            self.assertEqual(1, inspect.call_count)
            self.preview.describe(self.snapshot(rows[0]['locator_id'], categories=self.categories[:-1]))
            self.assertEqual(2, inspect.call_count)
        self.assertEqual(([1], [2]), (first['target_pages'], second['target_pages']))
        copy_path = self.root / 'changed.docx'
        copy_path.write_bytes(self.raw)
        self.preview.describe(self.snapshot(rows[0]['locator_id'], template=copy_path))
        Document().save(copy_path)
        with self.assertRaisesRegex(ValueError, '已变化'):
            self.preview.describe(self.snapshot(rows[0]['locator_id'], template=copy_path))

    def test_invalid_locator_stale_hash_and_bad_page_fail_explicitly(self):
        with self.assertRaisesRegex(ValueError, '已失效'):
            self.preview.describe(self.snapshot('p:0:0000000000000000'))
        snapshot = self.snapshot(self.scan['headings'][0]['id'])
        snapshot['template_hash'] = '0' * 64
        with self.assertRaisesRegex(ValueError, '已变化'):
            self.preview.describe(snapshot)
        row = next(r for r in self.scan['mappings'] if r['category_id'] == 'company_profile')
        result = self.preview.describe(self.snapshot(row['locator_id']))
        with self.assertRaisesRegex(ValueError, '超出范围'):
            self.preview.page(result['preview_id'], result['page_count'] + 1)


class LocationAPITests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(prefix='bid-location-api-')
        source = Path(self.tmp.name) / 'source.docx'
        make_fixture(source)
        self.httpd = server.create_server(port=0, data_dir=Path(self.tmp.name) / 'data')
        self.thread = threading.Thread(target=self.httpd.serve_forever, daemon=True)
        self.thread.start()
        self.base = f'http://127.0.0.1:{self.httpd.server_port}'
        self.project = self.httpd.store.create_project({
            'name': '预览核对', 'filename': source.name,
            'data': base64.b64encode(source.read_bytes()).decode()})

    def tearDown(self):
        self.httpd.shutdown()
        self.httpd.server_close()
        self.thread.join()
        self.httpd.store.close()
        self.tmp.cleanup()

    def request(self, method, suffix, body=None, origin=None):
        request = urllib.request.Request(self.base + suffix, method=method,
            data=json.dumps(body).encode() if body is not None else None,
            headers={'Content-Type': 'application/json', 'Origin': origin or self.base})
        with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(request, timeout=20) as response:
            return json.loads(response.read())

    @unittest.skipUnless(find_converter(), 'Local LibreOffice is required for page rendering')
    def test_special_preview_before_and_after_unsaved_position_selection(self):
        store = self.httpd.store
        store.update_mappings(self.project['id'], {'mappings': [
            {'id': row['id'], 'action': 'skip', 'confirmed': True}
            for row in self.project['mappings']]})
        project = store.select_special_categories(self.project['id'], {'industry_ids': ['insurance']})
        row = next(r for r in project['mappings'] if r['category_id'] == 'I01')
        material = store.material({'name': '跨行业关联证明', 'category_id': 'F01',
            'linked_category_ids': ['I01'], 'subject': server.SUBJECT,
            'mode': 'text', 'content': '用于核对的专用材料', 'valid_until': '长期'})
        original = copy.deepcopy(store.require('project', project['id']))
        path = f'/api/projects/{project["id"]}/location-preview'
        body = {'mapping_id': row['id'], 'locator_id': '', 'action': 'fill',
                'material_ids': [material['id']]}
        unselected = self.request('POST', path, body)
        self.assertEqual([], unselected['markers'])
        self.assertEqual(1, unselected['initial_page'])
        locator = [h for h in project['headings'] if h['title'] == '公司介绍'][1]['id']
        located = self.request('POST', path, dict(body, locator_id=locator))
        self.assertEqual([2], located['target_pages'])
        self.assertEqual('after', located['markers'][0]['kind'])
        self.assertEqual(unselected['preview_id'], located['preview_id'])
        url = self.base + path + '/' + located['preview_id'] + '/pages/2'
        with urllib.request.build_opener(urllib.request.ProxyHandler({})).open(url, timeout=5) as response:
            self.assertEqual('image/png', response.headers.get_content_type())
            self.assertTrue(response.read().startswith(b'\x89PNG'))
        self.assertEqual(original, store.require('project', project['id']))

    def test_unsaved_position_preview_is_read_only_and_releases_database_lock(self):
        store = self.httpd.store
        rows = [m for m in self.project['mappings'] if m['category_id'] == 'company_profile']
        original = copy.deepcopy(store.require('project', self.project['id']))
        body = {'mapping_id': rows[0]['id'], 'locator_id': rows[1]['locator_id'], 'action': 'fill', 'material_ids': []}
        entered, release = threading.Event(), threading.Event()
        observed, responses, errors = [], [], []
        def render(snapshot):
            observed.append(snapshot)
            entered.set()
            if not release.wait(4):
                raise ValueError('Preview blocked unrelated requests')
            return {'title': '公司介绍'}
        def post():
            try:
                responses.append(self.request('POST', f'/api/projects/{self.project["id"]}/location-preview', body))
            except Exception as error:
                errors.append(error)
        with patch.object(store.location_preview, 'describe', side_effect=render):
            worker = threading.Thread(target=post)
            worker.start()
            self.assertTrue(entered.wait(2))
            try:
                self.assertEqual(1, self.request('GET', '/api/bootstrap')['stats']['project_count'])
            finally:
                release.set()
                worker.join(5)
        self.assertEqual([], errors)
        self.assertEqual([{'title': '公司介绍'}], responses)
        self.assertEqual(rows[1]['locator_id'], observed[0]['locator_id'])
        self.assertEqual(original, store.require('project', self.project['id']))

    def test_preview_rejects_unknown_mapping_foreign_locator_and_cross_origin(self):
        path = f'/api/projects/{self.project["id"]}/location-preview'
        foreign = self.httpd.store.material({'name': '其他行业材料', 'category_id': 'F01',
            'subject': server.SUBJECT, 'mode': 'text', 'content': '其他行业证明', 'valid_until': '长期'})
        for body in ({'mapping_id': 'foreign'},
                     {'mapping_id': self.project['mappings'][0]['id'], 'locator_id': 'foreign'},
                     {'mapping_id': self.project['mappings'][0]['id'], 'locator_id': '',
                      'material_ids': [foreign['id']]}):
            with self.assertRaises(urllib.error.HTTPError) as error:
                self.request('POST', path, body)
            self.assertEqual(400, error.exception.code)
        with self.assertRaises(urllib.error.HTTPError) as error:
            self.request('POST', path, {}, origin='https://example.com')
        self.assertEqual(403, error.exception.code)


if __name__ == '__main__':
    unittest.main()
