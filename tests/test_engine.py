import copy
import hashlib
import io
import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import engine
from credit_service import SOURCES as CREDIT_SOURCES


class EngineTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.root = Path(self.temp.name)
        self.template = self.root / "source.docx"
        self.output = self.root / "output.docx"

    def tearDown(self):
        self.temp.cleanup()

    def categories(self, *ids):
        return [copy.deepcopy(c) for c in engine.DEFAULT_CATEGORIES if c["id"] in ids]

    def doc(self, *titles):
        document = Document()
        for title in titles:
            document.add_heading(title, level=1)
        document.save(self.template)
        return document

    def material(self, category="company_profile", **updates):
        value = dict(id="material", name="公司介绍", category_id=category,
                     subject=engine.EXPECTED_SUBJECT, mode="text", content="确认后的公司介绍内容",
                     years=[], valid_until="长期", files=[], revision=1)
        value.update(updates)
        return value

    def scan(self, categories, materials):
        return engine.scan_document(self.template, categories, materials)

    def confirmed(self, rows):
        return [dict(row, confirmed=True) for row in rows]

    def test_default_seven_and_distinct_repeated_locations(self):
        document = Document()
        document.add_paragraph("公司介绍\t8", style="TOC Heading")
        document.add_heading("公司介绍", 1)
        document.add_paragraph("旧版一")
        document.add_heading("公司介绍", 1)
        document.add_paragraph("旧版二")
        document.save(self.template)
        scanned = self.scan(engine.DEFAULT_CATEGORIES, [])
        company = [r for r in scanned["mappings"] if r["category_id"] == "company_profile"]
        self.assertEqual(len(company), 2)
        self.assertEqual(len(scanned["mappings"]), 8)
        self.assertEqual(len({r["locator_id"] for r in company}), 2)
        self.assertEqual(scanned["mappings"], self.scan(engine.DEFAULT_CATEGORIES, [])["mappings"])

    def test_only_enabled_common_and_extension(self):
        self.doc("自定义证明", "公司介绍")
        categories = self.categories("company_profile")
        categories[0]["enabled"] = False
        categories += [dict(id="custom", name="自定义证明", aliases=[], enabled=True, layer="common"),
                       dict(id="insurance", name="公司介绍", aliases=[], enabled=True, layer="special")]
        rows = self.scan(categories, [])["mappings"]
        self.assertEqual([r["category_id"] for r in rows], ["custom"])

    def test_year_completeness_and_metadata_override(self):
        self.doc("2023-2024年财务审计报告")
        categories = self.categories("audit_report")
        first = self.material("audit_report", id="2023", name="2023-2024年财务审计报告", years=[2023])
        second = self.material("audit_report", id="2024", years=[2024])
        scanned = self.scan(categories, [first, second])
        row = scanned["mappings"][0]
        self.assertEqual(row["status"], "ready")
        self.assertEqual(row["material_ids"], ["2023", "2024"])
        row.update(material_ids=["2023"], required_years=[2023])
        checked = engine.validate_mappings([row], [first], scanned["headings"], categories)[0]
        self.assertEqual(checked["required_years"], [2023, 2024])
        self.assertEqual(checked["status"], "blocked")
        self.assertTrue(any("2024" in x for x in checked["issues"]))
        first["years"] = [2025]
        checked = engine.validate_mappings([row], [first], scanned["headings"], categories)[0]
        self.assertEqual(checked["status"], "blocked")

    def test_unknown_audit_year_is_not_inferred_from_name(self):
        self.doc("审计报告")
        categories = self.categories("audit_report")
        material = self.material("audit_report", name="2024年审计报告", years=[])
        scanned = self.scan(categories, [material])
        row = dict(scanned["mappings"][0], action="fill", material_ids=[material["id"]])
        row = engine.validate_mappings([row], [material], scanned["headings"], categories)[0]
        self.assertEqual(row["required_years"], [])
        self.assertEqual(row["status"], "blocked")
        self.assertTrue(any("不默认" in w for w in row["warnings"]))

    def test_category_subject_expiry_and_missing_id_are_hard_constraints(self):
        self.doc("公司介绍")
        categories = self.categories("company_profile")
        material = self.material()
        scanned = self.scan(categories, [material])
        original = scanned["mappings"][0]
        for field, value in [("category_id", "business_license"), ("subject", "其他公司"), ("subject", ""),
                             ("valid_until", "2001-01-01"), ("valid_until", "不详")]:
            altered = dict(material, **{field: value})
            checked = engine.validate_mappings([original], [altered], scanned["headings"], categories)[0]
            self.assertEqual(checked["status"], "blocked", (field, value))
        missing = dict(original, material_ids=["material", "missing"])
        checked = engine.validate_mappings([missing], [material], scanned["headings"], categories)[0]
        self.assertEqual(checked["status"], "blocked")
        no_expiry = dict(material, valid_until="")
        checked = engine.validate_mappings([original], [no_expiry], scanned["headings"], categories)[0]
        self.assertEqual(checked["status"], "ready")
        self.assertTrue(any("有效期未填写" in warning for warning in checked["warnings"]))
        annual = dict(material, valid_until="下一年度更新")
        checked = engine.validate_mappings([original], [annual], scanned["headings"], categories)[0]
        self.assertEqual(checked["status"], "ready")
        self.assertTrue(any("未提供明确有效期" in warning for warning in checked["warnings"]))

    def test_attachment_cannot_replace_proof_body(self):
        document = self.doc("法定代表人（负责人）身份证明")
        material = self.material("legal_identity", name="法人身份证", description="证件附件")
        categories = self.categories("legal_identity")
        scanned = self.scan(categories, [material])
        row = dict(scanned["mappings"][0], action="fill", material_ids=["material"])
        checked = engine.validate_mappings([row], [material], scanned["headings"], categories)[0]
        self.assertEqual(checked["status"], "blocked")
        document.add_paragraph("兹证明 某某 系北京腾宝科技有限公司的法定代表人。")
        document.save(self.template)
        scanned = self.scan(categories, [material])
        self.assertEqual(scanned["mappings"][0]["status"], "ready")
        self.assertTrue(any("附件" in warning for warning in scanned["mappings"][0]["warnings"]))

    def test_confirmation_and_blocked_fill_gate_even_for_draft(self):
        self.doc("公司介绍")
        material = self.material()
        scanned = self.scan(self.categories("company_profile"), [material])
        with self.assertRaisesRegex(ValueError, "确认"):
            engine.export_document(self.template, self.output, scanned["mappings"], [material])
        self.assertFalse(self.output.exists())
        row = dict(scanned["mappings"][0], material_ids=["material", "missing"], confirmed=True)
        with self.assertRaisesRegex(ValueError, "不能导出"):
            engine.export_document(self.template, self.output, [row], [material], "draft")
        self.assertFalse(self.output.exists())

    def test_draft_missing_markers_and_reviewed_refusal(self):
        self.doc("公司介绍")
        scanned = self.scan(self.categories("company_profile", "tax_social"), [])
        rows = self.confirmed(scanned["mappings"])
        result = engine.export_document(self.template, self.output, rows, [])
        self.assertEqual(result["missing_count"], 2)
        text = "\n".join(p.text for p in Document(self.output).paragraphs)
        self.assertIn("待补初稿", text)
        self.assertIn("【待补：公司介绍】", text)
        self.assertIn("原模板未定位", text)
        with self.assertRaisesRegex(ValueError, "材料齐全"):
            engine.export_document(self.template, self.root / "reviewed.docx", rows, [], "reviewed")

    def test_repeated_headings_keep_own_content_and_source_unchanged(self):
        document = Document()
        for old in ("第一处原有内容", "第二处原有内容"):
            document.add_heading("公司介绍", 1)
            document.add_paragraph(old)
        document.save(self.template)
        before = hashlib.sha256(self.template.read_bytes()).hexdigest()
        material = self.material(content="新公司介绍")
        scanned = self.scan(self.categories("company_profile"), [material])
        rows = self.confirmed(scanned["mappings"])
        result = engine.export_document(self.template, self.output, rows, [material], "reviewed")
        self.assertEqual(result["filled_count"], 2)
        text = [p.text for p in Document(self.output).paragraphs]
        self.assertEqual(text, ["公司介绍", "第一处原有内容", "新公司介绍", "公司介绍", "第二处原有内容", "新公司介绍"])
        self.assertEqual(before, hashlib.sha256(self.template.read_bytes()).hexdigest())

    def test_stale_locator_refused(self):
        document = self.doc("公司介绍")
        material = self.material()
        scanned = self.scan(self.categories("company_profile"), [material])
        document.paragraphs[0].text = "公司介绍（新模板）"
        document.save(self.template)
        with self.assertRaisesRegex(ValueError, "填充位置"):
            engine.export_document(self.template, self.output, self.confirmed(scanned["mappings"]), [material])

    def test_image_order_rotation_and_missing_file(self):
        self.doc("营业执照")
        red = self.root / "red.png"
        blue = self.root / "blue.png"
        Image.new("RGB", (30, 10), "red").save(red)
        Image.new("RGB", (20, 15), "blue").save(blue)
        old_red = red.read_bytes()
        material = self.material("business_license", mode="side_by_side", files=[
            {"id": "blue", "name": "blue.png", "path": str(blue), "rotation": 0},
            {"id": "red", "name": "red.png", "path": str(red), "rotation": 90}])
        scanned = self.scan(self.categories("business_license"), [material])
        result = engine.export_document(self.template, self.output, self.confirmed(scanned["mappings"]), [material], "reviewed")
        self.assertEqual(result["results"][0]["file_count"], 2)
        with zipfile.ZipFile(self.output) as archive:
            images = [Image.open(io.BytesIO(archive.read(name))) for name in archive.namelist() if name.startswith("word/media/")]
            self.assertEqual([i.size for i in images], [(20, 15), (10, 30)])
            self.assertEqual(images[0].getpixel((0, 0)), (0, 0, 255))
        self.assertEqual(red.read_bytes(), old_red)
        blue.unlink()
        checked = engine.validate_mappings(scanned["mappings"], [material], scanned["headings"])[0]
        self.assertEqual(checked["status"], "blocked")

    def test_multiple_placeholders_and_no_double_heading_fill(self):
        document = Document()
        document.add_heading("公司介绍", 1)
        paragraph = document.add_paragraph("前文 ")
        paragraph.add_run("{{公司")
        paragraph.add_run("介绍}} 与 {{公司介绍}} 后文")
        document.save(self.template)
        material = self.material(content="新介绍")
        scanned = self.scan(self.categories("company_profile"), [material])
        self.assertEqual(len(scanned["mappings"]), 2)
        self.assertTrue(all(r["locator_id"].startswith("ph:") for r in scanned["mappings"]))
        engine.export_document(self.template, self.output, self.confirmed(scanned["mappings"]), [material], "reviewed")
        text = "\n".join(p.text for p in Document(self.output).paragraphs)
        self.assertNotIn("{{", text)
        self.assertIn("前文  与  后文", text)
        self.assertEqual(text.count("新介绍"), 2)

    def test_side_by_side_uses_only_own_empty_table(self):
        document = self.doc("营业执照")
        table = document.add_table(rows=1, cols=2)
        document.add_heading("下一个章节", 1)
        document.add_table(rows=1, cols=2)
        document.save(self.template)
        image = self.root / "scan.png"
        Image.new("RGB", (50, 30), "green").save(image)
        material = self.material("business_license", mode="side_by_side", files=[
            {"id": "front", "name": "front.png", "path": str(image), "rotation": 0},
            {"id": "back", "name": "back.png", "path": str(image), "rotation": 0}])
        scanned = self.scan(self.categories("business_license"), [material])
        engine.export_document(self.template, self.output, self.confirmed(scanned["mappings"]), [material], "reviewed")
        rendered = Document(self.output)
        from docx.oxml.ns import qn
        self.assertEqual(len(list(rendered.tables[0]._tbl.iter(qn("w:drawing")))), 2)
        self.assertEqual(len(list(rendered.tables[1]._tbl.iter(qn("w:drawing")))), 0)
        self.assertEqual(len(rendered.tables), 2)

    def test_credit_aliases_export_three_images_under_exact_child_titles(self):
        titles = [
            "供应商未被列入失信被执行人查询截图",
            "供应商未被列入重大税收违法失信主体查询截图",
            "供应商未被列入政府采购严重违法失信行为记录名单查询截图",
        ]
        parent = "供应商未被列入失信被执行人、重大税收违法失信主体、政府采购严重违法失信行为记录名单查询截图"
        categories = [dict(copy.deepcopy(source), enabled=True, layer="common") for source in CREDIT_SOURCES]
        original_categories = copy.deepcopy(categories)
        colors = [(255, 0, 0), (0, 128, 0), (0, 0, 255)]
        materials = []
        for category, color in zip(categories, colors):
            image = self.root / (category["id"] + ".png")
            Image.new("RGB", (40, 20), color).save(image)
            materials.append(self.material(category["id"], id=category["id"], name=category["name"],
                mode="page", files=[{"id": category["id"], "name": image.name,
                                     "path": str(image), "rotation": 0}]))

        # Real templates can have visible titles without Word heading styles;
        # their category aliases must retain section boundaries during export.
        for styled in (True, False):
            with self.subTest(styled=styled):
                document = Document()
                document.add_heading(parent, 1)
                document.add_paragraph("本节仅保留查询说明。")
                for title in titles:
                    document.add_paragraph(title, style="Heading 2" if styled else "Normal")
                    document.add_paragraph("")
                document.add_heading("后续章节", 1)
                document.add_paragraph("后续正文保留。")
                document.save(self.template)
                source_hash = hashlib.sha256(self.template.read_bytes()).hexdigest()
                scanned = self.scan(categories, materials)
                self.assertEqual([row["title"] for row in scanned["mappings"]], titles)
                self.assertEqual(len({row["locator_id"] for row in scanned["mappings"]}), 3)
                self.assertTrue(all(row["status"] == "ready" for row in scanned["mappings"]))
                parent_heading = next(h for h in scanned["headings"] if h["title"] == parent)
                self.assertEqual(parent_heading["category_ids"], [])

                result = engine.export_document(self.template, self.output,
                    self.confirmed(scanned["mappings"]), materials, "reviewed", categories=categories)
                self.assertEqual(result["filled_count"], 3)
                self.assertEqual([row["file_count"] for row in result["results"]], [1, 1, 1])
                self.assertFalse(any("人工指定位置" in warning for warning in result["warnings"]))
                exported = Document(self.output)
                pictures = {title: [] for title in [parent, *titles, "后续章节"]}
                active = parent
                for paragraph in exported.paragraphs:
                    if paragraph.text in pictures:
                        active = paragraph.text
                    for blip in paragraph._p.iter(qn("a:blip")):
                        part = exported.part.related_parts[blip.get(qn("r:embed"))]
                        with Image.open(io.BytesIO(part.blob)) as image:
                            pictures[active].append(image.getpixel((0, 0)))
                self.assertEqual(pictures, {parent: [], **{title: [color] for title, color in zip(titles, colors)}, "后续章节": []})
                self.assertEqual(source_hash, hashlib.sha256(self.template.read_bytes()).hexdigest())
                self.assertEqual(categories, original_categories)

    def test_disabled_credit_categories_leave_default_seven_unchanged(self):
        categories = copy.deepcopy(engine.DEFAULT_CATEGORIES)
        categories += [dict(copy.deepcopy(source), enabled=False, layer="common") for source in CREDIT_SOURCES]
        self.doc(*(source["name"] for source in CREDIT_SOURCES))
        scanned = self.scan(categories, [])
        self.assertEqual([row["category_id"] for row in scanned["mappings"]],
                         [category["id"] for category in engine.DEFAULT_CATEGORIES])
        self.assertEqual(len(scanned["mappings"]), 7)
        result = engine.export_document(self.template, self.output,
            self.confirmed(scanned["mappings"]), [], categories=categories)
        self.assertEqual(result["filled_count"], 0)
        self.assertEqual(result["missing_count"], 7)
        self.assertEqual(len(Document(self.output).inline_shapes), 0)

    def test_export_current_categories_reject_disabled_or_removed_mapping(self):
        category = dict(copy.deepcopy(CREDIT_SOURCES[0]), enabled=True, layer="common")
        self.doc(category["name"])
        material = self.material(category["id"])
        rows = self.confirmed(self.scan([category], [material])["mappings"])
        for current in ([dict(category, enabled=False)], [], [dict(category, layer="special")]):
            with self.subTest(categories=current):
                with self.assertRaisesRegex(ValueError, "已停用、分类已变更或未纳入本项目"):
                    engine.export_document(self.template, self.output, rows, [material], categories=current)
                self.assertFalse(self.output.exists())


if __name__ == "__main__":
    unittest.main()
