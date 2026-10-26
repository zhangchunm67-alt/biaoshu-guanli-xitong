"""Review cases: repeated real sections, table directories and format suffixes."""
import copy
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

import engine


class TocBoundaryTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='bid-toc-boundary-')
        self.root = Path(self.temp.name)
        self.source = self.root / 'framework.docx'
        self.output = self.root / 'filled.docx'
        self.categories = [copy.deepcopy(c) for c in engine.DEFAULT_CATEGORIES
                           if c['id'] in {'authorization', 'business_license', 'company_profile'}]
        self.materials = []
        for cid, color in [('authorization', 'red'), ('business_license', 'green'), ('company_profile', 'blue')]:
            path = self.root / (cid + '.png')
            Image.new('RGB', (120, 80), color).save(path)
            self.materials.append(dict(id=cid, category_id=cid, name=cid + '图片',
                                       subject=engine.EXPECTED_SUBJECT, mode='page',
                                       files=[dict(path=str(path), name=path.name, rotation=0)],
                                       years=[], valid_until='长期'))

    def tearDown(self):
        self.temp.cleanup()

    def scan(self, document):
        document.save(self.source)
        return engine.scan_document(self.source, self.categories, self.materials)

    def locator(self, document, paragraph):
        return ('p:' + '.'.join(map(str, engine._xml_path(paragraph._p, document.element.body))) +
                ':' + engine._digest(paragraph._p))

    def export(self, scanned, category_ids):
        original = self.source.read_bytes()
        rows = [dict(row, action='fill', confirmed=True, material_ids=[row['category_id']])
                for row in scanned['mappings'] if row['category_id'] in category_ids]
        engine.export_document(self.source, self.output, rows, self.materials,
                               mode='reviewed', categories=self.categories)
        self.assertEqual(original, self.source.read_bytes())
        return Document(self.output)

    def test_empty_real_sections_are_not_mistaken_for_directory_of_later_section(self):
        document = Document()
        document.add_paragraph('一、商务部分')
        first_license = document.add_paragraph('营业执照')
        first_profile = document.add_paragraph('公司简介')
        document.add_paragraph('二、服务方案部分')
        second_license = document.add_paragraph('营业执照')
        document.add_paragraph('本节原有内容保持不变。')
        second_profile = document.add_paragraph('公司简介')
        document.add_paragraph('本节原有内容保持不变。')
        document.add_paragraph('三、结束附件')
        scanned = self.scan(document)
        for cid, paragraphs in [('business_license', [first_license, second_license]),
                                ('company_profile', [first_profile, second_profile])]:
            actual = {row['locator_id'] for row in scanned['mappings'] if row['category_id'] == cid}
            self.assertEqual({self.locator(document, paragraph) for paragraph in paragraphs}, actual)
        result = self.export(scanned, {'business_license', 'company_profile'})
        children = list(result.element.body)
        middle = next(i for i, node in enumerate(children) if engine._text(node) == '二、服务方案部分')
        self.assertEqual(2, sum(len(list(node.iter(qn('a:blip')))) for node in children[:middle]))
        self.assertEqual(2, sum(len(list(node.iter(qn('a:blip')))) for node in children[middle:]))

    def test_directory_table_cells_are_not_candidates_and_images_go_to_body(self):
        document = Document()
        document.add_paragraph('目录')
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = '授权委托书'
        table.cell(0, 1).text = '3'
        table.cell(1, 0).text = '营业执照'
        table.cell(1, 1).text = '4'
        document.add_paragraph('授权委托书格式')
        document.add_paragraph('本授权书声明：现委托李四办理投标事项。')
        agent_id = document.add_paragraph('附：被授权人身份证复印件')
        license_heading = document.add_paragraph('营业执照（图片粘贴位置）')
        document.add_paragraph('三、结束附件')
        scanned = self.scan(document)
        selectable = {heading['id'] for heading in scanned['headings']}
        selected = {row['locator_id'] for row in scanned['mappings']}
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    locator = self.locator(document, paragraph)
                    for target in [locator, *['at:' + c['id'] + ':' + locator for c in self.categories]]:
                        self.assertNotIn(target, selectable)
                        self.assertNotIn(target, selected)
        for cid, locator in [('authorization', 'at:authorization:' + self.locator(document, agent_id)),
                             ('business_license', self.locator(document, license_heading))]:
            matching = [row for row in scanned['mappings'] if row['category_id'] == cid]
            self.assertEqual(1, len(matching))
            self.assertEqual(locator, matching[0]['locator_id'])
        result = self.export(scanned, {'authorization', 'business_license'})
        self.assertFalse(list(result.tables[0]._tbl.iter(qn('a:blip'))))
        self.assertEqual([['授权委托书', '3'], ['营业执照', '4']],
                         [[cell.text for cell in row.cells] for row in result.tables[0].rows])
        children = list(result.element.body)
        start = next(i for i, node in enumerate(children) if engine._text(node) == '授权委托书格式')
        self.assertFalse(any(list(node.iter(qn('a:blip'))) for node in children[:start]))
        self.assertEqual(2, sum(len(list(node.iter(qn('a:blip')))) for node in children[start:]))

    def test_authorization_format_title_uses_real_body_instead_of_appendix(self):
        for directory_title in ('目录', '目录：', '一、目录'):
            with self.subTest(directory_title=directory_title):
                document = Document()
                document.add_paragraph(directory_title)
                directory = document.add_paragraph('授权委托书')
                document.add_paragraph('（四）授权委托书格式')
                document.add_paragraph('本授权书声明：现委托李四办理投标事项。')
                agent_id = document.add_paragraph('附：被授权人身份证复印件')
                document.add_paragraph('（五）结束附件')
                scanned = self.scan(document)
                matching = [row for row in scanned['mappings'] if row['category_id'] == 'authorization']
                self.assertEqual(1, len(matching))
                self.assertEqual('at:authorization:' + self.locator(document, agent_id), matching[0]['locator_id'])
                self.assertNotIn(self.locator(document, directory), {h['id'] for h in scanned['headings']})
                result = self.export(scanned, {'authorization'})
                children = list(result.element.body)
                start = next(i for i, node in enumerate(children) if engine._text(node) == '（四）授权委托书格式')
                end = next(i for i, node in enumerate(children) if engine._text(node) == '（五）结束附件')
                self.assertEqual(1, sum(len(list(node.iter(qn('a:blip')))) for node in children[start:end]))
                self.assertFalse(any('通用材料附件' in p.text for p in result.paragraphs))

    def test_contents_before_page_index_does_not_become_authorization_form(self):
        document = Document()
        document.add_paragraph('第六章 附件（竞价响应文件格式）')
        document.add_paragraph('年 月 日')
        entries = [document.add_paragraph(t) for t in ['报价一览表', '商务条款偏离表',
            '法定代表人授权委托书', '资格证明文件', '合同条款偏离表']]
        for text in ['项目人员团队', '十二、其他内容', '评审索引表', '序号', '评审因素',
                     '对应页码', '备注或说明', '注：格式仅供参考。', '附件1：报价格式',
                     '报价一览表', '此处保留报价内容。', '附件2：授权格式']:
            document.add_paragraph(text)
        document.add_paragraph('法定代表人授权委托书')
        document.add_paragraph('本授权书声明：现委托李四办理投标事项。')
        agent_id = document.add_paragraph('附：被授权人身份证复印件')
        scanned = self.scan(document)
        rows = [r for r in scanned['mappings'] if r['category_id'] == 'authorization']
        self.assertEqual(['at:authorization:' + self.locator(document, agent_id)], [r['locator_id'] for r in rows])
        self.assertTrue(all(self.locator(document, p) not in {h['id'] for h in scanned['headings']} for p in entries))


if __name__ == '__main__':
    unittest.main()
