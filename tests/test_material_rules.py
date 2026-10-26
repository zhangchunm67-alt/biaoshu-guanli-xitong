"""Evidence type, period and real output boundaries within shared categories."""
import copy
import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from PIL import Image

import engine
from material_rules import months


class MaterialRuleTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='bid-material-rules-')
        self.root = Path(self.temp.name)
        self.path = self.root / 'source.docx'
        self.categories = copy.deepcopy(engine.DEFAULT_CATEGORIES) + [dict(
            id='text_template', name='文书模板', aliases=['投标函', '授权书', '承诺函'],
            layer='special', enabled=True)]

    def tearDown(self):
        self.temp.cleanup()

    def material(self, name, cid='tax_social', **extra):
        path = self.root / (str(len(list(self.root.glob('*.png')))) + '.png')
        Image.new('RGB', (120, 80), (len(list(self.root.glob('*.png'))) * 30, 90, 120)).save(path)
        return dict(id=path.stem, name=name, category_id=cid, subject=engine.EXPECTED_SUBJECT,
                    mode='page', files=[dict(path=str(path), name=path.name)],
                    years=[2025, 2026], valid_until='长期', **extra)

    def scan(self, labels, materials):
        doc = Document()
        doc.add_heading('第五章 投标文件格式', 1)
        for label in labels:
            doc.add_heading(label, 2)
            doc.add_paragraph('本栏目原有内容。')
        doc.save(self.path)
        return engine.scan_document(self.path, self.categories, materials, ['text_template'])

    def tax_rows(self, scanned):
        return [r for r in scanned['mappings'] if r['category_id'] == 'tax_social']

    def test_tax_and_social_cannot_substitute_for_each_other(self):
        tax = self.material('2026年9月完税证明')
        social = self.material('2026年8月社保缴费凭证')
        scanned = self.scan(['纳税证明', '社保缴费证明'], [tax, social])
        rows = self.tax_rows(scanned)
        self.assertEqual(2, len(rows))
        for row, correct, wrong in zip(rows, [tax, social], [social, tax]):
            self.assertEqual([correct['id']], row['material_ids'])
            checked = engine.validate_mappings([dict(row, material_ids=[wrong['id']], action='fill')],
                [tax, social], scanned['headings'], self.categories)[0]
            self.assertEqual('blocked', checked['status'])
            self.assertTrue(any('类型不符' in s for s in checked['issues']))

    def test_combined_slot_requires_both_and_selects_unique_complement(self):
        tax = self.material('2026年9月完税证明')
        social = self.material('2026年8月社保缴费凭证')
        scanned = self.scan(['税收和社会保障资金缴纳凭证'], [tax, social])
        row = self.tax_rows(scanned)[0]
        self.assertEqual('ready', row['status'])
        self.assertEqual({tax['id'], social['id']}, set(row['material_ids']))
        checked = engine.validate_mappings([dict(row, material_ids=[tax['id']])],
            [tax], scanned['headings'], self.categories)[0]
        self.assertEqual('blocked', checked['status'])
        self.assertIn('缺少要求凭证：社保。', checked['issues'])
        partial = self.tax_rows(self.scan(['税收和社会保障资金缴纳凭证'], [tax]))[0]
        self.assertEqual([tax['id']], partial['material_ids'])
        self.assertIn('缺少要求凭证：社保。', partial['issues'])

    def test_specific_month_uses_period_not_year_or_payment_date(self):
        august = self.material('2026年8月社保缴费凭证', description='交易日期：2026年9月18日')
        september = self.material('2026年9月社保缴费凭证')
        scanned = self.scan(['2026年8月社保缴费证明'], [august, september])
        row = self.tax_rows(scanned)[0]
        self.assertEqual([august['id']], row['material_ids'])
        bad = engine.validate_mappings([dict(row, material_ids=[september['id']])],
            [august, september], scanned['headings'], self.categories)[0]
        self.assertEqual('blocked', bad['status'])
        unknown = self.material('社保缴费凭证')
        self.assertFalse(self.tax_rows(self.scan(['2026年8月社保缴费证明'], [unknown]))[0]['material_ids'])

    def test_period_range_and_partial_coverage(self):
        record = self.material('单位社保缴费证明（2025年9月至2026年8月）')
        row = self.tax_rows(self.scan(['2026年8月社保缴费证明'], [record]))[0]
        self.assertEqual([record['id']], row['material_ids'])
        august = self.material('2026年8月社保缴费凭证')
        september = self.material('2026年9月社保缴费凭证')
        scanned = self.scan(['2026年8月至9月社保缴费证明'], [august, september])
        self.assertEqual('ready', self.tax_rows(scanned)[0]['status'])
        self.assertEqual(12, len(months(record['name'])))
        self.assertEqual(2, len(months('2025年7、8月')))
        row = self.tax_rows(self.scan(['2026年8月至9月社保缴费证明'], [august]))[0]
        self.assertEqual(([august['id']], 'blocked'), (row['material_ids'], row['status']))
        self.assertIn('缺少要求月份：社保 2026年9月。', row['issues'])

    def test_short_month_ranges_share_the_year(self):
        for text in ['2026年8-9月', '2026年8至9月', '2026年8～9月', '2026年8月-9月']:
            self.assertEqual(2, len(months(text)), text)
        self.assertFalse(months('2026年9-8月'))
        right = self.material('2026年8-9月社保缴费凭证')
        wrong = self.material('2025年12月社保缴费凭证')
        scanned = self.scan(['2026年8-9月社保缴费证明'], [right, wrong])
        row = self.tax_rows(scanned)[0]
        self.assertEqual([right['id']], row['material_ids'])
        self.assertEqual('ready', row['status'])
        bad = engine.validate_mappings([dict(row, material_ids=[wrong['id']])],
            [right, wrong], scanned['headings'], self.categories)[0]
        self.assertIn('2025年12月社保缴费凭证：缴费所属月份与目标要求不符。', bad['issues'])

    def test_overlapping_social_versions_require_manual_selection(self):
        a = self.material('2026年8月社保缴费凭证')
        b = self.material('单位社保缴费证明（2025年9月至2026年8月）')
        row = self.tax_rows(self.scan(['社保缴费证明'], [a, b]))[0]
        self.assertFalse(row['material_ids'])

    def test_decimal_subsections_receive_their_own_evidence(self):
        doc = Document()
        for title in ['第五章 投标文件格式', '4.税收和社会保障资金缴纳凭证',
                      '4.1提供近6个月任意一个月的税收缴纳凭证，依法免税的，应提供证明文件。',
                      '4.2提供近6个月任意一个月的社会保障资金缴纳凭证，依法免缴的，应提供证明文件。',
                      '5.其他材料']:
            doc.add_paragraph(title)
        doc.save(self.path)
        mats = [self.material('2026年9月完税证明'), self.material('2026年8月社保缴费凭证')]
        scanned = engine.scan_document(self.path, self.categories, mats)
        rows = self.tax_rows(scanned)
        self.assertEqual(2, len(rows))
        self.assertTrue(all(not row['material_ids'] for row in rows))
        rows = [dict(row, action='fill', material_ids=[material['id']]) for row, material in zip(rows, mats)]
        original = self.path.read_bytes()
        output = self.root / 'result.docx'
        engine.export_document(self.path, output, rows, mats, categories=self.categories)
        result = Document(output)
        body = list(result.element.body)
        positions = [next(i for i, e in enumerate(body) if engine._text(e).startswith(prefix)) for prefix in ['4.1', '4.2', '5.']]
        for n, material in enumerate(mats):
            pictures = [b for e in body[positions[n]:positions[n + 1]] for b in e.iter(qn('a:blip'))]
            self.assertEqual(1, len(pictures))
            self.assertEqual(Path(material['files'][0]['path']).read_bytes(), result.part.related_parts[pictures[0].get(qn('r:embed'))].blob)
        self.assertEqual(original, self.path.read_bytes())

    def test_text_templates_keep_each_body_slot_and_reject_wrong_type(self):
        scanned = self.scan(['投标函', '授权书', '承诺函'], [])
        rows = [r for r in scanned['mappings'] if r['category_id'] == 'text_template']
        self.assertEqual(['投标函', '授权书', '承诺函'], [r['title'] for r in rows])
        material = dict(id='letter', category_id='text_template', name='投标函参考模板',
            subject=engine.EXPECTED_SUBJECT, mode='text', content='投标函测试正文', files=[], years=[], valid_until='长期')
        for row in rows:
            checked = engine.validate_mappings([dict(row, action='fill', material_ids=['letter'])],
                [material], scanned['headings'], self.categories)[0]
            self.assertEqual('ready' if row['title'] == '投标函' else 'blocked', checked['status'])

    def test_all_seven_categories_skip_directory_and_preserve_body_windows(self):
        doc = Document()
        doc.add_paragraph('目录')
        for category in engine.DEFAULT_CATEGORIES:
            doc.add_paragraph(category['name'])
        doc.add_heading('第五章 投标文件格式', 1)
        materials = []
        for category in engine.DEFAULT_CATEGORIES:
            doc.add_heading(category['name'], 2)
            doc.add_paragraph('本授权书声明：现委托李四办理投标事项。' if category['id'] == 'authorization'
                else '兹证明张三为本单位法定代表人。' if category['id'] == 'legal_identity' else '本栏目原有说明。')
            if category['id'] == 'authorization':
                doc.add_paragraph('附：被授权人身份证复印件')
            material = self.material(category['name'] + '测试图片', category['id'])
            material['years'] = [2023, 2024, 2025, 2026]
            materials.append(material)
        doc.add_heading('结束附件', 1)
        doc.save(self.path)
        original = self.path.read_bytes()
        scanned = engine.scan_document(self.path, self.categories, materials)
        self.assertEqual(7, len(scanned['mappings']))
        self.assertTrue(all(r['status'] == 'ready' for r in scanned['mappings']), scanned['mappings'])
        output = self.root / 'all.docx'
        engine.export_document(self.path, output, scanned['mappings'], materials, categories=self.categories)
        result = Document(output)
        nodes = list(result.element.body)
        start = next(i for i, e in enumerate(nodes) if engine._text(e) == '第五章 投标文件格式')
        self.assertFalse(any(list(e.iter(qn('a:blip'))) for e in nodes[:start]))
        positions = [next(i for i, e in enumerate(nodes) if i > start and engine._text(e) == c['name']) for c in engine.DEFAULT_CATEGORIES]
        positions.append(next(i for i, e in enumerate(nodes) if engine._text(e) == '结束附件'))
        for n, material in enumerate(materials):
            blips = [b for e in nodes[positions[n]:positions[n + 1]] for b in e.iter(qn('a:blip'))]
            self.assertEqual(1, len(blips))
            self.assertEqual(Path(material['files'][0]['path']).read_bytes(), result.part.related_parts[blips[0].get(qn('r:embed'))].blob)
        self.assertEqual(original, self.path.read_bytes())


if __name__ == '__main__':
    unittest.main()
