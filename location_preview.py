"""Render original template pages and locate export anchors without adding visible text."""
from __future__ import annotations

from collections import OrderedDict
import hashlib
import json
import os
from pathlib import Path
import shutil
import subprocess
import tempfile
import threading
import unicodedata

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.table import Table

import engine
from word_import import find_converter


# PDFium is not thread safe. Also serialize converter jobs to bound memory usage.
_RENDER_LOCK = threading.RLock()
VERSION = 'page-preview-v1'


def _bookmark(paragraph, index, bookmark_id, edge):
    name = f'BidPreview{index}{edge}'
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(bookmark_id))
    start.set(qn('w:name'), name)
    end = OxmlElement('w:bookmarkEnd')
    end.set(qn('w:id'), str(bookmark_id))
    if edge == 'S':
        offset = 1 if paragraph.find(qn('w:pPr')) is not None else 0
        paragraph.insert(offset, start)
        paragraph.insert(offset + 1, end)
    else:
        paragraph.append(start)
        paragraph.append(end)
    return name


def _render_pdf(source, directory):
    converter = find_converter()
    if not converter:
        raise ValueError('原页预览需要本机 LibreOffice，请安装后重试。')
    profile = directory / 'profile'
    (profile / 'user').mkdir(parents=True)
    (profile / 'user/registrymodifications.xcu').write_text(
        '<?xml version="1.0" encoding="UTF-8"?>'
        '<oor:items xmlns:oor="http://openoffice.org/2001/registry">'
        '<item oor:path="/org.openoffice.Office.Common/Security/Scripting">'
        '<prop oor:name="MacroSecurityLevel" oor:op="fuse"><value>3</value></prop>'
        '</item></oor:items>', encoding='utf-8')
    options = {name: {'type': 'boolean', 'value': 'true'} for name in
               ('ExportBookmarks', 'ExportBookmarksToPDFDestination')}
    env = os.environ.copy()
    # Use the Mac font configuration, including installed Chinese fonts.
    if not env.get('FONTCONFIG_FILE') and Path('/opt/homebrew/etc/fonts/fonts.conf').is_file():
        env['FONTCONFIG_FILE'] = '/opt/homebrew/etc/fonts/fonts.conf'
    try:
        result = subprocess.run([
            converter, f'-env:UserInstallation={profile.as_uri()}',
            '--headless', '--nologo', '--nodefault', '--norestore',
            '--convert-to', 'pdf:writer_pdf_Export:' + json.dumps(options),
            '--outdir', str(directory), str(source),
        ], stdin=subprocess.DEVNULL, stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL, timeout=90, check=False, env=env)
    except subprocess.TimeoutExpired:
        raise ValueError('标书页面渲染超时，请稍后重试。') from None
    except OSError:
        raise ValueError('本机页面渲染工具无法启动，请检查 LibreOffice。') from None
    finally:
        shutil.rmtree(profile, ignore_errors=True)
    output = directory / (source.stem + '.pdf')
    if result.returncode or not output.is_file() or not output.stat().st_size:
        raise ValueError('暂时无法渲染这份 Word 的原页，请检查文档是否可正常打开。')
    if output.stat().st_size > 160 * 1024 * 1024:
        raise ValueError('预览文件过大，请精简模板后重试。')
    return output


class LocationPreview:
    def __init__(self):
        # Runtime cache only: keep at most two rendered templates; remove on close.
        self._temporary = tempfile.TemporaryDirectory(prefix='bid-page-preview-')
        self.root = Path(self._temporary.name)
        self.cache = OrderedDict()
        # Inline previews ask once per review card. Hash each template version
        # and recognize its locations once, not on every card request.
        self._recognized = OrderedDict()
        self._digests = OrderedDict()
        self._recognized_lock = threading.Lock()

    def close(self):
        with _RENDER_LOCK:
            self.cache.clear()
            self._temporary.cleanup()
        with self._recognized_lock:
            self._recognized.clear()
            self._digests.clear()

    def _verify(self, source, template_hash):
        stat = Path(source).stat()
        identity = (str(source), stat.st_size, stat.st_mtime_ns)
        with self._recognized_lock:
            digest = self._digests.get(identity)
            if digest is None:
                digest = hashlib.sha256(Path(source).read_bytes()).hexdigest()
                self._digests[identity] = digest
                while len(self._digests) > 16:
                    self._digests.popitem(last=False)
        if digest != template_hash:
            raise ValueError('模板已变化，请重新上传或扫描后再预览。')

    def _recognize(self, source, template_hash, categories):
        """Shared read-only document and public records for one template and category set."""
        rules = hashlib.sha256(json.dumps(categories, sort_keys=True, ensure_ascii=False, default=str).encode()).hexdigest()
        key = (template_hash, rules)
        with self._recognized_lock:
            if key not in self._recognized:
                document = engine._read_template(source)
                records, _ = engine._inspect(document, categories)
                self._recognized[key] = (document, {r['id']: r for r in records if r['_public']})
                while len(self._recognized) > 2:
                    self._recognized.popitem(last=False)
            self._recognized.move_to_end(key)
            return self._recognized[key]

    @staticmethod
    def key(template_hash):
        return hashlib.sha256((VERSION + template_hash).encode()).hexdigest()

    def _build(self, source, template_hash):
        from pypdf import PdfReader
        key = self.key(template_hash)
        if key in self.cache:
            self.cache.move_to_end(key)
            return self.cache[key]
        directory = self.root / key
        directory.mkdir()
        try:
            raw = Path(source).read_bytes()
            if hashlib.sha256(raw).hexdigest() != template_hash:
                raise ValueError('模板已变化，请重新上传或扫描后再预览。')
            document = engine._read_template(source)
            paragraphs = list(document.element.body.iter(qn('w:p')))
            if len(paragraphs) > 15000:
                raise ValueError('模板段落过多，暂不支持原页预览。')
            used_ids = [int(e.get(qn('w:id'))) for e in document.element.iter(qn('w:bookmarkStart'))
                        if str(e.get(qn('w:id'), '')).isdigit()]
            offset = max(used_ids, default=0) + 1
            names = {}
            paths = [engine._xml_path(p, document.element.body) for p in paragraphs]
            for index, (paragraph, path) in enumerate(zip(paragraphs, paths)):
                # LibreOffice increases empty-cell height when a bookmark makes an
                # empty paragraph nonempty. Resolve those cells from the original
                # rendered table borders instead, without touching their XML.
                if not engine._text(paragraph).strip() and not engine._drawing(paragraph):
                    continue
                names[path] = {edge: _bookmark(paragraph, index, offset + index * 2 + j, edge)
                               for j, edge in enumerate(('S', 'E'))}
            marked = directory / 'template.docx'
            document.save(marked)
            pdf = _render_pdf(marked, directory)
            marked.unlink()
            reader = PdfReader(pdf)
            if not 0 < len(reader.pages) <= 600:
                raise ValueError('模板超过 600 页，暂不支持原页预览。')
            destinations = {}
            for name, destination in reader.named_destinations.items():
                name = name.lstrip('/')
                if not name.startswith('BidPreview'):
                    continue
                page_index = reader.get_destination_page_number(destination)
                if page_index is None or not 0 <= page_index < len(reader.pages):
                    continue
                page = reader.pages[page_index]
                if page.rotation:
                    continue  # Do not invent coordinates for unsupported rotation.
                box = page.mediabox
                if destination.get('/Top') is None or destination.get('/Left') is None:
                    continue
                destinations[name] = {
                    'page': page_index + 1,
                    'x': max(0, min(1, (float(destination['/Left']) - float(box.left)) / float(box.width))),
                    'y': max(0, min(1, (float(box.top) - float(destination['/Top'])) / float(box.height))),
                }
            entry = {'pdf': pdf, 'names': names, 'destinations': destinations, 'tables': {},
                     'page_count': len(reader.pages), 'preview_id': key}
            self.cache[key] = entry
            while len(self.cache) > 2:
                _, expired = self.cache.popitem(last=False)
                shutil.rmtree(expired['pdf'].parent)
            return entry
        except Exception:
            shutil.rmtree(directory, ignore_errors=True)
            raise

    def describe(self, snapshot):
        source = snapshot['template']
        self._verify(source, snapshot['template_hash'])
        if not snapshot['locator_id']:
            with _RENDER_LOCK:
                entry = self._build(source, snapshot['template_hash'])
                return {'preview_id': entry['preview_id'], 'page_count': entry['page_count'],
                        'initial_page': 1, 'target_pages': [], 'markers': [],
                        'title': '标书原页',
                        'description': '尚未选择插入位置。可先翻页查看原文，再从上方选择位置查看标记。',
                        'template_name': snapshot['template_name']}
        document, records = self._recognize(source, snapshot['template_hash'], snapshot['categories'])
        record = records.get(snapshot['locator_id'])
        if record is None:
            raise ValueError('所选插入位置已失效，请重新扫描后再选择。')
        if record['kind'] == 'common_appendix':
            with _RENDER_LOCK:
                entry = self._build(source, snapshot['template_hash'])
                return {'preview_id': entry['preview_id'], 'page_count': entry['page_count'],
                        'initial_page': entry['page_count'], 'target_pages': [], 'markers': [],
                        'title': record['title'],
                        'description': '原稿没有独立填报位置。导出时会在原文末尾另起一页新增通用材料区；此处展示原稿末页，新增材料不会插入采购要求或评分规则。',
                        'template_name': snapshot['template_name']}
        targets = []
        selected = snapshot['materials']
        use_anchor = not selected or snapshot['action'] != 'fill'
        if snapshot['action'] == 'fill':
            for material in selected:
                cells = (engine._vacant_cells(document, record, len(material.get('files', [])))
                         if material.get('mode') == 'side_by_side' and material.get('files') else None)
                if cells is not None:
                    for index, cell in enumerate(cells[:len(material['files'])]):
                        targets.append((cell.paragraphs[0]._p, 'S', 'cell', f'第 {index + 1} 张图片填入此单元格'))
                else:
                    use_anchor = True
        after_table = False
        table_anchor = None
        if use_anchor:
            anchor = record['_element'] if record['kind'] == 'placeholder' else record['_anchor']
            if anchor.tag == qn('w:tbl'):
                after_table = True
                table_anchor = anchor
                paragraphs = []
            else:
                paragraphs = [anchor]
            for paragraph in paragraphs:
                targets.append((paragraph, 'E', 'after', '材料从此处之后插入'))
        if snapshot['action'] == 'missing':
            description = '当前保留为缺项，导出初稿时会在此位置之后添加待补标记。'
        elif snapshot['action'] == 'skip':
            description = '当前选择跳过，不填入材料；导出初稿时会在此位置之后添加跳过说明。'
        elif record['kind'] == 'placeholder':
            description = '导出时移除本占位符，并在其所在段落之后插入材料。'
        elif any(t[2] == 'cell' for t in targets) and not use_anchor:
            description = '图片将填入本节现有的空白表格单元格。'
        elif after_table:
            description = '材料将在此表格结束后追加，原表格内容保留。'
        else:
            description = '材料将在本节原有内容结束后追加；原文保留。'
        if snapshot['action'] == 'fill' and any(m.get('mode') == 'page' for m in selected) and not record['_inside_table']:
            description += '整页图片材料会从此处之后另起一页。'
        with _RENDER_LOCK:
            entry = self._build(source, snapshot['template_hash'])
            markers, anchors = [], []
            for paragraph, edge, kind, label in targets:
                path = engine._xml_path(paragraph, document.element.body)
                name = entry['names'].get(path, {}).get(edge)
                location = entry['destinations'].get(name)
                if not location and kind == 'cell':
                    location = self._cell_location(document, paragraph, entry)
                if not location:
                    raise ValueError('本位置未能可靠定位到原页，请重新选择位置或检查模板。')
                marker = {**location, 'kind': kind, 'label': label}
                if kind == 'after':
                    marker['_paragraph_text'] = engine._text(paragraph)
                    start_name = entry['names'].get(path, {}).get('S')
                    marker['_start'] = entry['destinations'].get(start_name, location)
                (anchors if kind == 'after' else markers).append(marker)
            if table_anchor is not None:
                regions = self._table_regions(document, table_anchor, entry)
                last = regions[-1]
                markers.append({'page': last['page'], 'x': last['right'],
                                'y': min(.99, last['bottom'] + .003),
                                'kind': 'table', 'label': '本表格结束后插入'})
            if anchors:
                marker = max(anchors, key=lambda m: (m['page'], m['y'], m['x']))
                if after_table:
                    marker['label'] = '本表格结束后插入'
                markers.append(marker)
            if not markers:
                raise ValueError('该位置没有可显示的插入点。')
            self._place_below_text(entry, markers, inside_cell=record['_inside_table'])
            for marker in markers:
                marker.pop('_paragraph_text', None)
                marker.pop('_start', None)
            if record['_inside_table']:
                for marker in markers:
                    if marker['kind'] == 'after':
                        marker['kind'] = 'cell'
                        marker['label'] = '此单元格内，段落后插入'
            if snapshot['action'] != 'fill':
                note = '待补标记' if snapshot['action'] == 'missing' else '跳过说明'
                for marker in markers:
                    marker['label'] = ('此单元格内添加' if marker['kind'] == 'cell' else
                                       '本表格结束后添加' if marker['kind'] == 'table' else
                                       '此处之后添加') + note
            pages = sorted({m['page'] for m in markers})
            return {'preview_id': entry['preview_id'], 'page_count': entry['page_count'],
                    'initial_page': pages[0], 'target_pages': pages, 'markers': markers,
                    'title': record['title'], 'description': description,
                    'template_name': snapshot['template_name']}

    @staticmethod
    def _page_tables(entry, number):
        """Connected horizontal/vertical PDF border strokes identify real table boxes."""
        import pypdfium2 as pdfium
        if number in entry['tables']:
            return entry['tables'][number]
        pdf = pdfium.PdfDocument(entry['pdf'])
        page = pdf[number - 1]
        try:
            width, height = page.get_size()
            groups = []
            for obj in page.get_objects():
                if obj.type != 2:  # FPDF_PAGEOBJ_PATH
                    continue
                left, bottom, right, top = obj.get_bounds()
                horizontal = right - left > 10 and top - bottom < 3
                vertical = top - bottom > 5 and right - left < 3
                if not (horizontal or vertical):
                    continue
                box = [left, bottom, right, top]
                hits = [g for g in groups if not (
                    g['box'][2] + 1 < left or right + 1 < g['box'][0] or
                    g['box'][3] + 1 < bottom or top + 1 < g['box'][1])]
                lines = [(box, horizontal)]
                for group in hits:
                    groups.remove(group)
                    lines.extend(group['lines'])
                    a = group['box']
                    box = [min(box[0], a[0]), min(box[1], a[1]), max(box[2], a[2]), max(box[3], a[3])]
                groups.append({'box': box, 'lines': lines})
            result = []
            for group in groups:
                left, bottom, right, top = group['box']
                horizontal = [b for b, h in group['lines'] if h]
                vertical = [b for b, h in group['lines'] if not h]
                if right - left < 30 or top - bottom < 5 or len(horizontal) < 2 or len(vertical) < 2:
                    continue
                def distinct(values):
                    result = []
                    for value in sorted(values):
                        if not result or value - result[-1] > 2:
                            result.append(value)
                    return result
                result.append({'page': number, 'left': left / width, 'right': right / width,
                               'top': (height - top) / height, 'bottom': (height - bottom) / height,
                               'xs': [v / width for v in distinct([(b[0] + b[2]) / 2 for b in vertical])],
                               'ys': [(height - v) / height for v in reversed(distinct([(b[1] + b[3]) / 2 for b in horizontal]))]})
            entry['tables'][number] = sorted(result, key=lambda r: (r['top'], r['left']))
            return entry['tables'][number]
        finally:
            page.close()
            pdf.close()

    def _table_regions(self, document, table, entry):
        path = engine._xml_path(table, document.element.body)
        before = [p for p in entry['names'] if p < path]
        after = [p for p in entry['names'] if p > path and p[:len(path)] != path]
        start = entry['destinations'].get(entry['names'][max(before)]['E']) if before else None
        end = entry['destinations'].get(entry['names'][min(after)]['S']) if after else None
        first, last = start['page'] if start else 1, end['page'] if end else entry['page_count']
        if last - first > 50:
            raise ValueError('表格范围过大，无法可靠标记插入位置。')
        regions = []
        for page in range(first, last + 1):
            for box in self._page_tables(entry, page):
                if start and page == start['page'] and box['top'] < start['y'] - .005:
                    continue
                if end and page == end['page'] and box['bottom'] > end['y'] + .005:
                    continue
                regions.append(box)
        if not regions:
            raise ValueError('此表格没有可可靠识别的边界，暂不能在原页中标记；请在 Word 中核对。')
        if len({r['page'] for r in regions}) != len(regions):
            raise ValueError('此处有多个相邻表格，无法可靠区分插入区域；请在 Word 中核对。')
        return regions

    def _cell_location(self, document, paragraph, entry):
        cell = next((p for p in paragraph.iterancestors() if p.tag == qn('w:tc')), None)
        table_node = next((p for p in paragraph.iterancestors() if p.tag == qn('w:tbl')), None)
        if cell is None or table_node is None:
            raise ValueError('无法可靠定位这个空白插入位置。')
        regions = self._table_regions(document, table_node, entry)
        table = Table(table_node, document._body)
        if len(regions) != 1:
            raise ValueError('此空白表格跨页，暂不能可靠标记单元格；请在 Word 中核对。')
        box = regions[0]
        if len(box['ys']) != len(table.rows) + 1 or len(box['xs']) != len(table.columns) + 1:
            raise ValueError('此空白表格的单元格边界不完整，暂不能可靠标记。')
        for row_index, row in enumerate(table.rows):
            columns = [i for i, candidate in enumerate(row.cells) if candidate._tc is cell]
            if columns:
                return {'page': box['page'],
                        'x': (box['xs'][min(columns)] + box['xs'][max(columns) + 1]) / 2,
                        'y': (box['ys'][row_index] + box['ys'][row_index + 1]) / 2}
        raise ValueError('无法可靠定位这个空白单元格。')

    @staticmethod
    def _place_below_text(entry, markers, inside_cell=False):
        """Bookmarks sit at the line top. Draw an 'after' marker below its last glyph."""
        import pypdfium2 as pdfium
        pdf = pdfium.PdfDocument(entry['pdf'])
        try:
            for number in {m['page'] for m in markers if m['kind'] == 'after'}:
                page = pdf[number - 1]
                text = page.get_textpage()
                try:
                    width, height = page.get_size()
                    chars = [(i, text.get_text_range(i, 1)) for i in range(text.count_chars())]
                    boxes = [text.get_charbox(i) for i, char in chars if char.strip()]
                    def normalized(value):
                        return ''.join(c for c in unicodedata.normalize('NFKC', value)
                                       if not c.isspace() and c not in '\u00ad\u200b\ufeff')
                    stream, indices = [], []
                    for i, char in chars:
                        value = normalized(char)
                        stream.append(value)
                        indices.extend([i] * len(value))
                    rendered = ''.join(stream)
                    for marker in (m for m in markers if m['page'] == number and m['kind'] == 'after'):
                        x, top = marker['x'] * width, (1 - marker['y']) * height
                        nearby = [b for b in boxes if abs(b[2] - x) <= 30 and -4 <= top - b[3] <= 25]
                        target = normalized(marker.get('_paragraph_text', ''))
                        matches, offset = [], 0
                        while target and (offset := rendered.find(target, offset)) >= 0:
                            matches.append(offset)
                            offset += len(target)
                        box = None
                        if matches:
                            # A Word bookmark may point to the first line even at
                            # the paragraph end. Locate its actual final glyph,
                            # retaining the start bookmark to disambiguate repeats.
                            start = marker.get('_start', marker)
                            start_x, start_y = start['x'] * width, (1 - start['y']) * height
                            def distance(offset):
                                first = text.get_charbox(indices[offset])
                                return abs(first[0] - start_x) + abs(first[3] - start_y)
                            offset = min(matches, key=distance)
                            box = text.get_charbox(indices[offset + len(target) - 1])
                        elif nearby:
                            box = min(nearby, key=lambda b: abs(b[2] - x) + abs(top - b[3]))
                        if box:
                            if inside_cell:
                                # Keep the point inside the source paragraph's
                                # line, including rows with no space below it.
                                marker['x'] = (box[2] - 2) / width
                                marker['y'] = (height - (box[1] + box[3]) / 2) / height
                            else:
                                marker['y'] = min(.99, max(0, (height - box[1] + 2) / height))
                finally:
                    text.close()
                    page.close()
        finally:
            pdf.close()

    def page(self, key, page_number):
        import pypdfium2 as pdfium
        with _RENDER_LOCK:
            entry = self.cache.get(key)
            if not entry:
                raise ValueError('预览已过期，请关闭后重新打开插入位置预览。')
            if not 1 <= page_number <= entry['page_count']:
                raise ValueError('预览页码超出范围。')
            self.cache.move_to_end(key)
            output = entry['pdf'].parent / f'page-{page_number}.png'
            if not output.is_file():
                pdf = pdfium.PdfDocument(entry['pdf'])
                try:
                    page = pdf[page_number - 1]
                    try:
                        width, height = page.get_size()
                        scale = min(1.7, 1700 / width, (5_000_000 / (width * height)) ** .5)
                        bitmap = page.render(scale=scale)
                        try:
                            picture = bitmap.to_pil()
                            picture.save(output, format='PNG')
                        finally:
                            bitmap.close()
                    finally:
                        page.close()
                finally:
                    pdf.close()
            return output.read_bytes()
