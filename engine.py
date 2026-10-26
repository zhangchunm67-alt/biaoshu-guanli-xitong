"""Local, deterministic DOCX material assembly; no network or database access.

Locators name an XML path in the *original* template and a digest of that
paragraph.  They are resolved before edits, so repeated titles and inserted
pages cannot move a target.  This module never changes source documents/images.
"""
from __future__ import annotations

import copy
import hashlib
import io
import os
import re
import tempfile
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Pt, RGBColor
from docx.table import Table
from docx.text.paragraph import Paragraph
from PIL import Image, ImageOps


EXPECTED_SUBJECT = "北京腾宝科技有限公司"
DEFAULT_CATEGORIES = [
    {"id": "legal_identity", "name": "法定代表人（负责人）身份证明", "aliases": ["法定代表人身份证明", "法人身份证明"]},
    {"id": "authorization", "name": "法定代表人（负责人）授权委托书", "aliases": ["法定代表人授权委托书", "授权委托书"]},
    {"id": "business_license", "name": "营业执照", "aliases": ["企业法人营业执照", "营业执照副本"]},
    {"id": "audit_report", "name": "2023-2024年财务审计报告", "aliases": ["财务审计报告", "审计报告"]},
    {"id": "tax_social", "name": "税收和社会保障资金缴纳凭证", "aliases": ["税收和社保缴纳凭证", "税收和社会保障资金缴纳证明"]},
    {"id": "company_profile", "name": "公司介绍", "aliases": ["公司简介", "企业介绍", "企业简介"]},
    {"id": "iso_cert", "name": "ISO相关资质证书", "aliases": ["ISO认证证书", "ISO资质证书"]},
]
for _category in DEFAULT_CATEGORIES:
    _category.update(layer="common", enabled=True)

_TOKEN = re.compile(r"\{\{\s*([^{}]+?)\s*\}\}")
_YEAR = re.compile(r"(?<!\d)(20\d{2})(?!\d)")
_EXTENSIONS = {".jpg", ".jpeg", ".png", ".bmp", ".tif", ".tiff", ".webp", ".gif"}
SPECIAL_LOCATOR_VERSION = 9
COMMON_LOCATOR_VERSION = 11


def _common_title(text):
    value = _normal(text, remove_years=True)
    value = re.sub(r'^[•●◆◇★☆*·\-—]+', '', value)
    value = re.sub(r'[（(【\[].*?[）)】\]]', '', value)
    return value.strip('：:；;，,、。')


def _common_variant(title, category_id):
    target = _common_title(title)
    if category_id == 'tax_social':
        target = re.sub(r'\d{1,2}月?(?:至|到|[-—–~～])\d{1,2}月|\d{1,2}月', '', target)
    patterns = {
        'business_license': r'(?:(?:投标人|投标方|供应商|企业法人|企业|公司)的?)?营业执照(?:正副本|正本|副本)?(?:彩色)?(?:复印件|扫描件|影印件)?',
        'company_profile': r'(?:(?:投标人|投标方|供应商)的?)?(?:公司|企业)(?:介绍|简介|概况)',
        'legal_identity': r'(?:法定代表人|法人代表|法人|负责人)(?:身份证明书?|证明书|身份证(?:正反面)?(?:复印件|扫描件)?)',
        'authorization': r'(?:法定代表人|法人代表|法人|负责人)?授权(?:委托)?书(?:格式)?',
        'tax_social': r'(?:依法)?(?:缴纳)?(?:(?:税收|纳税|税款)(?:和|及|与)?(?:社会保障资金|社会保险|社保)?|社会保障资金|社会保险|社保)(?:缴纳|缴费)?(?:证明(?:材料)?|凭证|记录)|完税(?:证明|凭证)',
    }
    return bool(category_id in patterns and re.fullmatch(patterns[category_id], target))


def _special_score(title, category):
    """Match complete evidence-section titles, never a keyword in prose."""
    target = _normal(title, remove_years=True)
    if category['id'] == 'text_template' and re.fullmatch(
            r'(?:[（(]?(?:附件|附表|格式)[一二三四五六七八九十百零〇\d]+[）)]?[：:、.．\-]*)?'
            r'(?:(?:法定代表人)?授权(?:委托)?书(?:格式)?|投标函|承诺[函书])', target):
        return 130, '文书类型与正文标题'
    if target in {_normal(n, True) for n in [category['name'], *category.get('aliases', [])]}:
        return 130, '栏目名称或已配置别名'
    if len(target) > 50 or re.search(r'[。；;，,]', target):
        return 0, ''
    family = re.fullmatch(r'[IFBW]0([1-4])', category.get('code', category['id']))
    if not family:
        return 0, ''
    industries = {'insurance': '保险', 'fund': '基金', 'bank': '银行', 'wealth': '理财'}
    industry = industries.get(category.get('industry_id'))
    for name in industries.values():
        if target.startswith(name):
            if industry and name != industry:
                return 0, ''
            target = target[len(name):]
            break
    target = re.sub(r'^[（(]?(?:附件|附表|格式)[一二三四五六七八九十\d]+[）)、:：.．]*', '', target)
    target = re.sub(r'[（(](?:格式|参考格式|样式)[）)]$', '', target)
    if family[1] == '3' and target in {'资格证明文件', '资格证明材料', '资质证明文件', '资质证明材料'}:
        return 0, ''  # Broad parents also contain common licenses and financial evidence.
    rules = {
        '1': [
            (110, r'(?:同类|类似)?(?:项目)?(?:业绩|案例|合同)(?:及履约)?证明(?:材料|文件|附件)'),
            (100, r'(?:同类|类似|企业|公司|供应商|投标人)?(?:项目)?业绩(?:情况)?(?:汇总表|一览表|表)'),
            (85, r'(?:同类|类似)(?:项目)?(?:实施|服务)?(?:经验|业绩)(?:证明)?'),
            (80, r'(?:项目|服务|成功|典型)?案例(?:展示|介绍|汇总|证明材料)'),
            (85, r'(?:项目|历史|企业|公司)?业绩(?:情况|证明|材料)?'),
        ],
        '2': [
            (110, r'(?:拟派|项目|服务|主要|关键)?(?:团队成员|项目人员|人员|成员)(?:经历|经验|资质|资格|专业)(?:与专业)?证明(?:材料|文件)?'),
            (100, r'(?:拟派|拟投入)?(?:项目|服务)?(?:团队成员|项目人员|人员|成员|项目负责人|项目经理)(?:个人)?简历(?:表)?'),
            (80, r'(?:拟派|拟投入)?(?:项目|服务)?(?:团队成员|项目人员|人员|成员)(?:情况)?(?:汇总表|一览表|配置表|名单)'),
            (80, r'(?:拟派|拟投入)?(?:项目|服务)?(?:团队|人员)(?:介绍|配置|组成|情况|经历与专业证明)'),
        ],
        '3': [
            (110, r'(?:行业)?(?:专项|专业|业务|企业)?(?:资质|资格)(?:证书|证明|证明文件|证明材料)'),
            (110, r'(?:广播电视节目制作经营|增值电信业务经营|互联网新闻信息服务|广告经营)许可证'),
        ],
        '4': [
            (110, r'(?:平台|媒体|媒介|渠道|厂商|制造商|品牌|业务|合作)(?:资源)?(?:代理|合作)?(?:及|与)?授权(?:书|证明|证明材料|文件|资料)'),
            (100, r'(?:平台|媒体|媒介|渠道|业务)?(?:资源|合作资源)(?:与合作|与|及)?(?:授权|证明|优势|介绍|情况|清单|证明材料)'),
            (100, r'(?:平台|媒体|渠道)(?:代理|合作)(?:资质|证明|协议|授权资料)'),
        ],
    }
    for score, pattern in rules[family[1]]:
        if re.fullmatch(pattern, target):
            return score, '专用材料标题识别'
    return 0, ''


def _special_candidates(records, category):
    candidates = []
    for record in records:
        if category['id'] not in record['category_ids']:
            continue
        score, reason = record['_special_scores'][category['id']]
        if record['kind'] == 'placeholder':
            score += 30
        candidates.append({'locator_id': record['id'], 'title': record['title'],
                           'score': score, 'reason': reason})
    candidates.sort(key=lambda c: -c['score'])
    if not candidates:
        return '', {'status': 'not_found', 'reason': '模板未识别到明确的专用材料章节', 'candidates': []}
    if len(candidates) > 1 and candidates[0]['score'] - candidates[1]['score'] < 10:
        return '', {'status': 'ambiguous', 'reason': '识别到多个相近位置，请核对候选章节', 'candidates': candidates[:8]}
    chosen = next(r for r in records if r['id'] == candidates[0]['locator_id'])
    match = dict(chosen.get('location_match', {}), status='matched',
                 reason=candidates[0]['reason'], candidates=candidates[:8])
    return candidates[0]['locator_id'], match


def _unique(items):
    return list(dict.fromkeys(items))


def _years(text):
    values = {int(y) for y in _YEAR.findall(str(text or ""))}
    for a, b in re.findall(r"(20\d{2})\s*[-—–~至到]\s*(20\d{2})", str(text or "")):
        if int(a) <= int(b) and int(b) - int(a) <= 30:
            values.update(range(int(a), int(b) + 1))
    return sorted(values)


def _normal(text, remove_years=False):
    value = str(text or "").strip()
    if not re.match(r"^20\d{2}(?:\D|$)", value):
        value = re.sub(r"^\d+(?:[.．]\d+)*(?:[.．、]|\s+)\s*", "", value)
    value = re.sub(r"^(?:第[一二三四五六七八九十百\d]+(?:章|节|部分|篇)|[一二三四五六七八九十百]+[、.．]|[（(][一二三四五六七八九十\d]+[）)])\s*", "", value)
    value = re.sub(r"\s+", "", value).rstrip("：:")
    value = value.replace("（负责人）", "").replace("(负责人)", "")
    if remove_years:
        value = re.sub(r"20\d{2}(?:[-—–~至到]20\d{2})?年?度?", "", value)
    return value.lower()


def _category_ids(title, categories):
    target = _normal(title, remove_years=True)
    return [c["id"] for c in categories if target in {
        _normal(n, remove_years=True) for n in [c["name"], *c.get("aliases", [])]
    } or (c.get('layer', 'common') == 'common' and _common_variant(title, c['id']))]


def _text(element):
    return "".join(t.text or "" for t in element.iter(qn("w:t")))


def _drawing(element):
    return next(element.iter(qn("w:drawing")), None) is not None or next(element.iter(qn("w:pict")), None) is not None


def _style_level(paragraph):
    ppr = paragraph._p.find(qn("w:pPr"))
    if ppr is not None:
        node = ppr.find(qn("w:outlineLvl"))
        if node is not None and str(node.get(qn("w:val"), "9")).isdigit():
            value = int(node.get(qn("w:val"), "9"))
            if value < 9:
                return value + 1
    style = paragraph.style
    seen = set()
    while style is not None and style.style_id not in seen:
        seen.add(style.style_id)
        match = re.match(r"(?:heading|标题)\s*([1-9])", style.name or "", re.I)
        if match:
            return int(match.group(1))
        ppr = style.element.find(qn("w:pPr"))
        node = ppr.find(qn("w:outlineLvl")) if ppr is not None else None
        if node is not None and node.get(qn("w:val")) in [str(i) for i in range(9)]:
            return int(node.get(qn("w:val"))) + 1
        style = style.base_style
    return 0


def _number_level(text):
    if len(text) > 90 or re.search(r"[。；;，,]$", text):
        return 0
    if re.match(r"^第[一二三四五六七八九十百\d]+[章节部分]", text):
        return 1
    if re.match(r"^[一二三四五六七八九十百]+[、.．]", text):
        return 1
    if re.match(r"^[（(][一二三四五六七八九十]+[）)]", text):
        return 2
    if not re.match(r"^20\d{2}(?:\D|$)", text):
        match = re.match(r"^(\d+(?:[.．]\d+)*)(?:[.．、]|\s+)\S", text)
        if match:
            return min(9, len(re.split(r"[.．]", match.group(1))))
    return 0


def _is_toc(element, paragraph):
    name = (getattr(paragraph.style, 'name', '') or "").lower()
    if name.startswith(("toc", "目录")) or re.search(r"\t\s*\d+\s*$", paragraph.text):
        return True
    if re.search(r"[.．·…]{2,}\s*\d+\s*$", paragraph.text):
        return True
    for node in element.iter():
        if node.tag == qn("w:hyperlink") and node.get(qn("w:anchor"), "").startswith("_Toc"):
            return True
        if node.tag == qn("w:instrText") and re.search(r"\b(?:TOC|PAGEREF)\b", node.text or ""):
            return True
    for ancestor in element.iterancestors():
        if ancestor.tag == qn("w:sdt"):
            for node in ancestor.iter(qn("w:docPartGallery")):
                if node.get(qn("w:val"), "").lower() in {"table of contents", "目录"}:
                    return True
    return False


def _xml_path(element, body):
    path = []
    while element is not body:
        parent = element.getparent()
        path.append(parent.index(element))
        element = parent
    return tuple(reversed(path))


def _digest(element):
    return hashlib.sha256(element.xml.encode("utf-8")).hexdigest()[:16]


def _page_size(document, top_index):
    section_index = 0
    for index, element in enumerate(document.element.body):
        if index >= top_index:
            break
        if element.tag == qn("w:p") and element.find(".//" + qn("w:sectPr")) is not None:
            section_index += 1
    section = document.sections[min(section_index, len(document.sections) - 1)]
    return int(section.page_width - section.left_margin - section.right_margin), int(section.page_height - section.top_margin - section.bottom_margin)


def _inspect(document, categories, toc_elements=None):
    """Private records retain XML objects; public headings are JSON-only."""
    import recognition
    from document_toc import excluded_paragraphs
    body = document.element.body
    if toc_elements is None:
        toc_elements = excluded_paragraphs(document)
    records = []
    common = [c for c in categories if c.get('layer', 'common') != 'special']
    special = [c for c in categories if c.get('layer') == 'special']
    skipped_toc = 0
    for element in body.iter(qn("w:p")):
        paragraph = Paragraph(element, document._body)
        title = _text(element).strip()
        if not title:
            continue
        if element in toc_elements:
            skipped_toc += 1
            continue
        path = _xml_path(element, body)
        inside_table = any(a.tag == qn("w:tc") for a in element.iterancestors())
        style_level = _style_level(paragraph)
        level = style_level
        if not level and not inside_table:
            level = _number_level(title)
        cats = _category_ids(title, common) if len(title) <= 160 and not inside_table else []
        special_title = not inside_table and (any(_special_score(title, c)[0] for c in special) or bool(recognition.title_topics(title)))
        locator = "p:" + ".".join(map(str, path)) + ":" + _digest(element)
        base = {"id": locator, "title": title, "level": level or 9,
                "preview": "", "kind": "heading" if level else "paragraph",
                "category_ids": cats, "existing_content": False,
                "legal_body_present": False, "authorization_body_present": False,
                "_element": element, "_path": path, "_inside_table": inside_table,
                "_special_scores": {},
                "_raw_categories": list(cats),
                "_structural": recognition.structural(title, bool(style_level)),
                "_public": bool((level and not inside_table) or cats or special_title),
                "_anchor": element, "_own": []}
        records.append(base)
        for occurrence, token in enumerate(_TOKEN.finditer(title)):
            record = dict(base, id=locator.replace("p:", "ph:", 1) + f":{occurrence}",
                          title=token.group(1).strip(), kind="placeholder", level=9,
                          category_ids=_category_ids(token.group(1), common), _special_scores={},
                          _token=token.group(0), _occurrence=occurrence, _public=True)
            records.append(record)
    recognition.resolve(document, records, categories, toc_elements)
    # A token and a role-specific attachment retain their enclosing original
    # authorization text for name checks, without moving their write-back anchor.
    for record in records:
        if record['kind'] == 'placeholder' or record.get('_attachment'):
            parents = [r for r in records if r is not record and r['kind'] not in {'placeholder','attachment','common_appendix'}
                       and len(r['_path']) == 1 and r['_path'][0] <= record['_path'][0]
                       and record['_path'][0] - r['_path'][0] < 60
                       and (r.get('authorization_body_present') or r.get('legal_body_present'))]
            if parents:
                owner = max(parents, key=lambda r: r['_path'])
                for key in ('legal_body_present','authorization_body_present'):
                    record[key] = owner[key]
                if 'authorization' in record['category_ids']:
                    record['_authorization_context'] = '\n'.join(_text(e) for e in owner['_own'])
    import authorization
    for record in records:
        if 'authorization' in record['category_ids']:
            authorization.describe(record)
    return records, skipped_toc


def _public(record):
    return {k: v for k, v in record.items() if not k.startswith("_")}


def _read_template(path):
    try:
        return Document(path if hasattr(path, 'read') else str(path))
    except Exception as exc:
        raise ValueError("无法读取 Word 模板，请使用有效的 .docx 文件。") from exc


def _document_warnings(document, records, skipped_toc):
    warnings = ["仅处理已启用的通用栏目；原模板其他正文、姓名、日期和项目内容均保留，须由经办人终审。",
                "导出仅完成材料装配；未核验签章、招标条款、材料真实性或实际页码，不代表可直接投标。"]
    full_text = "\n".join(_text(e) for e in document.element.body)
    if re.search(r"20\d{2}\s*年\s*\d{1,2}\s*月|项目[/／]标包|项目名称|标包\d|投标人名称|兹证明", full_text):
        warnings.append("模板含已填写的日期、项目或人员/主体正文，可能沿用旧项目信息；本次不会自动改写。")
    if any(r["existing_content"] and r.get("category_ids") for r in records):
        warnings.append("部分匹配栏目已有正文、表格或图片；系统保留原内容并追加，需核对重复、过期及旧版本材料。")
    if skipped_toc:
        warnings.append(f"已排除 {skipped_toc} 处目录/页码引用，填充位置不依赖页码。")
    if len(document.sections) > 1:
        warnings.append("模板包含多个分节，图片按所在节纸张尺寸适配；导出后仍需在 Word 中检查分页与版式。")
    return warnings


def scan_document(template_path, categories, materials, special_category_ids=()):
    from recognition import family
    from document_toc import excluded_paragraphs, excluded_locator_ids
    selected = set(special_category_ids)
    enabled = [dict(c) for c in categories if
               (c.get("layer", "common") == "common" or c["id"] in selected)
               and c.get("enabled", True)]
    document = _read_template(template_path)
    toc_elements = excluded_paragraphs(document)
    records, skipped = _inspect(document, enabled, toc_elements)
    requirements = records[0].get('_requirements', {}) if records else {}
    headings = [_public(r) for r in records if r["_public"]]
    mappings = []
    locations = {c['id']: _special_candidates(records, c) for c in enabled if c.get('layer') == 'special'}
    paths = {r['id']: tuple(r['_path']) for r in records}
    families = {c['id']: family(c) for c in enabled}
    for cid, (locator, match) in list(locations.items()):
        if locator and sum(bool(other) and paths.get(other)==paths.get(locator) and families[other_cid]==families[cid]
                           for other_cid,(other,_) in locations.items()) > 1:
            # Different industries must not silently duplicate evidence at one location.
            match.update(status='ambiguous', reason='多个行业共用此章节，请核对各行业的插入位置')
    for cid, (locator, match) in list(locations.items()):
        if match['status'] == 'ambiguous':
            locations[cid] = ('', match)
    for category in enabled:
        hits = [r for r in records if category["id"] in r["category_ids"]]
        if category.get('layer') == 'special' and category['id'] != 'text_template':
            hits = [r for r in hits if r['id'] == locations[category['id']][0]]
        # An explicit token is the insertion point for its enclosing same-category section.
        tokens = [r for r in hits if r["kind"] == "placeholder"]
        hits = [r for r in hits if r["kind"] == "placeholder" or not any(
            t["_element"] is r["_element"] or any(t["_element"] is e or t["_element"] in e.iterdescendants() for e in r["_own"])
            for t in tokens)]
        for hit in hits or [None]:
            locator = hit["id"] if hit else ""
            row = {"id": "map_" + hashlib.sha256((str(category["id"]) + "|" + locator).encode()).hexdigest()[:20],
                   "category_id": category["id"], "category_name": category["name"],
                   "title": hit["title"] if hit else category["name"], "locator_id": locator,
                   "material_ids": [], "action": "missing", "confirmed": False, "note": "",
                   "required_years": _years(hit["title"]) if hit else [], "issues": [],
                   "status": "missing", "existing_content": hit["existing_content"] if hit else False}
            row['requirement'] = copy.deepcopy(requirements.get(category['id'], {'status':'not_detected','evidence':[]}))
            if category['id'] == 'audit_report':
                row['required_years'] = list(row['requirement'].get('required_years', []))
            if hit and hit.get('location_match') and category.get('layer', 'common') == 'common':
                row['location_match'] = copy.deepcopy(hit['location_match'])
            if not hit and category['id'] == 'authorization':
                row.update(title='被授权人身份证附件', location_match={'status': 'not_found', 'candidates': [],
                    'reason': '授权委托书正文由文字模板填写；原文未单独标注被授权人身份证附件位置，不插入图片。'})
            if category.get("layer") == "special":
                row.update(auto_match=False, action="fill",
                           note="", layer="special", location_match=(
                               copy.deepcopy(hit.get('location_match', {'status': 'matched', 'candidates': []}))
                               if hit and category['id'] == 'text_template' else locations[category['id']][1]))
            elif hit:
                from material_rules import partial_issue, automatic_ids
                candidates = []
                for material in materials:
                    if material.get("category_id") != category["id"]:
                        continue
                    probe = dict(row, material_ids=[material["id"]], action="fill")
                    checked = validate_mappings([probe], [material], headings, enabled)[0]
                    # In a multi-year column a single valid year may be a partial candidate.
                    severe = [s for s in checked["issues"] if not partial_issue(s)]
                    if not severe:
                        candidates.append(material)
                ids = automatic_ids(row, candidates, headings, enabled)
                if ids:
                    row.update(material_ids=ids, action='fill')
                elif candidates:
                    row['note'] = '存在重叠版本或需按截止日核对缴费期限，请人工选择材料。'
            mappings.append(row)
    mappings = validate_mappings(mappings, materials, headings, enabled)
    return {"headings": headings, "mappings": mappings, "requirements": requirements,
            "excluded_locator_ids": excluded_locator_ids(document, toc_elements, categories),
            "warnings": _document_warnings(document, records, skipped)}


def _material_years(material):
    result = set()
    for value in material.get("years") or []:
        if isinstance(value, bool):
            continue
        try:
            year = int(value)
            if str(value).strip() == str(year) and 1900 <= year <= 2200:
                result.add(year)
        except (ValueError, TypeError):
            pass
    return result


def _attachment_only(material):
    description = str(material.get("name", "")) + " " + str(material.get("description", ""))
    return bool(material.get("attachment_only") or "证件附件" in description or "身份证" in description.replace("身份证明", ""))


def material_form_feedback(material, file_count=0):
    """Intrinsic checks shared by the material form and save endpoint."""
    errors, warnings = {}, {}
    if not str(material.get('name') or '').strip():
        errors['name'] = '请填写材料名称。'
    if re.sub(r'\s+', '', str(material.get('subject') or '')) != EXPECTED_SUBJECT:
        errors['subject'] = f'材料主体须为{EXPECTED_SUBJECT}。'
    years = material.get('years', [])
    if not isinstance(years, list) or len(years) > 30 or any(type(y) is not int or not 1900 <= y <= 2200 for y in years):
        errors['years'] = '年度须为 1900—2200 之间的四位年份，多个年度用逗号分隔。'
    elif material.get('category_id') == 'audit_report' and not years:
        errors['years'] = '审计报告必须填写适用年度，不能只写在材料名称中。'
    value = str(material.get('valid_until') or '').strip()
    if not value:
        warnings['valid_until'] = '有效期未填写，请核对材料有效期；长期有效的材料可填写“长期”。'
    elif value in {'下一年度更新', '按年度更新', '每年更新'}:
        warnings['valid_until'] = '尚未提供明确有效期，请按项目要求核对。'
    elif value not in {'长期', '永久', '长期有效'}:
        try:
            if not re.fullmatch(r'\d{4}-\d{2}-\d{2}', value):
                raise ValueError()
            expiry = date.fromisoformat(value)
            if expiry < date.today():
                errors['valid_until'] = f'材料已过期（{value}），请上传有效材料并填写真实到期日期。'
            elif (expiry - date.today()).days <= 30:
                warnings['valid_until'] = '材料将在 30 天内到期，请核对投标截止日期。'
        except ValueError:
            errors['valid_until'] = '有效期格式无效，请填写真实日期 YYYY-MM-DD（如 2027-12-31）或“长期”。'
    mode = material.get('mode', 'page')
    if mode == 'text':
        if not str(material.get('content') or '').strip():
            errors['content'] = '请填写文字材料正文。'
    elif mode in {'page', 'side_by_side'}:
        if not file_count:
            errors['files'] = '请至少添加一张材料图片；已有材料可先取消编辑，在详情中追加图片。'
    else:
        errors['mode'] = '请选择有效的排版方式。'
    if material.get('category_id') in {'legal_identity', 'authorization'}:
        warnings['description'] = '身份证附件不能替代证明或授权正文；请同时核对正文、人员、签章和日期。'
    return {'errors': errors, 'warnings': warnings}


def _check_material(material, category_id, required, heading, special=False):
    errors, warnings = [], []
    name = str(material.get("name") or material.get("id") or "材料")
    if category_id == 'tax_social':
        from material_rules import tax_errors
        errors.extend(tax_errors(material, heading))
    if category_id == 'text_template':
        from material_rules import text_kind
        wanted, actual = text_kind(heading.get('title', '')), text_kind(name)
        if wanted and actual != wanted:
            errors.append(f'{name}：文书类型与目标栏目不符，请选择对应模板。')
    if material.get("category_id") != category_id and not (
            special and category_id in material.get("linked_category_ids", [])):
        errors.append(f"{name}：材料所属栏目不符。")
    subject = re.sub(r"\s+", "", str(material.get("subject") or ""))
    if special:
        if not subject or subject in {"待核对", "未知"}:
            errors.append(f"{name}：请先在材料库登记实际证明主体。")
        if material.get("archived"):
            errors.append(f"{name}：材料已归档停用，请改选可用材料。")
        warnings.append(f"{name}：请核对主体“{subject}”、项目要求、签章及授权范围。")
        if material.get("review_status", "pending") != "checked":
            warnings.append(f"{name}：材料库尚未标记核对通过，本次使用须逐项核验原件。")
    elif subject != EXPECTED_SUBJECT:
        errors.append(f"{name}：主体必须核实为{EXPECTED_SUBJECT}，当前为{subject or '未填写'}。")
    years = _material_years(material)
    if (required or category_id == "audit_report") and not years:
        errors.append(f"{name}：缺少年度元数据，名称和别名不能替代年度。")
    if required and years and not years.intersection(required):
        errors.append(f"{name}：年度与目标要求不符。")
    valid_until = str(material.get("valid_until") or "").strip()
    if not valid_until:
        warnings.append(f"{name}：有效期未填写，须人工核对。")
    elif valid_until in {"下一年度更新", "按年度更新", "每年更新"}:
        warnings.append(f"{name}：原台账标注“{valid_until}”，未提供明确有效期，须按本项目要求人工核对。")
    elif valid_until not in {"长期", "永久", "长期有效"}:
        try:
            if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", valid_until):
                raise ValueError()
            expiry = date.fromisoformat(valid_until)
            if expiry < date.today():
                target = warnings if special and str(category_id).endswith("01") else errors
                target.append(f"{name}：材料已过期（{expiry.isoformat()}）。" +
                              ("作为历史业绩使用时，请核对招标要求的业绩期限。" if target is warnings else ""))
            elif (expiry - date.today()).days <= 30:
                warnings.append(f"{name}：将在30天内到期，请核对投标截止日。")
        except ValueError:
            errors.append(f"{name}：有效期格式无效，请填写 YYYY-MM-DD 或长期。")
    mode = material.get("mode")
    if mode == "text":
        if not str(material.get("content") or "").strip():
            errors.append(f"{name}：文字内容为空。")
    elif mode in {"page", "side_by_side"}:
        files = material.get("files") or []
        if not files:
            errors.append(f"{name}：没有材料图片。")
        for file in files:
            path = Path(str(file.get("path") or ""))
            label = str(file.get("name") or "图片")
            if not path.is_file() or path.suffix.lower() not in _EXTENSIONS:
                errors.append(f"{name}：图片不可用（{label}）。")
                continue
            try:
                rotation = int(file.get("rotation") or 0)
                if rotation % 90:
                    raise ValueError("rotation")
                with Image.open(path) as image:
                    image.verify()
            except Exception:
                errors.append(f"{name}：图片损坏或旋转值无效（{label}）。")
    else:
        errors.append(f"{name}：插入方式无效。")
    if category_id in {"legal_identity", "authorization"} and _attachment_only(material):
        body_key = "legal_body_present" if category_id == "legal_identity" else "authorization_body_present"
        if category_id == 'authorization' or (category_id == 'legal_identity' and heading.get('identity_attachment_target')):
            warnings.append(f"{name}：本次仅装配身份证附件，不代表身份证明或授权正文已填写；人员、日期和签章由商务核对。")
        elif not heading.get(body_key):
            errors.append(f"{name}：身份证附件不能替代证明/授权正文，目标节未识别到相应文书正文。")
        else:
            warnings.append(f"{name}：仅作为身份证附件；原证明/授权正文、人员、日期和签章须单独核对。")
        if category_id == 'authorization' and heading.get('authorization_text'):
            import authorization
            _, original_agent = authorization._names(re.sub(r'\s+', '', heading['authorization_text']))
            selected_agent = authorization.person(material)
            if original_agent and selected_agent and original_agent != selected_agent:
                errors.append(f'原授权正文代理人为{original_agent}，所选证件为{selected_agent}；请核对人员，原文保留。')
    return errors, warnings


def validate_mappings(mappings, materials, headings, categories=None):
    by_material = {m["id"]: m for m in materials}
    if len(by_material) != len(materials):
        raise ValueError("材料 ID 重复，无法可靠核对。")
    by_heading = {h["id"]: h for h in headings}
    by_category = {c["id"]: c for c in categories} if categories is not None else None
    seen_ids, used_locations = set(), {}
    result = []
    for original in mappings:
        row = copy.deepcopy(original)
        mid = row.get("id")
        if not mid or mid in seen_ids:
            raise ValueError("栏目记录 ID 缺失或重复，请重新扫描模板。")
        seen_ids.add(mid)
        errors, warnings = [], []
        category_id = row.get("category_id")
        category = by_category.get(category_id) if by_category is not None else None
        special = bool(category and category.get("layer") == "special")
        row["layer"] = "special" if special else "common"
        if by_category is not None:
            if (not category or category.get("layer", "common") not in {"common", "special"}
                    or not category.get("enabled", True)
                    or original.get("layer", "common") != category.get("layer", "common")):
                errors.append("该栏目已停用、分类已变更或未纳入本项目，请重新核对分类。")
        heading = by_heading.get(row.get("locator_id"))
        # Retire legacy text-generation settings; old confirmations covered a
        # different output and must be reviewed once under image-only filling.
        legacy_authorization = row.pop('authorization', None)
        if isinstance(legacy_authorization, dict) and legacy_authorization.get('enabled') is True:
            row['confirmed'] = False
        row.pop('authorization_preview', None)
        if heading:
            row["title"] = heading["title"]
            row["existing_content"] = bool(heading.get("existing_content"))
            if heading.get("category_ids") and category_id not in heading["category_ids"]:
                errors.append("选定位置与栏目类别不符。")
            if not heading.get("category_ids"):
                warnings.append("当前为人工指定位置，请核对栏目上下文。")
            if heading.get('kind') == 'common_appendix':
                warnings.append('原文有材料要求但无独立位置，将在文末新增通用材料区；请核对是否需要提交及材料适用性。')
        requirement = (heading or {}).get('requirement') or row.get('requirement') or {}
        row["required_years"] = (requirement.get('required_years') or _years(heading["title"])) if heading else []
        # Alternative financial years are not a requirement to supply every year.
        year_options = requirement.get('year_options', []) if category_id == 'audit_report' else []
        if year_options:
            row['required_years'] = []
        if category_id == 'tax_social' and heading:
            from material_rules import period_issue
            # Shown before any material is chosen, not only on a selected one.
            if period_issue(heading):
                warnings.append(period_issue(heading))
        action = row.get("action")
        if action not in {"fill", "missing", "skip"}:
            errors.append("处理方式无效。")
        if action == "fill":
            if not heading:
                errors.append("未找到有效填充位置，请选择原模板中的位置。")
            location = row.get("locator_id")
            if location and location in used_locations:
                errors.append("同一位置被多个栏目重复使用，请核对映射。")
                used_locations[location]["issues"].append("同一位置被多个栏目重复使用，请核对映射。")
                used_locations[location]["status"] = "blocked"
            if location:
                used_locations[location] = row
            ids = row.get("material_ids") or []
            if len(ids) != len(set(ids)):
                errors.append("同一栏目重复选择同一材料。")
            if not ids:
                errors.append("存在多个材料版本，请人工选择。" if sum(m.get("category_id") == category_id for m in materials) > 1 else "尚未选择材料。")
            if category_id == 'authorization' and heading and heading.get('authorization_has_images'):
                errors.append('原授权书已含证件图片，系统保留原件；请先核对已有附件，避免重复或混用人员。')
            covered = set()
            valid_materials = []
            for material_id in ids:
                material = by_material.get(material_id)
                if not material:
                    errors.append("指定材料不存在或已删除。")
                    continue
                target = heading or {}
                check_category = category_id
                if category_id == 'authorization':
                    if material.get('mode') not in {'page', 'side_by_side'}:
                        errors.append('授权书栏目仅支持填充图片材料。')
                    if material.get('category_id') == 'legal_identity':
                        check_category = 'legal_identity'
                        target = dict(target, identity_attachment_target=True)
                accepted_years = sorted({y for option in year_options for y in option}) if year_options else row['required_years']
                errs, warns = _check_material(material, check_category, accepted_years, target, special)
                errors.extend(errs)
                warnings.extend(warns)
                if not errs:
                    covered.update(_material_years(material))
                    valid_materials.append(material)
            if category_id == 'tax_social' and heading:
                from material_rules import tax_coverage_errors
                errors.extend(tax_coverage_errors(valid_materials, heading))
            missing_years = sorted(set(row["required_years"]) - covered)
            if missing_years:
                errors.append("缺少要求年度：" + "、".join(map(str, missing_years)) + "。")
            if year_options and not any(set(option) <= covered for option in year_options):
                errors.append('缺少可接受年度：' + ' 或 '.join('、'.join(map(str, option)) for option in year_options) + '。')
            row["status"] = "blocked" if errors else "ready"
        elif action == "skip":
            row["status"] = "skipped" if not errors else "blocked"
            warnings.append("本栏目人工跳过，不视为材料齐全。")
        else:
            row["status"] = "missing" if not errors else "blocked"
            errors.append("本栏目待补材料。" if heading else "模板未识别到本栏目位置，需人工核对；初稿将在末尾列出。")
        if category_id == "audit_report" and not row["required_years"] and not year_options:
            warnings.append("目标审计栏目未注明年度，不默认采用2023/2024年；请按本项目要求核对。")
        if row.get("existing_content"):
            warnings.append("本节已有内容，将保留并追加，请检查旧版或重复材料。")
        row["issues"], row["warnings"] = _unique(errors), _unique(warnings)
        row["confirmed"] = row.get("confirmed") is True
        result.append(row)
    return result


def _paragraph_after(anchor, document, text="", red=False):
    element = OxmlElement("w:p")
    anchor.addnext(element)
    paragraph = Paragraph(element, document._body)
    formatting = paragraph.paragraph_format
    formatting.space_before = formatting.space_after = Pt(0)
    formatting.first_line_indent = Pt(0)
    formatting.left_indent = formatting.right_indent = Pt(0)
    formatting.line_spacing_rule = WD_LINE_SPACING.SINGLE
    if text:
        run = paragraph.add_run(text)
        if red:
            run.font.color.rgb = RGBColor(192, 0, 0)
            run.font.bold = True
    return paragraph


def _image(file):
    try:
        with Image.open(file["path"]) as source:
            image = ImageOps.exif_transpose(source).copy()
        rotation = int(file.get("rotation") or 0) % 360
        if rotation:
            image = image.rotate(-rotation, expand=True)
        if image.mode not in {"RGB", "RGBA", "L", "LA"}:
            image = image.convert("RGBA")
        buffer = io.BytesIO()
        image.save(buffer, format="PNG")
        buffer.seek(0)
        return buffer, image.size
    except Exception as exc:
        raise ValueError("图片读取失败：" + str(file.get("name") or "材料图片")) from exc


def _add_picture(paragraph, file, max_width, max_height):
    buffer, (width, height) = _image(file)
    ratio = min(max_width / width, max_height / height)
    paragraph.add_run().add_picture(buffer, width=Emu(int(width * ratio)), height=Emu(int(height * ratio)))


def _remove_token(record):
    element = record["_element"]
    tokens = list(_TOKEN.finditer(_text(element)))
    occurrence = record["_occurrence"]
    if occurrence >= len(tokens) or tokens[occurrence].group(1).strip() != record["title"]:
        raise ValueError("占位符已变化，请重新扫描原模板。")
    start, end = tokens[occurrence].span()
    cursor = 0
    for text_node in element.iter(qn("w:t")):
        value = text_node.text or ""
        node_end = cursor + len(value)
        if cursor < end and node_end > start:
            text_node.text = value[:max(0, start - cursor)] + value[max(0, end - cursor):]
        cursor = node_end


def _vacant_cells(document, record, file_count):
    """Share the exact empty-table target between export and location preview."""
    if record["kind"] == "placeholder":
        return None
    for block in record["_own"]:
        if block.tag != qn("w:tbl") or _text(block).strip() or _drawing(block):
            continue
        table = Table(block, document._body)
        cells, seen = [], set()
        for row in table.rows:
            for cell in row.cells:
                if cell._tc not in seen:
                    cells.append(cell)
                    seen.add(cell._tc)
        if len(cells) >= file_count:
            return cells
    return None


def _insert_materials(document, record, selected):
    anchor = record["_anchor"]
    if record["kind"] == "placeholder":
        _remove_token(record)
        anchor = record["_element"]
    max_width, max_height = _page_size(document, record["_path"][0])
    in_cell = record["_inside_table"]
    if in_cell:
        for ancestor in record["_element"].iterancestors():
            if ancestor.tag == qn("w:tc"):
                width_node = ancestor.find("./" + qn("w:tcPr") + "/" + qn("w:tcW"))
                if width_node is not None and width_node.get(qn("w:type")) == "dxa":
                    max_width = min(max_width, int(width_node.get(qn("w:w"), "5000")) * 635 - int(Cm(0.4)))
                break
    page_material = False
    for material in selected:
        if material["mode"] == "text":
            for line in str(material["content"]).splitlines() or [""]:
                paragraph = _paragraph_after(anchor, document, line)
                anchor = paragraph._p
            continue
        files = material["files"]
        if material["mode"] == "side_by_side":
            # Use a genuinely empty picture area in this exact section.  The
            # two cells keep front/back images together even when the template
            # has unusual character spacing or paragraph inheritance.
            vacant_cells = _vacant_cells(document, record, len(files))
            if vacant_cells is not None:
                for file, cell in zip(files, vacant_cells):
                    paragraph = cell.paragraphs[0]
                    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    fmt = paragraph.paragraph_format
                    fmt.first_line_indent = fmt.left_indent = fmt.right_indent = Pt(0)
                    fmt.space_before = fmt.space_after = Pt(0)
                    fmt.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    fmt.keep_together = True
                    width = int(cell.width or max_width // 2) - int(Cm(0.6))
                    _add_picture(paragraph, file, width, int(Cm(5.5)))
                continue
            for index in range(0, len(files), 2):
                paragraph = _paragraph_after(anchor, document)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                paragraph.paragraph_format.keep_together = True
                pair = files[index:index + 2]
                each_width = (max_width - int(Cm(0.5))) // len(pair)
                for i, file in enumerate(pair):
                    if i:
                        paragraph.add_run("  ").font.size = Pt(1)
                    _add_picture(paragraph, file, each_width, int(Cm(5.5)))
                anchor = paragraph._p
        else:
            for file in files:
                first_appendix_image = record['kind'] == 'common_appendix' and anchor is record['_anchor']
                paragraph = _paragraph_after(anchor, document)
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                if not in_cell:
                    paragraph.paragraph_format.page_break_before = not first_appendix_image
                    page_material = True
                image_height = max_height - int(Cm(4.8) if first_appendix_image else Cm(1.2))
                _add_picture(paragraph, file, max_width - int(Cm(0.2)), min(image_height, int(Cm(20)) if in_cell else max_height))
                anchor = paragraph._p
    if page_material:
        following = anchor.getnext()
        already_break = following is None or following.tag == qn("w:sectPr")
        if following is not None:
            already_break = already_break or any(n.get(qn("w:type")) == "page" for n in following.iter(qn("w:br")))
            already_break = already_break or following.find(".//" + qn("w:sectPr")) is not None
            if following.tag == qn("w:p"):
                already_break = already_break or bool(Paragraph(following, document._body).paragraph_format.page_break_before)
        if not already_break:
            paragraph = _paragraph_after(anchor, document)
            paragraph.paragraph_format.page_break_before = True


def _update_fields(document):
    settings = document.settings.element
    element = settings.find(qn("w:updateFields"))
    if element is None:
        element = OxmlElement("w:updateFields")
        after = {"hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids", "mathPr", "themeFontLang", "clrSchemeMapping", "doNotAutoCompressPictures", "decimalSymbol", "listSeparator"}
        next_element = next((e for e in settings if e.tag.split("}")[-1] in after), None)
        if next_element is None:
            settings.append(element)
        else:
            next_element.addprevious(element)
    element.set(qn("w:val"), "true")


def export_document(template_path, output_path, mappings, materials, mode="draft", categories=None, progress=None):
    if mode not in {"draft", "reviewed"}:
        raise ValueError("导出模式无效。")
    if Path(template_path).resolve() == Path(output_path).resolve():
        raise ValueError("不能覆盖原始 Word 模板。")
    if not mappings:
        raise ValueError("没有可导出的栏目，请先扫描模板。")
    if mode == "reviewed" and any(row.get("confirmed") is not True for row in mappings):
        raise ValueError("所有栏目必须先逐项确认后才能导出。")
    report = progress or (lambda percent, stage: None)
    report(5, "正在读取模板")
    document = _read_template(template_path)
    if categories is None:
        # Compatibility for callers without a category store. Store-backed callers
        # must pass current settings to preserve aliases and enforce disablement.
        categories = list({r["category_id"]: {"id": r["category_id"], "name": r.get("category_name") or r["title"], "aliases": [], "enabled": True, "layer": "common"} for r in mappings}.values())
        defaults = {c["id"]: c for c in DEFAULT_CATEGORIES}
        for category in categories:
            if category["id"] in defaults:
                category["aliases"] = [defaults[category["id"]]["name"], *defaults[category["id"]]["aliases"]]
    mapping_categories = {r["category_id"] for r in mappings}
    enabled = [c for c in categories if
               (c.get("layer", "common") == "common" or c["id"] in mapping_categories)
               and c.get("enabled", True)]
    records, skipped = _inspect(document, enabled)
    by_locator = {r["id"]: r for r in records}
    public = [_public(r) for r in records]
    checked = validate_mappings(mappings, materials, public, categories)
    if mode == 'reviewed' and any(row.get('confirmed') is not True for row in checked):
        raise ValueError('填充方式已变更，请重新核对并确认栏目后再导出。')
    invalid = [r for r in checked if r["status"] == "blocked"]
    if invalid:
        raise ValueError("不能导出：" + "；".join(r["title"] + "：" + "、".join(r["issues"]) for r in invalid))
    if mode == "reviewed" and any(r["status"] != "ready" for r in checked):
        raise ValueError("复核版要求所有栏目材料齐全；存在待补或跳过栏目，请导出待补初稿。")
    appendices = [by_locator[row['locator_id']] for row in checked
                  if row.get('locator_id') in by_locator
                  and by_locator[row['locator_id']]['kind'] == 'common_appendix'
                  and row['action'] != 'skip']
    if appendices:
        def appendix_heading(text, level):
            # Converted tenders may omit built-in heading styles. Check before
            # add_heading, which otherwise leaves a half-created paragraph.
            if f'Heading {level}' in document.styles:
                paragraph = document.add_heading(text, level=level)
            else:
                paragraph = document.add_paragraph(text)
                outline = OxmlElement('w:outlineLvl')
                outline.set(qn('w:val'), str(level-1))
                paragraph._p.get_or_add_pPr().append(outline)
                for run in paragraph.runs:
                    run.font.bold = True
                    run.font.size = Pt(16 if level == 1 else 14)
            for run in paragraph.runs:
                run.font.color.rgb = RGBColor(0, 0, 0)
            return paragraph
        document.add_page_break()
        appendix_heading('通用材料附件（请商务核对）', 1)
        document.add_paragraph('依据原文材料要求补充以下附件。原稿未提供可靠独立位置；正文、人员、日期、签章及材料适用性仍须商务核对。')
        for index, record in enumerate(appendices):
            heading = appendix_heading(record['title'].replace('文末新增材料区 · ', ''), 2)
            heading.paragraph_format.page_break_before = index > 0
            heading.paragraph_format.keep_with_next = True
            record['_element'] = record['_anchor'] = heading._p
    by_material = {m["id"]: m for m in materials}
    warnings = _document_warnings(document, records, skipped)
    warnings.append("Word 目录域已设置打开时更新；人工写死的目录和页码不会自动更新，请更新目录并核对。")
    results = []
    # Resolve every target against the original tree. Reverse token order keeps
    # multiple same-paragraph token occurrences stable while they are removed.
    ordered = sorted(checked, key=lambda r: (by_locator.get(r.get("locator_id"), {}).get("_path", (999999,)), by_locator.get(r.get("locator_id"), {}).get("_occurrence", -1)), reverse=True)
    report(20, "正在填入材料")
    for index, row in enumerate(ordered):
        report(20 + int(65 * index / max(1, len(ordered))), f"正在填入材料 {index + 1}/{len(ordered)}")
        record = by_locator.get(row.get("locator_id"))
        selected = [by_material[i] for i in row.get("material_ids", [])] if row["action"] == "fill" else []
        if row["status"] == "ready":
            _insert_materials(document, record, selected)
        else:
            label = "待补" if row["status"] == "missing" else "人工跳过"
            text = f"【{label}：{row['title']}】" + (" " + row["note"] if row.get("note") else "")
            if record:
                if record["kind"] == "placeholder":
                    _remove_token(record)
                _paragraph_after(record["_anchor"], document, text, red=True)
            else:
                paragraph = document.add_paragraph()
                run = paragraph.add_run(text + "（原模板未定位，在此列出待处理项）")
                run.font.color.rgb = RGBColor(192, 0, 0)
                run.font.bold = True
        warnings.extend(row.get("warnings", []))
        results.append({"mapping_id": row["id"], "category_id": row["category_id"], "title": row["title"],
                        "locator_id": row.get("locator_id", ""), "status": row["status"], "action": row["action"],
                        "material_ids": [m["id"] for m in selected], "material_names": [m["name"] for m in selected],
                        "file_count": sum(len(m.get("files", [])) for m in selected if m["mode"] != "text"),
                        "issues": row["issues"], "warnings": row.get("warnings", []), "note": row.get("note", "")})
    if mode == "draft":
        notice = OxmlElement("w:p")
        document.element.body.insert(0, notice)
        run = Paragraph(notice, document._body).add_run("待补初稿｜仅供内部核对，不可直接投标。请处理红色标记并检查原模板残留信息、签章及页码。")
        run.font.color.rgb = RGBColor(192, 0, 0)
        run.font.bold = True
    report(88, "正在保存 Word 文件")
    _update_fields(document)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = None
    try:
        with tempfile.NamedTemporaryFile(dir=destination.parent, suffix=".docx", delete=False) as file:
            temporary = file.name
        document.save(temporary)
        report(95, "正在校验文件并生成核对报告")
        Document(temporary)
        os.replace(temporary, destination)
    except Exception as exc:
        if temporary and os.path.exists(temporary):
            os.unlink(temporary)
        raise ValueError("导出 Word 失败，未生成完整结果。") from exc
    order = {r["id"]: index for index, r in enumerate(checked)}
    results.sort(key=lambda r: order[r["mapping_id"]])
    return {"results": results, "filled_count": sum(r["status"] == "ready" for r in checked),
            "missing_count": sum(r["status"] == "missing" for r in checked),
            "skipped_count": sum(r["status"] == "skipped" for r in checked), "warnings": _unique(warnings)}
